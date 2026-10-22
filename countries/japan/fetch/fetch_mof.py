#!/usr/bin/env python3
"""Refresh the Ministry of Finance series into sidecar CSVs.

Usage:  python3 -I countries/japan/fetch/fetch_mof.py

Five publications, all on www.mof.go.jp, none behind a key:

  fiscal/tax_revenue_from_mof.csv      zeisyu.xls — general-account tax revenue since FY1979
  fiscal/general_account_from_mof.csv  財政統計 tables 1, 4 and 20 (budget and settlement)
  fiscal/central_debt_from_mof.csv     国債及び借入金並びに政府保証債務現在高 (quarterly)
  markets/jgb_from_mof.csv             国債金利情報 jgbcm_all.csv + jgbcm.csv (daily yields)
  international/reserves_from_mof.csv  外貨準備等の状況 historical.csv (monthly)

Things that are easy to get wrong here, and how each is handled:

**Japanese era dates.** The JGB file dates every row in wareki
(S49.9.24, H18.1.4, R8.9.30); the reserves file writes 平成18年 on the
first month of each year only; the settlement tables key on 平成/令和 year
numbers with 元 for year one; the debt tables head their columns R3.9末.
All go through one explicit table (WAREKI below) that also knows where each
era starts and ends — S64 is 1–7 January 1989 and H31 ends 30 April 2019
— and rejects a date outside its era rather than guessing.

**Units.** zeisyu is in 億円 (¥100 million), table 4 in 百万円 (¥ million),
tables 1 and 20 in 千円 (¥ thousand), the debt tables in 億円, reserves in
US$ million. The unit cell of every table is asserted, and every yen
amount is converted explicitly to ¥ BILLION in the output.

**Fiscal years** run April–March and are keyed on the STARTING year
(fiscal_year 2024 = April 2024 – March 2025), the Canada/UK convention.

**Which tax number.** zeisyu.xls runs to the latest fiscal year whose
accounts have closed, but nothing in it says whether its last row is the
settlement (決算) or a budget estimate. The fetcher reads the latest
決算概要 page, takes the settlement tax revenue from its 予算決算比較表
PDF (via pdftotext) and asserts that zeisyu's last row equals it.

**Three debts.** central_debt_from_mof.csv is CENTRAL government at FACE
value. The BoJ flow-of-funds general-government liabilities (market value)
and the IMF gross general-government debt are other objects, in other
files; none of the three is ever spliced onto another.

**The debt history.** The live suii.xls covers only the last five years.
Earlier editions survive in the National Diet Library's web archive
(WARP), which the live page links to: the fetcher follows those links,
reads each archived suii.xls, and for the 2006–2009 quarters (older than
any archived workbook) reads the per-quarter HTML pages that MOF published
before suii.xls existed, from the archived 2011 index. Every quarter is
taken from the NEWEST edition that has it; where editions overlap, they are
asserted to agree, and each row says which edition it came from.

Needs openpyxl and poppler's `pdftotext` (the settlement-summary PDF). The
legacy .xls files (zeisyu.xls, every suii.xls) are read by the small BIFF8
reader below, so xlrd is not needed.

Outputs are `_from_mof` sidecars. Nothing here overwrites a curated file.
"""
from __future__ import annotations

import csv
import html as htmllib
import io
import json
import re
import shutil
import statistics
import struct
import subprocess
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from datetime import date
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
CACHE = ROOT / "cache"
MOF = "https://www.mof.go.jp"
WARP = "https://warp.ndl.go.jp"
UA = {"User-Agent": "Mozilla/5.0 (worldreports data fetcher)"}
Y0, Y1 = 2006, 2026

URL = {
    "zeisyu": f"{MOF}/tax_policy/summary/condition/zeisyu.xls",
    "stats_index": f"{MOF}/policy/budget/reference/statistics/data.htm",
    "account_index": f"{MOF}/policy/budget/budger_workflow/account/index.html",
    "gbb_index": f"{MOF}/jgbs/reference/gbb/data.htm",
    "jgb_all": f"{MOF}/jgbs/reference/interest_rate/data/jgbcm_all.csv",
    "jgb_now": f"{MOF}/jgbs/reference/interest_rate/jgbcm.csv",
    "reserves": f"{MOF}/policy/international_policy/reference/official_reserve_assets/historical.csv",
    "imf_debt": "https://www.imf.org/external/datamapper/api/v1/GGXWDG_NGDP/JPN",
    "wb_reserves": ("https://api.worldbank.org/v2/country/JPN/indicator/"
                    "FI.RES.TOTL.CD?format=json&date=2024"),
    "oecd_10y": "https://fred.stlouisfed.org/graph/fredgraph.csv?id=IRLTLT01JPM156N",
}

# The archived index from before suii.xls existed, when MOF published one
# HTML page per quarter-end. No archived page links to it (the 2016
# capture's "older" link is absent), so it is reached by asking WARP for
# its capture nearest this date; the code then asserts what it finds.
WARP_QUARTERLY_INDEX = f"{WARP}/web/20110301000000/www.mof.go.jp/jgbs/reference/gbb/data.htm"

# Nominal GDP, calendar 2024, ¥bn: ESRI QE 2026Q2-2 (gaku-mcy2622), as
# recorded in DISCOVERY.md. Used only as the denominator of a magnitude
# assertion against the IMF ratio — never written to a CSV.
ESRI_NGDP_2024_BN = 636893.3
WARP_PAUSE = 6          # seconds between NDL archive requests


# --- wareki ---------------------------------------------------------------

WAREKI = {  # era: (offset, first day, last day)
    "M": (1867, date(1868, 1, 25), date(1912, 7, 29)),
    "T": (1911, date(1912, 7, 30), date(1926, 12, 24)),
    "S": (1925, date(1926, 12, 25), date(1989, 1, 7)),
    "H": (1988, date(1989, 1, 8), date(2019, 4, 30)),
    "R": (2018, date(2019, 5, 1), date(9999, 12, 31)),
}
ERA_NAME = {"明治": "M", "大正": "T", "昭和": "S", "平成": "H", "令和": "R"}


def era_year(era: str, n: str | int | float) -> int:
    """Western year of era year n ('元' is year one)."""
    era = ERA_NAME.get(era, era)
    if era not in WAREKI:
        raise ValueError(f"unknown era {era!r}")
    n = 1 if str(n).strip() in ("元", "元年") else int(float(n))
    return WAREKI[era][0] + n


def wareki_date(text: str) -> date:
    """'S49.9.24' / 'H18.1.4' / 'R8.9.30' to a date, checked against the era."""
    m = re.fullmatch(r"([MTSHR])(\d{1,2})\.(\d{1,2})\.(\d{1,2})", text.strip())
    if not m:
        raise ValueError(f"not a wareki date: {text!r}")
    era, y, mo, d = m.group(1), *map(int, m.groups()[1:])
    out = date(era_year(era, y), mo, d)
    first, last = WAREKI[era][1:]
    if not first <= out <= last:
        raise ValueError(f"{text!r} = {out} is outside the {era} era")
    return out


def _wareki_self_test() -> None:
    for text, want in (("S49.9.24", date(1974, 9, 24)), ("S64.1.6", date(1989, 1, 6)),
                       ("H1.1.9", date(1989, 1, 9)), ("H18.1.4", date(2006, 1, 4)),
                       ("H31.4.26", date(2019, 4, 26)), ("R1.5.7", date(2019, 5, 7)),
                       ("R8.9.30", date(2026, 9, 30))):
        if wareki_date(text) != want:
            raise RuntimeError(f"wareki: {text} -> {wareki_date(text)}, expected {want}")
    for bad in ("S64.1.9", "H31.5.7", "R1.4.1"):
        try:
            wareki_date(bad)
        except ValueError:
            continue
        raise RuntimeError(f"wareki: {bad} should have been rejected")
    if era_year("令和", "元") != 2019 or era_year("平成", 18) != 2006:
        raise RuntimeError("wareki: era_year table is wrong")


# --- fetching -------------------------------------------------------------

def get(url: str, retries=4) -> bytes:
    last = None
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, headers=UA)
            with urllib.request.urlopen(req, timeout=180) as r:
                return r.read()
        except urllib.error.HTTPError as e:
            last = f"HTTP {e.code}"
            if e.code == 404:
                break
            # WARP answers 403 "high traffic" when asked too fast: back off hard.
            time.sleep(45 * (attempt + 1) if e.code in (403, 429) else 3 * (attempt + 1))
        except Exception as e:  # noqa: BLE001
            last = e
            time.sleep(3 * (attempt + 1))
    raise RuntimeError(f"mof: {url}: {last}")


def cached(name: str, url: str, shape, pause=1.0) -> bytes:
    """Fetch once into cache/mof_<name>; `shape(blob)` must pass first.

    A cached error page survives the retry that would have fixed it, so
    nothing is written to the cache until it looks like what it claims.
    """
    CACHE.mkdir(exist_ok=True)
    path = CACHE / f"mof_{name}"
    if path.exists():
        return path.read_bytes()
    blob = get(url)
    problem = shape(blob)
    if problem:
        raise RuntimeError(f"mof: {url}: {problem}")
    path.write_bytes(blob)
    time.sleep(pause)
    return blob



# --- minimal BIFF8 (.xls) reader ----------------------------------------------
#
# zeisyu.xls and every edition of suii.xls are Excel 97 .xls, which openpyxl
# cannot read, and xlrd is not importable under `python3 -I`. This is the
# reader from fetch_isa.py (fetchers are standalone scripts, so it is copied,
# not imported): it reads shared strings, NUMBER/RK/MULRK, LABEL and cached
# formula numbers straight from the OLE compound file. Checked cell-for-cell
# against xlrd 2.0.2 on zeisyu.xls and the four suii.xls editions (live and
# WARP 2016/2019/2023) at consolidation, 2026-10-22.

def _cfb_stream(blob: bytes, want: str) -> bytes:
    if blob[:8] != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise ValueError("not an OLE compound file")
    ssz = 1 << struct.unpack_from("<H", blob, 30)[0]
    mssz = 1 << struct.unpack_from("<H", blob, 32)[0]
    dir_start, = struct.unpack_from("<I", blob, 48)
    cutoff, = struct.unpack_from("<I", blob, 56)
    minifat_start, = struct.unpack_from("<I", blob, 60)
    difat_start, difat_n = struct.unpack_from("<II", blob, 68)

    def sector(i):
        return blob[(i + 1) * ssz:(i + 2) * ssz]

    difat = list(struct.unpack_from("<109I", blob, 76))
    s = difat_start
    for _ in range(difat_n):
        if s >= 0xFFFFFFFA:
            break
        vals = struct.unpack(f"<{ssz // 4}I", sector(s))
        difat += vals[:-1]
        s = vals[-1]
    fat = []
    for f in difat:
        if f < 0xFFFFFFFA:
            fat += struct.unpack(f"<{ssz // 4}I", sector(f))

    def chain(start):
        out, s, seen = [], start, set()
        while s < 0xFFFFFFFA:
            if s in seen:
                raise ValueError("FAT loop")
            seen.add(s)
            out.append(sector(s))
            s = fat[s]
        return b"".join(out)

    d = chain(dir_start)
    entries = []
    for off in range(0, len(d), 128):
        e = d[off:off + 128]
        nlen, = struct.unpack_from("<H", e, 64)
        name = e[:max(nlen - 2, 0)].decode("utf-16-le", "replace")
        start, size = struct.unpack_from("<IQ", e, 116)
        entries.append((name, e[66], start, size & 0xFFFFFFFF))
    for name, typ, start, size in entries:
        if name == want and typ == 2:
            if size >= cutoff:
                return chain(start)[:size]
            ministream = chain(entries[0][2])
            mf = chain(minifat_start) if minifat_start < 0xFFFFFFFA else b""
            minifat = struct.unpack(f"<{len(mf) // 4}I", mf)
            out, s = [], start
            while s < 0xFFFFFFFA:
                out.append(ministream[s * mssz:(s + 1) * mssz])
                s = minifat[s]
            return b"".join(out)[:size]
    raise ValueError(f"no {want} stream")


def _rk(v: int) -> float:
    if v & 2:
        x = float(v >> 2 if v < 0x80000000 else (v >> 2) - (1 << 30))
    else:
        x = struct.unpack("<d", struct.pack("<Q", (v & 0xFFFFFFFC) << 32))[0]
    return x / 100 if v & 1 else x


def _parse_sst(chunks: list[bytes]) -> list[str]:
    ci, pos, data = 0, 8, chunks[0]
    _, n = struct.unpack_from("<II", data, 0)
    out = []
    for _ in range(n):
        if pos >= len(data):
            ci, pos = ci + 1, 0
            data = chunks[ci]
        cch, = struct.unpack_from("<H", data, pos)
        flags = data[pos + 2]
        pos += 3
        rich = ext = 0
        if flags & 0x08:
            rich, = struct.unpack_from("<H", data, pos)
            pos += 2
        if flags & 0x04:
            ext, = struct.unpack_from("<I", data, pos)
            pos += 4
        wide, parts, left = flags & 1, [], cch
        while left:
            if pos >= len(data):
                ci += 1
                data = chunks[ci]
                wide, pos = data[0] & 1, 1
            take = min(left, (len(data) - pos) // (2 if wide else 1))
            nb = take * (2 if wide else 1)
            seg = data[pos:pos + nb]
            parts.append(seg.decode("utf-16-le") if wide
                         else seg.decode("latin-1"))
            pos += nb
            left -= take
        skip = 4 * rich + ext
        while skip:
            if pos >= len(data):
                ci, pos = ci + 1, 0
                data = chunks[ci]
            t = min(skip, len(data) - pos)
            pos += t
            skip -= t
        out.append("".join(parts))
    return out


def xls_rows(blob: bytes, sheet: int = 0) -> list[list]:
    wb = None
    for name in ("Workbook", "Book"):
        try:
            wb = _cfb_stream(blob, name)
            break
        except ValueError:
            continue
    if wb is None:
        raise ValueError("no Workbook stream")
    recs, pos = [], 0
    while pos + 4 <= len(wb):
        rid, ln = struct.unpack_from("<HH", wb, pos)
        recs.append((rid, pos, wb[pos + 4:pos + 4 + ln]))
        pos += 4 + ln
    sst, sheets = [], []
    for k, (rid, pos, data) in enumerate(recs):
        if rid == 0x0085:
            sheets.append(struct.unpack_from("<I", data, 0)[0])
        elif rid == 0x00FC:
            chunks, j = [data], k + 1
            while j < len(recs) and recs[j][0] == 0x003C:
                chunks.append(recs[j][2])
                j += 1
            sst = _parse_sst(chunks)
        elif rid == 0x000A:
            break
    start, cells, on = sheets[sheet], {}, False
    for rid, pos, data in recs:
        if pos == start:
            on = True
            continue
        if not on:
            continue
        if rid == 0x000A:
            break
        if rid == 0x00FD:
            r, c, _, i = struct.unpack_from("<HHHI", data, 0)
            cells[(r, c)] = sst[i]
        elif rid == 0x0203:
            r, c, _, v = struct.unpack_from("<HHHd", data, 0)
            cells[(r, c)] = v
        elif rid == 0x027E:
            r, c, _, v = struct.unpack_from("<HHHI", data, 0)
            cells[(r, c)] = _rk(v)
        elif rid == 0x00BD:
            r, c0 = struct.unpack_from("<HH", data, 0)
            for i in range((len(data) - 6) // 6):
                _, v = struct.unpack_from("<HI", data, 4 + 6 * i)
                cells[(r, c0 + i)] = _rk(v)
        elif rid == 0x0006:
            r, c = struct.unpack_from("<HH", data, 0)
            if data[12:14] != b"\xff\xff":
                cells[(r, c)] = struct.unpack_from("<d", data, 6)[0]
        elif rid == 0x0204:
            r, c, _, n = struct.unpack_from("<HHHH", data, 0)
            raw = data[9:]
            cells[(r, c)] = (raw[:2 * n].decode("utf-16-le") if data[8] & 1
                             else raw[:n].decode("latin-1"))
    if not cells:
        return []
    rows = [[None] * (max(c for _, c in cells) + 1)
            for _ in range(max(r for r, _ in cells) + 1)]
    for (r, c), v in cells.items():
        rows[r][c] = v
    return rows


class XlsSheet:
    """The first sheet of an .xls, shaped like the xlrd sheet this file used
    to read: `nrows`, and `row_values(i)` with empty cells as "" (xls_rows
    already gives every row the same width)."""

    def __init__(self, blob: bytes):
        self._rows = [["" if v is None else v for v in r] for r in xls_rows(blob)]
        self.nrows = len(self._rows)

    def row_values(self, i: int) -> list:
        return self._rows[i]


def is_xls(b):
    return None if b[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" else "not an .xls (OLE2) file"


def is_xlsx(b):
    return None if b[:2] == b"PK" else "not an .xlsx (zip) file"


def is_pdf(b):
    return None if b[:5] == b"%PDF-" else "not a PDF"


def has(*needles, enc="utf-8"):
    def check(b):
        t = b.decode(enc, "replace")
        missing = [n for n in needles if n not in t]
        return f"missing {missing}" if missing else None
    return check


def text_of(b: bytes) -> str:
    for enc in ("utf-8", "cp932", "euc-jp"):
        try:
            return b.decode(enc)
        except UnicodeDecodeError:
            continue
    return b.decode("utf-8", "replace")


def norm(s) -> str:
    """NFKC (full-width digits to ASCII), no spaces of either width."""
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(s or "")))


def links(html: str):
    for m in re.finditer(r'href="([^"]+)"[^>]*>(.*?)</a>', html, re.S | re.I):
        yield m.group(1), norm(re.sub(r"<[^>]+>", "", m.group(2)))


def num(v):
    """A table cell as a float, or None for the dashes MOF uses for nil."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("", "-", "－", "ー", "―", "…", "△"):
        return None
    s = s.replace("△", "-").replace("▲", "-")
    return float(s)


def write(path: Path, comment: str, header: list[str], rows: list[list]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        for line in comment.strip().splitlines():
            fh.write(f"# {line.strip()}\n")
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(rows)
    print(f"  {path.relative_to(ROOT)}  {len(rows)} rows")


def expect(label, got, want, tol):
    if got in ("", None) or abs(float(got) - want) > tol:
        raise RuntimeError(f"spot check {label}: got {got!r}, expected {want} ± {tol}")
    print(f"    ok  {label} = {got}")


def r1(v, scale):
    return "" if v is None else round(v * scale, 1)


# --- tax revenue ----------------------------------------------------------

def latest_settlement_tax() -> tuple[int, float]:
    """(fiscal year, settlement tax revenue in 億円) from the newest 決算概要."""
    idx = text_of(cached("account_index.html", URL["account_index"], has("決算")))
    years = []
    for href, label in links(idx):
        m = re.fullmatch(r"令和(\d+|元)年度決算", label)
        if m and "/account/fy" in href:
            years.append((era_year("R", m.group(1)), urllib.parse.urljoin(URL["account_index"], href)))
    if not years:
        raise RuntimeError("mof: no 令和N年度決算 links on the settlement index")
    fy, page_url = max(years)
    page = text_of(cached(f"account_fy{fy}.html", page_url, has("決算")))
    pdf_url = next((urllib.parse.urljoin(page_url, h) for h, t in links(page)
                    if "予算決算比較表" in t and h.lower().endswith(".pdf")), None)
    if not pdf_url:
        raise RuntimeError(f"mof: no 予算決算比較表 PDF on {page_url}")
    blob = cached(f"account_fy{fy}_hikaku.pdf", pdf_url, is_pdf)
    if not shutil.which("pdftotext"):
        raise RuntimeError("mof: pdftotext (poppler) is needed to read the settlement PDF")
    pdf_path = CACHE / f"mof_account_fy{fy}_hikaku.pdf"
    out = subprocess.run(["pdftotext", "-layout", str(pdf_path), "-"],
                         capture_output=True, check=True).stdout.decode("utf-8")
    flat = norm(out)
    if f"令和{fy - 2018}年度一般会計決算概要" not in flat and \
            f"令和{'元' if fy == 2019 else fy - 2018}年度一般会計決算概要" not in flat:
        raise RuntimeError(f"mof: {pdf_url} is not the FY{fy} settlement summary")
    # Two 税収 figures on the page: budget (補正後) first, settlement second.
    hits = re.findall(r"税\s*収\s+([\d,]+)", out)
    if len(hits) < 2:
        raise RuntimeError(f"mof: could not find budget and settlement 税収 in {pdf_url}")
    return fy, float(hits[1].replace(",", ""))


def tax_revenue() -> dict[int, dict]:
    blob = cached("zeisyu.xls", URL["zeisyu"], is_xls)
    sh = XlsSheet(blob)
    cells = [[norm(c) for c in sh.row_values(i)] for i in range(sh.nrows)]
    if not any("(億円)" in c for row in cells[:5] for c in row):
        raise RuntimeError("zeisyu: unit cell (億円) not found")
    head = next(i for i, row in enumerate(cells) if "一般会計税収" in row)
    col = {name: cells[head].index(name) for name in
           ("一般会計税収", "所得税収", "法人税収", "消費税収")}
    out, era = {}, None
    for i in range(head + 1, sh.nrows):
        row = sh.row_values(i)
        era = ERA_NAME.get(norm(row[1]), era)
        if era is None or norm(row[2]) == "":
            continue
        fy = era_year(era, norm(row[2]))
        if int(float(row[3])) != fy:      # the sheet carries both: they must agree
            raise RuntimeError(f"zeisyu: row {i} wareki {era}{row[2]} != {row[3]}")
        out[fy] = {k: num(row[c]) for k, c in col.items()}
    return out


# --- general account ------------------------------------------------------

def stats_table(number: int) -> bytes:
    idx = text_of(cached("stats_index.html", URL["stats_index"], has("統計表一覧")))
    hits = [(h, t) for h, t in links(idx) if t.startswith(f"第{number}表")]
    if len(hits) != 1:
        raise RuntimeError(f"mof stats: expected one link for 第{number}表, got {hits}")
    href, title = hits[0]
    print(f"    第{number}表 -> {href}  {title[:40]}")
    return cached(f"stats_{number:02d}.xlsx", urllib.parse.urljoin(URL["stats_index"], href), is_xlsx)


def table1() -> dict[int, dict]:
    """Budget and settlement totals, 千円, from the 平成 and 令和 sheets."""
    wb = openpyxl.load_workbook(io.BytesIO(stats_table(1)), data_only=True)
    out = {}
    for ws in wb.worksheets:
        if not ("平成" in ws.title or "令和" in ws.title):
            continue
        rows = list(ws.iter_rows(values_only=True))
        if not any("(単位:千円)" in norm(c) for c in rows[0] if c):
            raise RuntimeError(f"table 1 {ws.title}: unit is not 千円")
        if [norm(c) for c in rows[2][2:6]] != ["予算額", "決算額", "予算額", "決算額"]:
            raise RuntimeError(f"table 1 {ws.title}: unexpected header {rows[2][:6]}")
        era, fy = None, None
        for r in rows[3:]:
            c0 = norm(r[0])
            if c0 in ERA_NAME:
                era = ERA_NAME[c0]
                continue
            m = re.fullmatch(r"(元|\d+)(年度)?", c0)
            if m and era:
                fy = era_year(era, m.group(1))
            if fy and norm(r[1]).startswith("計"):
                out[fy] = {"rev_budget": num(r[2]), "rev_settle": num(r[3]),
                           "exp_budget": num(r[4]), "exp_settle": num(r[5])}
    return out


def table4() -> dict[int, dict]:
    """Settlement revenue by main item, 百万円."""
    wb = openpyxl.load_workbook(io.BytesIO(stats_table(4)), data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if "一般会計歳入主要科目別決算" not in norm(rows[1][0]):
        raise RuntimeError("table 4: title is not 一般会計歳入主要科目別決算")
    if not any("(単位:百万円)" in norm(c) for c in rows[2] if c):
        raise RuntimeError("table 4: unit is not 百万円")
    h3, h4 = [norm(c) for c in rows[3]], [norm(c) for c in rows[4]]
    c_tax0 = h3.index("租税及印紙収入")
    c_tax = h4.index("計", c_tax0)
    c_bond, c_total = h3.index("公債金"), h3.index("合計")
    out, era = {}, None
    for r in rows[5:]:
        c0 = norm(r[0])
        if c0 in ERA_NAME:
            era = ERA_NAME[c0]
        y = norm(r[1])
        if era and re.fullmatch(r"元|\d+", y):
            out[era_year(era, y)] = {"tax": num(r[c_tax]), "bonds": num(r[c_bond]),
                                     "total": num(r[c_total])}
    return out


def table20() -> dict[int, dict]:
    """Settlement expenditure by main item, 千円, one sheet per fiscal year."""
    wb = openpyxl.load_workbook(io.BytesIO(stats_table(20)), data_only=True)
    out = {}
    for ws in wb.worksheets:
        m = re.fullmatch(r"(平成|令和)(元|\d+)", norm(ws.title))
        if not m:
            continue
        fy = era_year(m.group(1), m.group(2))
        rows = list(ws.iter_rows(values_only=True))
        if not any("(単位:千円)" in norm(c) for r in rows[:3] for c in r if c):
            raise RuntimeError(f"table 20 {ws.title}: unit is not 千円")
        if norm(rows[2][0]) != f"{m.group(1)}{m.group(2)}年度":
            raise RuntimeError(f"table 20 {ws.title}: title row says {rows[2][0]!r}")
        head = [norm(c) for c in rows[3]]
        c_settle = head.index("決算額")
        labels = [norm(r[1]) for r in rows]
        debt = labels.index("国債費")
        ss = labels.index("(社会保障関係費)")
        ss_total = labels.index("計", ss)
        out[fy] = {"debt_service": num(rows[debt][c_settle]),
                   "social_security": num(rows[ss_total][c_settle])}
    return out


# --- central government debt ----------------------------------------------

DEBT_ROWS = {  # normalised label -> column
    "内国債": "jgb", "普通国債": "general_bonds",
    "財政投融資特別会計国債": "filp_bonds", "財政融資資金特別会計国債": "filp_bonds",
    "借入金": "borrowings", "政府短期証券": "financing_bills",
    "合計": "total", "政府保証債務": "guaranteed",
}
DEBT_COLS = ["jgb", "general_bonds", "filp_bonds", "borrowings",
             "financing_bills", "total", "guaranteed"]


def parse_suii(blob: bytes, edition: str) -> dict[tuple[int, int], dict]:
    """{(year, month): {column: 億円}} from one edition of suii.xls."""
    sh = XlsSheet(blob)
    if not any("億円" in str(c) for i in range(min(3, sh.nrows)) for c in sh.row_values(i)):
        raise RuntimeError(f"suii {edition}: unit 億円 not found")
    out: dict[tuple[int, int], dict] = defaultdict(dict)
    quarters = None
    for i in range(sh.nrows):
        row = sh.row_values(i)
        label = norm(str(row[0]).split("\n")[0]) or norm(str(row[1]).split("\n")[0])
        if label.startswith("区分"):
            quarters = []
            for cell in row[3:]:
                if not str(cell).strip():
                    quarters.append(None)
                    continue
                first, *rest = str(cell).split("\n")
                m = re.fullmatch(r"([HR])(\d+)\.(\d+)末", norm(first))
                if not m:
                    raise RuntimeError(f"suii {edition}: column head {cell!r}")
                y, mo = era_year(m.group(1), m.group(2)), int(m.group(3))
                if rest and not rest[0].strip().startswith(str(y)):
                    raise RuntimeError(f"suii {edition}: {cell!r} wareki/western disagree")
                quarters.append((y, mo))
            continue
        col = DEBT_ROWS.get(label)
        if col and quarters:
            for q, v in zip(quarters, row[3:]):
                if q is not None and num(v) is not None:
                    out[q][col] = num(v)
    if not out:
        raise RuntimeError(f"suii {edition}: no quarters read")
    return dict(out)


def parse_quarter_page(html: str, y: int, mo: int, edition: str) -> dict:
    """One archived per-quarter page (e.g. 1812.htm = end-December 2006)."""
    flat = norm(htmllib.unescape(re.sub(r"<[^>]+>", "", html)))
    if f"(平成{y - 1988}年{mo}月末現在)" not in flat:
        raise RuntimeError(f"debt page {edition}: is not headed 平成{y - 1988}年{mo}月末現在 "
                           "— wrong quarter?")
    out = {}
    for tr in re.findall(r"<tr.*?</tr>", html, re.S | re.I):
        cells = [norm(htmllib.unescape(re.sub(r"<[^>]+>", "", c)))
                 for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", tr, re.S | re.I)]
        cells = [c for c in cells if c]
        if len(cells) >= 2 and cells[0] in DEBT_ROWS:
            out[DEBT_ROWS[cells[0]]] = num(cells[1])
    if "total" not in out or "guaranteed" not in out:
        raise RuntimeError(f"debt page {edition}: total/guaranteed not found")
    return out


def warp_frame(wrapper_url: str, name: str) -> str:
    """WARP serves a wrapper page; the archived page is its pywb iframe."""
    html = text_of(cached(f"warp_{name}.html", wrapper_url, has("pywb-frame"), pause=2))
    m = re.search(r'id="pywb-frame"[^>]*\ssrc="([^"]+)"', html)
    if not m:
        raise RuntimeError(f"warp: no replay frame in {wrapper_url}")
    return m.group(1)                      # /<collection>/<timestamp>/www.mof.go.jp/...


def central_debt():
    editions = []                          # (capture timestamp, label, {quarter: row})
    live_idx = text_of(cached("gbb_index.html", URL["gbb_index"], has("suii.xls")))
    live = cached("gbb_suii.xls", urllib.parse.urljoin(URL["gbb_index"], "suii.xls"), is_xls)
    editions.append(("99999999999999", "mof.go.jp suii.xls (live)", parse_suii(live, "live")))

    # Follow the archive links: live page -> WARP captures -> their links.
    queue = [h for h, _ in links(live_idx) if "warp" in h and "gbb/data.htm" in h]
    seen_frames, pending = set(), []
    while queue:
        wrapper = queue.pop(0)
        inner = re.sub(r"^https?://warp\.ndl\.go\.jp/\d+/\d+/", "", wrapper)
        if inner.startswith("http://warp") or inner.startswith("warp"):
            wrapper = "https://" + inner.split("://", 1)[-1]
        try:
            frame = warp_frame(wrapper, re.sub(r"\W+", "_", wrapper.split("ndl.go.jp/")[-1])[:80])
            if frame in seen_frames:
                continue
            seen_frames.add(frame)
            prefix = frame.split("/www.mof.go.jp")[0]
            ts = prefix.rsplit("/", 1)[-1]
            page = text_of(cached(f"warp_{ts}_data.htm", WARP + frame, has("gbb"),
                                  pause=WARP_PAUSE))
            if any("suii.xls" in h for h, _ in links(page)):
                blob = cached(f"warp_{ts}_suii.xls",
                              f"{WARP}{prefix}/www.mof.go.jp/jgbs/reference/gbb/suii.xls",
                              is_xls, pause=WARP_PAUSE)
                editions.append((ts, f"NDL WARP {ts} suii.xls", parse_suii(blob, ts)))
        except RuntimeError as e:          # the archive is down or throttling: say so
            pending.append(f"{wrapper}: {e}")
            continue
        for h, _ in links(page):
            if "warp.ndl.go.jp" in h and "gbb/data.htm" in h and h not in queue:
                queue.append(h)

    have = set().union(*(e[2] for e in editions))
    first = min(have)
    want = [(y, m) for y in range(Y0, Y1 + 1) for m in (3, 6, 9, 12)
            if (y, m) < first]
    pages = {}
    try:
        if want:
            frame = warp_frame(WARP_QUARTERLY_INDEX, "quarterly_index")
            prefix = frame.split("/www.mof.go.jp")[0]
            ts = prefix.rsplit("/", 1)[-1]
            idx = text_of(cached(f"warp_{ts}_data.htm", WARP + frame, has("gbb"),
                                 pause=WARP_PAUSE))
            listed = {h for h, _ in links(idx)}
        for y, m in want:
            name = f"{y - 1988:02d}{m:02d}.htm"           # 平成 YYMM
            if name not in listed:
                raise RuntimeError(f"warp {ts}: index does not list {name}")
            html = text_of(cached(f"warp_{ts}_{name}",
                                  f"{WARP}{prefix}/www.mof.go.jp/jgbs/reference/gbb/{name}",
                                  has("政府保証債務", enc="cp932"), pause=WARP_PAUSE))
            # A page that is there but says the wrong quarter is a real failure.
            pages[(y, m)] = parse_quarter_page(html, y, m, f"{ts} {name}")
    except RuntimeError as e:
        if "wrong quarter" in str(e):
            raise
        pending.append(f"quarterly pages: {e}")
    if pages:
        editions.append((ts, f"NDL WARP {ts} quarterly pages", pages))

    editions.sort(reverse=True)            # newest first
    rows, worst = {}, 0.0
    for ts, label, data in editions:
        for q, vals in data.items():
            if q in rows:
                newer = rows[q][0]
                for c in ("total", "guaranteed"):
                    a, b = newer.get(c), vals.get(c)
                    if a is not None and b is not None:
                        worst = max(worst, abs(a - b))
                        if abs(a - b) > 1:                 # one 億円 of rounding
                            raise RuntimeError(f"debt {q} {c}: {label} says {b}, newer "
                                               f"edition {rows[q][1]} says {a}")
            else:
                rows[q] = (vals, label)
    print(f"    {len(editions)} editions; overlapping quarters agree to within {worst:.0f} 億円")
    return rows, pending


# --- JGB yields -----------------------------------------------------------

def jgb():
    def parse(blob, label):
        t = blob.decode("cp932")
        rows = list(csv.reader(io.StringIO(t)))
        head = next(i for i, r in enumerate(rows) if r and r[0] == "基準日")
        tenors = rows[head][1:]
        out = {}
        for r in rows[head + 1:]:
            if not r or not r[0].strip() or r[0].startswith("※"):
                continue
            d = wareki_date(r[0])
            out[d] = {t: num(v) for t, v in zip(tenors, r[1:])}
        if not out:
            raise RuntimeError(f"jgb {label}: no rows")
        return tenors, out
    tenors, daily = parse(cached("jgbcm_all.csv", URL["jgb_all"], has("基準日", "10年", enc="cp932")), "all")
    t2, now = parse(cached("jgbcm.csv", URL["jgb_now"], has("基準日", enc="cp932")), "current")
    if t2 != tenors:
        raise RuntimeError("jgb: the running-month file has different tenors")
    for d, v in now.items():
        if d in daily and daily[d] != v:
            raise RuntimeError(f"jgb: {d} differs between the archive and the running month")
        daily[d] = v
    dates = sorted(daily)
    if dates != sorted(set(dates)):
        raise RuntimeError("jgb: duplicate dates")
    return tenors, daily


# --- reserves -------------------------------------------------------------

def reserves():
    t = cached("reserves.csv", URL["reserves"],
               has("外貨準備", "Official reserve assets", enc="cp932")).decode("cp932")
    rows = list(csv.reader(io.StringIO(t)))
    want = {4: "A. Official reserve assets", 5: "(1) Foreign currency reserves",
            16: "(4) Gold"}
    for col, text in want.items():
        if not any(len(r) > col and r[col].strip().startswith(text) for r in rows[:20]):
            raise RuntimeError(f"reserves: column {col} is no longer {text!r}")
    if not any("百万ドル" in r[0] for r in rows[:8] if r):
        raise RuntimeError("reserves: unit 百万ドル not found")
    out, year = {}, None
    for r in rows:
        if len(r) < 17:
            continue
        m = re.fullmatch(r"(平成|令和)(元|\d+)年", norm(r[0]))
        if m:
            year = era_year(m.group(1), m.group(2))
            if r[2].strip() and int(r[2]) != year:
                raise RuntimeError(f"reserves: {r[0]} next to {r[2]}")
        mm = re.fullmatch(r"(\d{1,2})月", norm(r[1]))
        if not (mm and year):
            continue
        out[(year, int(mm.group(1)))] = {"total": num(r[4]), "fx": num(r[5]), "gold": num(r[16])}
    return out


# --- main -----------------------------------------------------------------

def main() -> None:
    print("Ministry of Finance")
    _wareki_self_test()

    # Tax revenue -------------------------------------------------------
    tax = tax_revenue()
    fy_settled, settled = latest_settlement_tax()
    last = max(tax)
    if last > fy_settled:
        print(f"    zeisyu runs to FY{last}, beyond the latest settlement FY{fy_settled}: "
              "the later rows are not outturns and are dropped")
        tax = {k: v for k, v in tax.items() if k <= fy_settled}
    if abs(tax[fy_settled]["一般会計税収"] - settled) > 0.5:
        raise RuntimeError(f"zeisyu FY{fy_settled} {tax[fy_settled]['一般会計税収']} is not "
                           f"the settlement figure {settled} in the 決算概要")
    print(f"    zeisyu FY{fy_settled} = 決算概要 settlement {settled:.0f} 億円")
    rows = [[fy, r1(v["一般会計税収"], 0.1), r1(v["所得税収"], 0.1),
             r1(v["法人税収"], 0.1), r1(v["消費税収"], 0.1)] for fy, v in sorted(tax.items())]
    write(DATA / "fiscal" / "tax_revenue_from_mof.csv", f"""
        Ministry of Finance, 昭和54年度以降の税収の推移 (zeisyu.xls): NATIONAL general-account tax revenue, settlement basis.
        fiscal_year is the STARTING year of the April-March fiscal year (2024 = April 2024 - March 2025).
        Published in 100 million yen (億円); converted here to YEN BILLION (divided by 10).
        FY{fy_settled} was checked against the 令和{fy_settled - 2018}年度決算概要 budget/settlement table (as of the closing of the books).
        Not local taxes, not social contributions, not general government. Consumption tax starts FY1989 (blank before).
    """, ["fiscal_year", "tax_revenue_bn", "income_tax_bn", "corporation_tax_bn",
          "consumption_tax_bn"], rows)
    print("    spot checks:")
    t = {r[0]: r for r in rows}
    expect("tax FY2006 (¥bn)", t[2006][1], 49069.1, 0.05)
    expect("tax FY2024 (¥bn)", t[2024][1], 75232.1, 0.05)
    expect("tax FY2025 (¥bn)", t[2025][1], 84222.6, 0.05)

    # General account -----------------------------------------------------
    t1, t4, t20 = table1(), table4(), table20()
    for fy in sorted(set(t1) & set(t4)):
        a, b = t1[fy]["rev_settle"], t4[fy]["total"]
        if a is not None and b is not None and abs(a / 1000 - b) > 1.0:
            raise RuntimeError(f"general account FY{fy}: table 1 revenue {a} 千円 != table 4 {b} 百万円")
    for fy in sorted(set(t4) & set(tax)):
        a, b = t4[fy]["tax"], tax[fy]["一般会計税収"]
        if a is not None and abs(a / 100 - b) > 1.0:
            raise RuntimeError(f"FY{fy}: table 4 tax {a} 百万円 vs zeisyu {b} 億円")
    rows = []
    for fy in range(Y0, Y1 + 1):
        a, b, c = t1.get(fy, {}), t4.get(fy, {}), t20.get(fy, {})
        if not (a or b or c):
            continue
        exp = a.get("exp_settle")
        bonds = b.get("bonds")
        dep = "" if exp is None or bonds is None else round(100 * (bonds * 1000) / exp, 1)
        rows.append([fy, r1(a.get("rev_budget"), 1e-6), r1(a.get("rev_settle"), 1e-6),
                     r1(a.get("exp_budget"), 1e-6), r1(exp, 1e-6),
                     r1(b.get("tax"), 1e-3), r1(bonds, 1e-3),
                     r1(c.get("debt_service"), 1e-6), r1(c.get("social_security"), 1e-6), dep])
    write(DATA / "fiscal" / "general_account_from_mof.csv", """
        Ministry of Finance, 財政統計: national GENERAL ACCOUNT only (not special accounts, not local government, not general government).
        fiscal_year is the STARTING year of the April-March fiscal year. All amounts converted to YEN BILLION.
        *_budget = 予算額 after supplementary budgets, *_settlement = 決算額 (table 1, published in 千円 = ¥ thousand).
        tax_settlement (租税及印紙収入, taxes plus stamp revenue) and bond_revenue_settlement (公債金, new JGB issuance
        booked as revenue) are table 4 (published in 百万円 = ¥ million). debt_service_settlement (国債費) and
        social_security_settlement (社会保障関係費) are table 20 (千円). bond_dependency_pct = bond revenue / settled expenditure.
        Settlement columns are blank for years whose accounts are not yet in the statistics tables; budget columns are
        budgets, never outturns, and the plates draw settlement. FY2005's settlement is missing from table 1 itself.
    """, ["fiscal_year", "revenue_budget_bn", "revenue_settlement_bn",
          "expenditure_budget_bn", "expenditure_settlement_bn", "tax_settlement_bn",
          "bond_revenue_settlement_bn", "debt_service_settlement_bn",
          "social_security_settlement_bn", "bond_dependency_pct"], rows)
    print("    spot checks:")
    g = {r[0]: r for r in rows}
    expect("revenue settlement FY2024 (¥bn)", g[2024][2], 135980.9, 0.05)
    expect("expenditure settlement FY2024 (¥bn)", g[2024][4], 123024.0, 0.05)
    expect("debt service settlement FY2024 (¥bn)", g[2024][7], 25689.4, 0.05)
    expect("expenditure settlement FY2006 = table 20 total (¥bn)", g[2006][4], 81445.5, 0.05)

    # Central-government debt -----------------------------------------------
    debt, pending = central_debt()
    rows = []
    for y in range(Y0, Y1 + 1):
        for m in (3, 6, 9, 12):
            if (y, m) > max(debt):
                break
            vals, label = debt.get((y, m), ({}, ""))     # a hole stays a hole
            rows.append([y, f"{y}-{m:02d}", *[r1(vals.get(c), 0.1) for c in DEBT_COLS], label])
    missing = [r[1] for r in rows if not r[-1]]
    if missing:
        print(f"    PENDING: central debt quarters not staged (left blank): {missing}")
        for p in pending:
            print(f"      {p}")
    write(DATA / "fiscal" / "central_debt_from_mof.csv", """
        Ministry of Finance, 国債及び借入金並びに政府保証債務現在高: CENTRAL GOVERNMENT debt at FACE value, quarter-ends
        (quarter_end = end of that month; year = its calendar year). Published in 100 million yen; converted to YEN BILLION.
        total = jgb (内国債: general bonds incl. reconstruction and GX bonds, FILP bonds, subsidy and other bonds)
        + borrowings (借入金) + financing_bills (政府短期証券). general_bonds and filp_bonds are PARTS of jgb.
        guaranteed (政府保証債務, government-guaranteed debt) is NOT in total: a contingent liability, a separate column.
        The live file holds five years; earlier quarters are MOF's own earlier editions, read from the National Diet
        Library web archive (WARP): edition names the file each row came from, the newest edition wins, and every
        overlap between editions was checked to agree. Not general government (IMF) and not market value (BoJ flow of funds).
    """, ["year", "quarter_end", *[f"{c}_bn" for c in DEBT_COLS], "edition"], rows)
    print("    spot checks:")
    d = {r[1]: r for r in rows}
    expect("central debt Jun-2026 (¥bn)", d["2026-06"][7], 1346683.3, 0.05)
    if d["2006-12"][-1]:
        expect("central debt Dec-2006 (¥bn)", d["2006-12"][7], 832263.1, 0.05)

    # The magnitude, against the IMF: central face-value debt over ESRI
    # nominal GDP must land near the IMF's general-government ratio.
    imf = json.loads(cached("imf_ggxwdg.json", URL["imf_debt"], has("GGXWDG_NGDP")))
    imf_jpn = imf["values"]["GGXWDG_NGDP"]["JPN"]
    expect("IMF gross GG debt 2024 (% GDP)", imf_jpn["2024"], 213.6, 0.05)
    ratio = 100 * d["2024-12"][7] / ESRI_NGDP_2024_BN
    print(f"    MOF central debt Dec-2024 / ESRI CY2024 GDP = {ratio:.1f}% vs IMF GG {imf_jpn['2024']}%")
    if not 0.85 < ratio / imf_jpn["2024"] < 1.05:
        raise RuntimeError("central debt: magnitude disagrees with the IMF ratio")

    # JGB yields -----------------------------------------------------------
    tenors, daily = jgb()
    for day, want in ((date(2006, 1, 4), 1.446), (date(2016, 7, 1), -0.262),
                      (date(2026, 9, 30), 3.057)):
        expect(f"JGB 10y {day}", daily[day]["10年"], want, 0.0005)
    run = max(daily).year
    cols = ["2年", "5年", "10年", "20年", "30年", "40年"]
    by_year: dict[int, dict[str, list]] = defaultdict(lambda: defaultdict(list))
    months: dict[int, dict[str, set]] = defaultdict(lambda: defaultdict(set))
    for day in sorted(daily):
        for c in cols:
            v = daily[day].get(c)
            if v is not None:
                by_year[day.year][c].append(v)
                months[day.year][c].add(day.month)
    rows = []
    for y in range(Y0, run + 1):
        out = [y]
        for c in cols:
            vals = by_year[y][c]
            full = len(months[y][c]) == 12 or y == run
            out.append(round(statistics.fmean(vals), 3) if vals and full else "")
        ten = [daily[dd]["10年"] for dd in sorted(daily) if dd.year == y and daily[dd]["10年"] is not None]
        out.insert(4, ten[-1] if ten else "")
        out += [len(months[y]["10年"]), max(dd for dd in daily if dd.year == y).isoformat()]
        rows.append(out)
    write(DATA / "markets" / "jgb_from_mof.csv", """
        Ministry of Finance, 国債金利情報: JGB constant-maturity yields (% p.a.), daily, from jgbcm_all.csv plus the running
        month's jgbcm.csv. Dates in the source are Japanese-era (wareki) strings, converted with an explicit era table.
        *_avg = mean of the year's daily readings; jgb_10y_eoy = last reading of the year. Calendar years.
        A tenor is blank in a past year it did not cover in full (40y starts Nov-2007; MOF writes "-" before issuance).
        The running year is partial: months and last_date say how far it runs. YCC band: Sep-2016 to Mar-2024.
    """, ["year", "jgb_2y_avg", "jgb_5y_avg", "jgb_10y_avg", "jgb_10y_eoy", "jgb_20y_avg",
          "jgb_30y_avg", "jgb_40y_avg", "months", "last_date"], rows)

    # The 10-year against a second source: OECD MEI long-term rate (via FRED).
    oecd = defaultdict(list)
    for line in cached("oecd_irltlt01jpm.csv", URL["oecd_10y"], has("IRLTLT01JPM156N")).decode().splitlines()[1:]:
        day, v = line.split(",")
        if v.strip() and v != ".":
            oecd[int(day[:4])].append(float(v))
    j = {r[0]: r[3] for r in rows}
    gaps = {y: round(j[y] - statistics.fmean(oecd[y]), 3) for y in range(Y0, run)
            if len(oecd.get(y, [])) == 12}
    worst = max(gaps.items(), key=lambda kv: abs(kv[1]))
    print(f"    JGB 10y annual mean vs OECD MEI IRLTLT01JPM156N, {min(gaps)}-{max(gaps)}: "
          f"worst gap {worst[1]:+.3f}pp in {worst[0]}")
    if abs(worst[1]) > 0.06:
        raise RuntimeError(f"JGB 10y disagrees with OECD by {worst[1]}pp in {worst[0]}")

    # Reserves -------------------------------------------------------------
    res = reserves()
    run = max(res)[0]
    rows = []
    for y in range(Y0, Y1 + 1):
        ms = [m for (yy, m) in res if yy == y]
        if not ms:
            continue
        m = 12 if y < run else max(ms)
        v = res.get((y, m))
        if v is None:
            rows.append([y, "", "", "", ""])
            continue
        rows.append([y, round(v["total"]), round(v["fx"]), round(v["gold"]), f"{y}-{m:02d}"])
    write(DATA / "international" / "reserves_from_mof.csv", """
        Ministry of Finance, 外貨準備等の状況 (International Reserves/Foreign Currency Liquidity, IMF template), historical.csv.
        US$ MILLION, end of month. Each year is its December stock; the running year is its latest month (asof_month).
        reserves_total = I.A official reserve assets; fx_reserves (I.A.1 foreign currency reserves) and gold (I.A.4, at market
        value) are PARTS of it. Stocks, never summed. The year is written in wareki on each January row and carried down.
    """, ["year", "reserves_total_musd", "fx_reserves_musd", "gold_musd", "asof_month"], rows)
    print("    spot checks:")
    expect("reserves Dec-2006", res[(2006, 12)]["total"], 895320, 0.5)
    expect("reserves Dec-2024", res[(2024, 12)]["total"], 1230715, 0.5)
    expect("reserves Dec-2025", res[(2025, 12)]["total"], 1369775, 0.5)
    expect("reserves Sep-2026", res[(2026, 9)]["total"], 1178396, 0.5)
    wb = json.loads(cached("wb_reserves_2024.json", URL["wb_reserves"], has("FI.RES.TOTL.CD")))
    wb_v = wb[1][0]["value"] / 1e6
    print(f"    World Bank FI.RES.TOTL.CD 2024 = {wb_v:,.0f} US$mn vs MOF Dec-2024 {res[(2024, 12)]['total']:,.0f}")
    if abs(wb_v / res[(2024, 12)]["total"] - 1) > 0.01:
        raise RuntimeError("reserves: MOF and the World Bank disagree by more than 1%")


if __name__ == "__main__":
    main()
