#!/usr/bin/env python3
"""Refresh the Cabinet Office (ESRI) national accounts for Japan.

Usage:  python3 -I countries/japan/fetch/fetch_esri.py

Two products, both on esri.cao.go.jp, both Shift_JIS/xlsx file layers
with no API:

**The Quarterly Estimates (QE, 四半期別GDP速報).** One menu page per
release, `sna/data/data_list/sokuhou/files/<YYYY>/qe<YYQ>_<n>/
gdemenuja.html`, listing CSVs named `<kind>-<series><suffix>.csv`.
THE RELEASE DIRECTORY AND THE SUFFIX MOVE WITH EVERY QE (qe262_2 / 2622
is the 2026 Apr–Jun second preliminary of 2026-09-08), so both are read
off `sna/sokuhou/sokuhou_top.html` on every run and never written here.
Taken from each release:

  gaku-mcy  nominal, calendar year        gaku-jcy  real, calendar year
  gaku-mfy  nominal, fiscal year          gaku-jfy  real, fiscal year
  gaku-mk   nominal, SA, quarterly        gaku-jk   real, SA, quarterly
  ritu-mk / ritu-jk    published q/q % change of the SA series
  nritu-jk             published annualised q/q % change (real)

**The quarterly SA levels are at annual rates.** Every SA file ends
with "＊年率で表示している。" (shown at an annual rate): 2026 Q2 real GDP
is ¥598,950.5bn, a year's worth of output, not a quarter's. The annual
figures here therefore come from the calendar- and fiscal-year tables,
never from summing quarters (the sum is about four times GDP); the
fetcher asserts both that the mean of the four SA quarters sits on the
annual table and that the sum does not.

**Calendar and fiscal years are different files and different columns.**
The fiscal year runs April–March and is keyed on its starting year
(`2024/4-3.` = FY2024 = Apr 2024 – Mar 2025).

**The annual accounts (年次推計, 確報).** One edition per year, linked
from `sna/menu.html` as `kakuhou/files/<FY>/<FY>_kaku_top.html`; the
latest edition is FY2024 (published December 2025, the FY2025 edition
is due December 2026). Two tables are taken from its top page:

  (6-2) 一般政府の部門別勘定（GFS） — general-government accounts in
        the IMF GFSM 2014 layout, fiscal years, one block of five
        columns per year (central, local, social security funds,
        consolidation, general government).
  (7)   一般政府の機能別支出（COFOG） — expenditure by function, one sheet
        per fiscal year named in wareki (平成１７ … 令和６).

**The COFOG table has no total column.** It is a matrix of functions
by transaction (final consumption, subsidies, social benefits other
than in kind, other current transfers, GFCF, inventories, capital
transfers). Interest and land purchases are not in it. The per-function
totals written here are the SUM OF THOSE SEVEN PUBLISHED COLUMNS,
computed by this script and labelled as such; the fetcher asserts that
the ten functions add up to ESRI's own 合計 row in every column.

Spot values (DISCOVERY.md, release qe262_2): nominal CY2024 ¥636,893.3bn,
CY2025 ¥666,440.7bn, CY2006 ¥540,566.0bn; real SA 2026 Q2 ¥598,950.5bn.
They are asserted to the decimal while the release is qe262_2 and as
magnitudes after a newer QE revises them. Cross-check: nominal CY2024
over the BoJ's 2024 mean USD/JPY (151.50, FM08 FXERM07) must land within
1.5% of the World Bank's US$4.190tn.

Outputs are `_from_esri` sidecars under data/. Nothing here overwrites
a curated file.
"""
from __future__ import annotations

import csv
import io
import re
import sys
import time
import unicodedata
import urllib.error
import urllib.request
import warnings
import zipfile
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install -r requirements-fetch.txt")

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
CACHE = ROOT / "cache"

ESRI = "https://www.esri.cao.go.jp"
QE_TOP = f"{ESRI}/jp/sna/sokuhou/sokuhou_top.html"
SNA_MENU = f"{ESRI}/jp/sna/menu.html"
UA = {"User-Agent": "worldreports-fetch"}
GAP = 1.5                 # seconds between requests
LAST_FULL = 2025          # country.yaml last_full_year

# The release whose figures DISCOVERY.md recorded; exact spot values
# only hold for it.
VERIFIED_RELEASE = "qe262_2"

# Cross-check against a second source (both recorded in DISCOVERY.md).
WB_GDP_2024_USD_BN = 4190.0     # World Bank NY.GDP.MKTP.CD, JPN, 2024
BOJ_USDJPY_2024 = 151.50        # BoJ FM08 FXERM07, 2024 monthly mean


# ---------------------------------------------------------------- transport

_last = [0.0]


def _get(url: str) -> bytes:
    """One request, spaced, retried, backing off on 429/403."""
    last = None
    for attempt in range(4):
        wait = _last[0] + GAP - time.monotonic()
        if wait > 0:
            time.sleep(wait)
        try:
            req = urllib.request.Request(url, headers=UA)
            with urllib.request.urlopen(req, timeout=300) as r:
                blob = r.read()
            _last[0] = time.monotonic()
            return blob
        except urllib.error.HTTPError as e:
            _last[0] = time.monotonic()
            last = e
            if e.code == 404:
                break  # a wrong path does not heal on retry
            time.sleep(30 * (attempt + 1) if e.code in (429, 403) else 3)
        except OSError as e:
            _last[0] = time.monotonic()
            last = e
            time.sleep(3 * (attempt + 1))
    raise RuntimeError(f"{url}: {last}")


def fetch(url: str, name: str, check, fresh: bool = False) -> bytes:
    """Cached download; `check(blob)` raises before anything is cached.

    Index pages are fetched `fresh` on every run (that is how a new
    release is found); data files are cached under release-specific
    names, so a new release is a new file.
    """
    CACHE.mkdir(exist_ok=True)
    local = CACHE / name
    if local.exists() and not fresh:
        return local.read_bytes()
    print(f"  downloading {url}")
    blob = _get(url)
    check(blob)
    local.write_bytes(blob)
    return blob


def is_html(marker: bytes):
    def check(blob: bytes) -> None:
        if b"<html" not in blob[:2000].lower() or marker not in blob:
            raise RuntimeError(f"page lacks {marker!r} (moved or walled?)")
    return check


def is_xlsx(blob: bytes) -> None:
    if blob[:2] != b"PK":
        raise RuntimeError("payload is not an xlsx (error page?)")
    bad = zipfile.ZipFile(io.BytesIO(blob)).testzip()
    if bad:
        raise RuntimeError(f"xlsx member {bad} is corrupt")


def is_qe_csv(blob: bytes) -> None:
    if blob[:200].lower().lstrip().startswith(b"<"):
        raise RuntimeError("QE payload is HTML, not CSV (error page?)")
    text = blob.decode("cp932")  # raises on a non-Shift_JIS payload
    if "GDP(Expenditure Approach)" not in text and "国内総生産" not in text:
        raise RuntimeError("QE CSV has no GDP column")
    if len(re.findall(r"^\d{4}/", text, re.M)) < 20:
        raise RuntimeError("QE CSV has fewer than 20 dated rows")


# ---------------------------------------------------------------- QE

def qe_release() -> tuple[str, str, dict[str, str]]:
    """(release dir, file suffix, {table key: url}) for the current QE."""
    top = fetch(QE_TOP, "esri_sokuhou_top.html", is_html(b"gdemenuja"),
                fresh=True).decode("utf-8", "replace")
    menus = re.findall(r'href="(/jp/sna/data/data_list/sokuhou/files/'
                       r'(\d{4})/(qe\d{3}_\d)/gdemenuja\.html)"', top)
    if len(set(menus)) != 1:
        raise RuntimeError(f"sokuhou_top lists {len(set(menus))} QE menus; "
                           "expected exactly one current release")
    path, _, release = menus[0]
    menu_url = ESRI + path
    menu = fetch(menu_url, f"esri_{release}_menu.html",
                 is_html(b"gaku-jk"), fresh=True).decode("utf-8", "replace")
    links = dict()
    for name, kind, ser, suffix in re.findall(
            r'href="tables/(([a-z]+)-([a-z]+)(\d{4})\.csv)"', menu):
        links.setdefault(f"{kind}-{ser}", (suffix, name))
    suffixes = {s for s, _ in links.values()}
    if len(suffixes) != 1:
        raise RuntimeError(f"{release}: mixed file suffixes {suffixes}")
    suffix = suffixes.pop()
    base = menu_url.rsplit("/", 1)[0] + "/tables/"
    return release, suffix, {k: base + n for k, (_, n) in links.items()}


def num(v) -> float | None:
    """A published cell as a number: commas stripped; blanks, '***'
    (not computable) and '-' are no value."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).replace(",", "").strip()
    if re.fullmatch(r"-?\d+(\.\d+)?", t):
        return float(t)
    return None


def qe_table(release: str, links: dict, key: str) -> dict:
    """{'unit', 'cols': {english name: index}, 'A': {y: row}, 'Q':
    {(y, q): row}} for one QE CSV."""
    if key not in links:
        raise RuntimeError(f"{release}: {key} not on the menu page")
    url = links[key]
    blob = fetch(url, f"esri_{release}_{url.rsplit('/', 1)[-1]}", is_qe_csv)
    rows = list(csv.reader(io.StringIO(blob.decode("cp932"))))
    unit = " ".join(c for c in rows[1] if c.startswith("(")) or ""
    # Column names: the English header row holds most names, the row
    # below it the export/import split.
    hi = next(i for i, r in enumerate(rows[:10])
              if "GDP(Expenditure Approach)" in r)
    cols = {}
    for r in (rows[hi], rows[hi + 1]):
        for i, c in enumerate(r):
            c = c.strip()
            if c and c not in cols:
                cols[c] = i
    out = {"unit": unit, "cols": cols, "A": {}, "Q": {},
           "foot": [r[0] for r in rows if r and r[0].startswith("＊")]}
    year = None
    for r in rows[hi + 2:]:
        lab = r[0].strip() if r else ""
        if m := re.fullmatch(r"(\d{4})/(1-12|4-3)\.", lab):
            out["A"][int(m[1])] = r
        elif m := re.fullmatch(r"(?:(\d{4})/)?\s*(\d+)-\s*(\d+)\.", lab):
            if m[1]:
                year = int(m[1])
            q = {"1": 1, "4": 2, "7": 3, "10": 4}.get(m[2])
            if year is None or q is None:
                raise RuntimeError(f"{key}: unreadable quarter {lab!r}")
            if (q == 1) != bool(m[1]):
                # The year is printed on Q1 rows only; anything else
                # means the layout changed.
                raise RuntimeError(f"{key}: year label off Q1 at {lab!r}")
            out["Q"][(year, q)] = r
        elif lab and not lab.startswith("＊") and re.match(r"\d", lab):
            raise RuntimeError(f"{key}: unrecognised period {lab!r}")
    return out


def col(t: dict, name: str, key: str) -> int:
    if name not in t["cols"]:
        raise RuntimeError(f"{key}: no column {name!r} "
                           f"(have {sorted(t['cols'])[:12]}…)")
    return t["cols"][name]


def series(t: dict, name: str, key: str, freq: str) -> dict:
    i = col(t, name, key)
    return {k: num(r[i]) if i < len(r) else None
            for k, r in t[freq].items()}


# ---------------------------------------------------------------- writers

def write(rel: str, comment: str, header: list[str], rows: list[list]) -> None:
    path = DATA / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        for line in comment.strip().splitlines():
            fh.write(f"# {line}".rstrip() + "\n")
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(rows)
    span = f"{rows[0][0]}–{rows[-1][0]}" if rows else "EMPTY"
    print(f"  data/{rel}  {len(rows)} rows ({span})")


def fmt(v, nd: int = 1):
    return "" if v is None else round(v, nd)


def near(a, b, tol, what):
    if a is None or b is None or abs(a - b) > tol:
        raise RuntimeError(f"{what}: {a} vs {b} (tolerance {tol})")


# ---------------------------------------------------------------- GDP

COMPONENTS = [
    ("PrivateConsumption", "private_consumption"),
    ("PrivateResidentialInvestment", "private_residential_investment"),
    ("Private Non-Resi.Investment", "private_nonresidential_investment"),
    ("GovernmentConsumption", "government_consumption"),
    ("PublicInvestment", "public_investment"),
    ("Exports", "exports"),
    ("Imports", "imports"),
    ("GNI", "gni"),
]
GDP = "GDP(Expenditure Approach)"


def gdp() -> dict:
    release, suffix, links = qe_release()
    print(f"  QE release {release} (suffix {suffix})")
    t = {k: qe_table(release, links, k) for k in (
        "gaku-mcy", "gaku-mfy", "gaku-jcy", "gaku-jfy", "gaku-mk",
        "gaku-jk", "ritu-mk", "ritu-jk", "nritu-jk")}

    # Units and the chain base, read off the files themselves.
    for k in ("gaku-mcy", "gaku-mfy", "gaku-mk"):
        if "Billion Yen" not in t[k]["unit"]:
            raise RuntimeError(f"{k}: unit is {t[k]['unit']!r}, not ¥bn")
    chain = re.search(r"Chained \((\d{4})\) Yen", t["gaku-jk"]["unit"])
    if not chain:
        raise RuntimeError(f"gaku-jk unit {t['gaku-jk']['unit']!r}")
    for k in ("gaku-jcy", "gaku-jfy"):
        if f"Chained ({chain[1]}) Yen" not in t[k]["unit"]:
            raise RuntimeError(f"{k}: chain base differs from gaku-jk")
    for k in ("gaku-mk", "gaku-jk"):
        if not any("年率" in f for f in t[k]["foot"]):
            raise RuntimeError(f"{k}: the annual-rate footnote is gone; "
                               "check whether levels are still annualised")
    for k in ("ritu-mk", "ritu-jk", "nritu-jk"):
        if "%" not in t[k]["unit"]:
            raise RuntimeError(f"{k}: unit {t[k]['unit']!r} is not %")

    ncy = series(t["gaku-mcy"], GDP, "gaku-mcy", "A")
    rcy = series(t["gaku-jcy"], GDP, "gaku-jcy", "A")
    nfy = series(t["gaku-mfy"], GDP, "gaku-mfy", "A")
    rfy = series(t["gaku-jfy"], GDP, "gaku-jfy", "A")
    nq = series(t["gaku-mk"], GDP, "gaku-mk", "Q")
    rq = series(t["gaku-jk"], GDP, "gaku-jk", "Q")
    nqq = series(t["ritu-mk"], GDP, "ritu-mk", "Q")
    rqq = series(t["ritu-jk"], GDP, "ritu-jk", "Q")
    rqa = series(t["nritu-jk"], GDP, "nritu-jk", "Q")

    # --- spot values
    if release == VERIFIED_RELEASE:
        for got, want, what in ((ncy[2024], 636_893.3, "nominal CY2024"),
                                (ncy[2025], 666_440.7, "nominal CY2025"),
                                (ncy[2006], 540_566.0, "nominal CY2006"),
                                (rq[(2026, 2)], 598_950.5, "real SA 2026Q2")):
            near(got, want, 0.05, f"{release} {what}")
    else:
        print(f"  !! release {release} is newer than {VERIFIED_RELEASE}: "
              "spot values checked as magnitudes only", file=sys.stderr)
        near(ncy[2024], 636_893.3, 15_000, "nominal CY2024 magnitude")
        near(ncy[2006], 540_566.0, 15_000, "nominal CY2006 magnitude")

    # --- second source: World Bank US$ GDP at the BoJ's 2024 mean rate
    usd = ncy[2024] / BOJ_USDJPY_2024
    if abs(usd / WB_GDP_2024_USD_BN - 1) > 0.015:
        raise RuntimeError(f"CY2024 ¥{ncy[2024]:,.1f}bn / {BOJ_USDJPY_2024} "
                           f"= US${usd:,.0f}bn, World Bank says "
                           f"US${WB_GDP_2024_USD_BN:,.0f}bn")
    print(f"  cross-check: ¥{ncy[2024]:,.1f}bn / {BOJ_USDJPY_2024} = "
          f"US${usd / 1000:.3f}tn vs World Bank US$"
          f"{WB_GDP_2024_USD_BN / 1000:.3f}tn "
          f"({(usd / WB_GDP_2024_USD_BN - 1) * 100:+.2f}%)")

    # --- annualised quarterly levels: mean ≈ annual, sum ≈ 4 × annual
    years = sorted(y for y in ncy if all((y, q) in nq for q in range(1, 5)))
    for lv, ann, what in ((nq, ncy, "nominal"), (rq, rcy, "real")):
        for y in years:
            qs = [lv[(y, q)] for q in range(1, 5)]
            mean = sum(qs) / 4
            if abs(mean / ann[y] - 1) > 0.01:
                raise RuntimeError(f"{what} {y}: mean of SA quarters "
                                   f"{mean:,.1f} vs annual {ann[y]:,.1f}")
            if sum(qs) / ann[y] < 3.9:
                raise RuntimeError(f"{what} {y}: quarters are not at "
                                   "annual rates any more")

    # --- published q/q against the levels (one decimal, rounding slack)
    for lv, pub, what in ((rq, rqq, "ritu-jk"), (nq, nqq, "ritu-mk")):
        for (y, q), v in pub.items():
            prev = (y, q - 1) if q > 1 else (y - 1, 4)
            if v is None or prev not in lv:
                continue
            near(v, (lv[(y, q)] / lv[prev] - 1) * 100, 0.06,
                 f"{what} {y}Q{q} vs levels")

    # --- annual table
    comp = {name: series(t["gaku-mcy"], en, "gaku-mcy", "A")
            for en, name in COMPONENTS}
    years = sorted(set(ncy) | set(nfy))
    write("macro/gdp_from_esri.csv", f"""
Cabinet Office ESRI, Quarterly Estimates release {release} (QE menu
resolved from sokuhou_top.html), {chain[1]} benchmark, 2008 SNA,
expenditure approach. Units: ¥ BILLION (10億円). Real = chained
{chain[1]} yen, ¥ billion.
*_cy columns: CALENDAR year (gaku-mcy / gaku-jcy, rows "YYYY/1-12.").
*_fy columns: FISCAL year April-March keyed on its STARTING year
(gaku-mfy / gaku-jfy, rows "YYYY/4-3."; fy 2024 = Apr 2024 - Mar 2025).
Calendar and fiscal columns are different objects; never mix them.
Annual figures are ESRI's annual tables, never summed quarters (the
quarterly SA levels are annualised). Component columns are nominal
calendar-year: exports and imports are goods and services; gni is
gross national income. A blank is a year the table does not cover.
Spot values (release {VERIFIED_RELEASE}): CY2024 636,893.3; CY2025
666,440.7; CY2006 540,566.0. Cross-check: CY2024 / 151.50 JPY per USD
(BoJ 2024 mean) = US${usd / 1000:.3f}tn vs World Bank US$4.190tn.
""", ["year", "gdp_nominal_cy_bn", "gdp_real_cy_bn", "gdp_nominal_fy_bn",
      "gdp_real_fy_bn"] + [f"{n}_nominal_cy_bn" for _, n in COMPONENTS],
          [[y, fmt(ncy.get(y)), fmt(rcy.get(y)), fmt(nfy.get(y)),
            fmt(rfy.get(y))] + [fmt(comp[n].get(y)) for _, n in COMPONENTS]
           for y in years])

    quarters = sorted(set(rq) | set(nq))
    write("macro/gdp_quarterly_from_esri.csv", f"""
Cabinet Office ESRI, Quarterly Estimates release {release}, seasonally
adjusted, ¥ BILLION AT ANNUAL RATES ("年率で表示"): each quarter's level
is a year's worth of output, so a calendar year is the MEAN of its four
quarters, never the sum (and annual figures belong to
gdp_from_esri.csv). gdp_real = gaku-jk, chained {chain[1]} yen;
gdp_nominal = gaku-mk. qoq columns are ESRI's PUBLISHED quarter-on-
quarter % changes (ritu-jk, ritu-mk); qoq_annualised is ESRI's published
annualised real rate (nritu-jk). Quarter = calendar quarter (Q1 = Jan-
Mar). This is the vintage the country.yaml bands were dated from; every
QE revises history, so re-date bands only on purpose.
""", ["quarter", "gdp_real_sa_saar_bn", "gdp_nominal_sa_saar_bn",
      "gdp_real_qoq_pct", "gdp_real_qoq_annualised_pct",
      "gdp_nominal_qoq_pct"],
          [[f"{y}-Q{q}", fmt(rq.get((y, q))), fmt(nq.get((y, q))),
            fmt(rqq.get((y, q))), fmt(rqa.get((y, q))),
            fmt(nqq.get((y, q)))] for y, q in quarters])

    gov_fy = series(t["gaku-mfy"], "GovernmentConsumption", "gaku-mfy", "A")
    return {"nominal_fy": nfy, "gov_consumption_fy": gov_fy,
            "release": release}


# ---------------------------------------------------------------- annual accounts

def kakuhou() -> tuple[int, dict[str, str]]:
    """(edition FY, {table: url}) for the latest annual accounts."""
    menu = fetch(SNA_MENU, "esri_sna_menu.html", is_html(b"kaku_top"),
                 fresh=True).decode("utf-8", "replace")
    eds = {int(y) for y, y2 in re.findall(
        r'kakuhou/files/(\d{4})/(\d{4})_kaku_top\.html', menu) if y == y2}
    if not eds:
        raise RuntimeError("sna/menu.html links no annual-accounts edition")
    fy = max(eds)
    url = f"{ESRI}/jp/sna/data/data_list/kakuhou/files/{fy}/{fy}_kaku_top.html"
    top = fetch(url, f"esri_kaku{fy}_top.html", is_html(b"_jp.xlsx"),
                fresh=True).decode("utf-8", "replace")
    out = {}
    for key, pat in (("s6_2", rf"tables/{fy}s6_2_jp\.xlsx"),
                     ("s7", rf"tables/{fy}s7_jp\.xlsx")):
        if not re.search(f'href="{pat}"', top):
            raise RuntimeError(f"FY{fy} top page does not link {pat}")
        out[key] = url.rsplit("/", 1)[0] + "/" + pat.replace("\\", "")
    return fy, out


def book(url: str, name: str):
    blob = fetch(url, name, is_xlsx)
    return openpyxl.load_workbook(io.BytesIO(blob), read_only=True,
                                  data_only=True)


def label(v) -> str:
    return unicodedata.normalize("NFKC", str(v or "")).strip()


SECTORS = ["中央政府", "地方政府", "社会保障基金", "部門間調整", "一般政府"]


def fy_blocks(rows: list[tuple], what: str) -> dict[int, dict[str, int]]:
    """{fiscal year: {sector: column}} from the two header rows."""
    hy = next(i for i, r in enumerate(rows[:8])
              if any(re.search(r"年度[（(]\d{4}[)）]", label(c)) for c in r))
    years, sect = rows[hy], rows[hy + 1]
    out = {}
    for i, c in enumerate(years):
        m = re.search(r"年度\((\d{4})\)", label(c))
        if not m:
            continue
        # The year label is centred over its block (it sits on the
        # block's second column), so look for the block start nearby.
        starts = [j for j in range(max(i - 2, 1), i + 1)
                  if [label(s) for s in sect[j:j + 5]] == SECTORS]
        if len(starts) != 1:
            raise RuntimeError(f"{what} FY{m[1]}: no five-sector block "
                               f"near column {i}")
        out[int(m[1])] = {s: starts[0] + k for k, s in enumerate(SECTORS)}
    return out


def find_row(rows, pattern: str, what: str, start: int = 0) -> tuple:
    hits = [r for r in rows[start:] if re.fullmatch(pattern, label(r[0]))]
    if len(hits) != 1:
        raise RuntimeError(f"{what}: {len(hits)} rows match {pattern!r}")
    return hits[0]


def gg_accounts(fy_ed: int, url: str, nominal_fy: dict) -> None:
    wb = book(url, f"esri_{url.rsplit('/', 1)[-1]}")
    for s in ("経常・資本取引", "金融資産・負債"):
        if s not in wb.sheetnames:
            raise RuntimeError(f"{url}: sheet {s!r} missing")
    cur = list(wb["経常・資本取引"].iter_rows(values_only=True))
    fin = list(wb["金融資産・負債"].iter_rows(values_only=True))
    if "10億円" not in label(cur[1][0]) and "10億円" not in label(fin[2][0]):
        raise RuntimeError("6-2: unit is not ¥bn (10億円)")
    bc, bf = fy_blocks(cur, "6-2 current"), fy_blocks(fin, "6-2 financial")
    if sorted(bc) != sorted(bf) or max(bc) != fy_ed:
        raise RuntimeError(f"6-2: years {min(bc)}–{max(bc)} vs edition {fy_ed}")

    want = {
        "revenue": (cur, r"1 収入"),
        "taxes": (cur, r"11 税"),
        "social_contributions": (cur, r"12 社会負担"),
        "expense": (cur, r"2 支出"),
        "interest": (cur, r"24 利子"),
        "social_benefits": (cur, r"27 社会給付"),
        "net_operating_balance": (cur, r"純業務収支.*"),
        "net_acq_nonfinancial_assets": (cur, r"31 非金融資産の純取得"),
        "net_lending": (cur, r"純貸出\(\+\)/純借入\(-\)"),
        "financial_assets": (fin, r"62 金融資産"),
        "liabilities": (fin, r"63 負債"),
        "net_financial_worth": (fin, r"金融資産負債差額"),
    }
    got = {k: find_row(sheet, pat, f"6-2 {k}") for k, (sheet, pat) in want.items()}

    def v(k, fy, sector="一般政府"):
        blocks = bc if want[k][0] is cur else bf
        return num(got[k][blocks[fy][sector]])

    years = sorted(bc)
    for fy in years:
        near(v("revenue", fy) - v("expense", fy), v("net_operating_balance", fy),
             0.3, f"6-2 FY{fy} revenue − expense vs net operating balance")
        near(v("net_operating_balance", fy) - v("net_acq_nonfinancial_assets", fy),
             v("net_lending", fy), 0.3, f"6-2 FY{fy} NOB − NANFA vs net lending")
    # Magnitudes: the GG deficit in % of FY GDP, and gross liabilities.
    nl = v("net_lending", fy_ed) / nominal_fy[fy_ed] * 100
    li = v("liabilities", fy_ed) / nominal_fy[fy_ed] * 100
    print(f"  FY{fy_ed}: GG net lending {nl:+.1f}% of FY GDP, "
          f"GFS liabilities {li:.0f}% of FY GDP")
    if not -12 < nl < 3:
        raise RuntimeError(f"GG net lending FY{fy_ed} = {nl:.1f}% of GDP")
    if not 180 < li < 280:
        raise RuntimeError(f"GG liabilities FY{fy_ed} = {li:.0f}% of GDP, "
                           "outside the IMF's 200-230% neighbourhood")

    cols = list(want)
    write("fiscal/gg_accounts_from_esri.csv", f"""
Cabinet Office ESRI, annual national accounts FY{fy_ed} edition (2020
benchmark), table 6-2 一般政府の部門別勘定（GFS）: general government in the
IMF GFSM 2014 layout. FISCAL YEAR April-March keyed on its STARTING
year (fiscal_year 2024 = Apr 2024 - Mar 2025). Units: ¥ BILLION.
gg_* columns are consolidated general government (一般政府).
net_operating_balance = revenue - expense (expense includes consumption
of fixed capital); net_lending = net operating balance - net
acquisition of non-financial assets, from the non-financial accounts
(the financial account's own net-lending row differs by a statistical
discrepancy and is not used). financial_assets, liabilities and
net_financial_worth are END-OF-YEAR STOCKS at market value (GFS
balance sheet), not flows; pension entitlements of social security
funds are not recorded as liabilities (ESRI note 4). central/local/ssf
net lending are the unconsolidated subsectors. This is the national-
accounts general government; MOF's general account is a narrower, cash
object, and the two are never spliced.
""", ["fiscal_year"] + [f"gg_{c}_bn" for c in cols]
          + ["central_net_lending_bn", "local_net_lending_bn",
             "ssf_net_lending_bn"],
          [[fy] + [fmt(v(c, fy)) for c in cols]
           + [fmt(v("net_lending", fy, s)) for s in SECTORS[:3]]
           for fy in years])


ERA = {"平成": 1988, "令和": 2018}   # 平成1 = 1989, 令和1 = 2019

FUNCTIONS = [
    (1, "general_public_services"), (2, "defence"), (3, "public_order"),
    (4, "economic_affairs"), (5, "environment"), (6, "housing_community"),
    (7, "health"), (8, "recreation_culture"), (9, "education"),
    (10, "social_protection"),
]
# Transaction columns of table 7 (0-based), read off the header rows and
# asserted. Column 1 is final consumption, columns 2-3 its split into
# individual and collective consumption (not added again).
TRANSACTIONS = {1: "最終消費支出", 4: "補助金", 5: "現物社会移転", 6: "その他の",
                7: "総固定資本形成", 8: "在庫変動", 9: "資本移転"}


def wareki_fy(sheet: str) -> int:
    s = label(sheet).replace("元", "1")
    m = re.fullmatch(r"(平成|令和)(\d+)", s)
    if not m:
        raise RuntimeError(f"COFOG sheet {sheet!r} is not a wareki year")
    return ERA[m[1]] + int(m[2])


def cofog(fy_ed: int, url: str, gov_fy: dict) -> None:
    if wareki_fy("令和元") != 2019 or wareki_fy("平成１７") != 2005:
        raise RuntimeError("wareki table is wrong")
    wb = book(url, f"esri_{url.rsplit('/', 1)[-1]}")
    out = []
    for sheet in wb.sheetnames:
        fy = wareki_fy(sheet)
        rows = list(wb[sheet].iter_rows(values_only=True))
        if "10億円" not in label(rows[3][0]):
            raise RuntimeError(f"COFOG {sheet}: unit {rows[3][0]!r}")
        m = re.search(r"\((\d{4})\)", label(rows[3][1]))
        if not m or int(m[1]) != fy:
            raise RuntimeError(f"COFOG {sheet}: header year {rows[3][1]!r} "
                               f"≠ wareki FY{fy}")
        head = [label(a) + label(b) for a, b in zip(rows[4], rows[5])]
        for i, name in TRANSACTIONS.items():
            if not head[i].startswith(name):
                raise RuntimeError(f"COFOG {sheet}: column {i} is "
                                   f"{head[i]!r}, expected {name!r}")
        top, total = {}, None
        for r in rows[7:]:
            lab = label(r[0])
            if m := re.fullmatch(r"(\d+)\.(\D.*)", lab):
                top[int(m[1])] = r
            elif re.fullmatch(r"合\s*計", lab):
                total = r
        if sorted(top) != list(range(1, 11)) or total is None:
            raise RuntimeError(f"COFOG {sheet}: functions {sorted(top)}")
        for i in TRANSACTIONS:
            near(sum(num(top[f][i]) or 0 for f in top), num(total[i]), 0.6,
                 f"COFOG FY{fy} column {i}: functions vs 合計")

        def tot(r):
            return sum(num(r[i]) or 0 for i in TRANSACTIONS)

        out.append([fy] + [fmt(tot(top[f])) for f, _ in FUNCTIONS]
                   + [fmt(tot(total)), fmt(num(total[1])),
                      fmt(num(top[10][5]))])
        if fy == fy_ed and fy in gov_fy:
            # Second source within ESRI: final consumption of general
            # government in COFOG vs the QE fiscal-year table.
            if abs(num(total[1]) / gov_fy[fy] - 1) > 0.01:
                raise RuntimeError(f"COFOG FY{fy} final consumption "
                                   f"{num(total[1])} vs QE {gov_fy[fy]}")
    out.sort()
    if [r[0] for r in out] != list(range(out[0][0], out[-1][0] + 1)):
        raise RuntimeError("COFOG: fiscal years are not contiguous")
    write("fiscal/cofog_from_esri.csv", f"""
Cabinet Office ESRI, annual national accounts FY{fy_ed} edition, table 7
一般政府の機能別支出（COFOG）, one sheet per fiscal year (sheet names in
wareki, 平成17 = FY2005 ... 令和6 = FY2024). FISCAL YEAR April-March keyed
on its STARTING year. Units: ¥ BILLION. COMPUTED HERE: each function's
figure is the sum of the seven published transaction columns (final
consumption + subsidies + social benefits other than social transfers
in kind + other current transfers + gross fixed capital formation +
change in inventories + capital transfers); ESRI publishes no total
column. Interest (property income) and net land purchases are NOT in
the table while gross fixed capital formation IS, so total_expenditure
here is a different aggregate from the GFS expense of
gg_accounts_from_esri.csv (higher than it in recent years).
The ten functions add up to ESRI's 合計 row in every column (asserted).
final_consumption_total is ESRI's published 合計 of column 1.
social_benefits_cash_social_protection is the published cell for
function 10, column 5 (pensions and other cash benefits).
""", ["fiscal_year"] + [f"{n}_bn" for _, n in FUNCTIONS]
          + ["total_expenditure_bn", "final_consumption_total_bn",
             "social_benefits_cash_social_protection_bn"], out)


def main() -> None:
    print("Cabinet Office ESRI — QE CSVs and annual accounts")
    got = gdp()
    fy_ed, tables = kakuhou()
    print(f"  annual accounts edition FY{fy_ed}")
    gg_accounts(fy_ed, tables["s6_2"], got["nominal_fy"])
    cofog(fy_ed, tables["s7"], got["gov_consumption_fy"])
    print("Done. Diff the sidecars before adopting.")


if __name__ == "__main__":
    main()
