#!/usr/bin/env python3
"""Homicide and violence series from IPEA's Atlas da Violência and FBSP.

Usage:  python3 scripts/fetch_violence.py

Outputs:
  data/social/homicides.csv      IPEA Atlas (health system / SIM), 2006-
  data/social/violence_fbsp.csv  FBSP Anuário (police records), 2012-

The two are different universes and must not be spliced:

  Atlas  = SIM death certificates, CID-10 X85-Y09 (agressões) + Y35-Y36
           (intervenção legal), counted by victim's RESIDENCE.
  FBSP   = state police records, MVI = homicídio doloso + latrocínio +
           lesão corporal seguida de morte + mortes por intervenção policial,
           counted by OCCURRENCE.

Atlas runs the full 2006- window; FBSP only starts in 2012 but publishes
~18 months sooner, so it is the only source for the most recent year.

Two traps this encodes:
  - "registered" vs "estimated" homicides are both official. Registered is
    SIM as filed; estimated adds a modeled share of deaths filed as
    undetermined intent (MVCI). The gap reached 17% in 2024.
  - police killings: use FBSP, never SIM's Y35. SIM captures roughly a third
    of them, because the ICD code follows the death certificate rather than
    the police classification.
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import requests

ROOT = Path(__file__).resolve().parent.parent
CACHE = ROOT / "cache"
DATA = ROOT / "data"
(CACHE / "ipea").mkdir(parents=True, exist_ok=True)
(CACHE / "fbsp").mkdir(parents=True, exist_ok=True)

# IPEA's Atlas API is undocumented but stable and needs no auth.
IPEA = "https://www.ipea.gov.br/dados-api/series-values/{sid}/1"
BRAZIL = 1076          # regiao_id for Brazil at abrangencia 1 (countries)
SERIES = {
    "count": 328,      # homicídios registrados
    "rate": 20,        # taxa por 100 mil
    "hidden": 365,     # homicídios ocultos (modeled, from MVCI)
    "est_count": 366,  # homicídios estimados = registrados + ocultos
    "est_rate": 368,
}
YEAR_FROM = 2006

FBSP_URL = "https://forumseguranca.org.br/wp-content/uploads/2026/07/anuario-2026.xlsx"


def ipea_series(sid: int) -> dict[int, float]:
    cache_file = CACHE / "ipea" / f"series_{sid}.json"
    if cache_file.exists():
        payload = json.loads(cache_file.read_text())
    else:
        r = requests.get(IPEA.format(sid=sid), timeout=180,
                         headers={"Accept": "application/json"})
        r.raise_for_status()
        payload = r.json()
        cache_file.write_text(json.dumps(payload))
    return {int(row["periodo"][:4]): row["valor"] for row in payload
            if row.get("regiao_id") == BRAZIL and row.get("valor") is not None}


def fetch_atlas() -> None:
    data = {name: ipea_series(sid) for name, sid in SERIES.items()}
    for name, sid in SERIES.items():
        yrs = sorted(data[name])
        print(f"  [ipea {sid}] {name}: {len(yrs)} years {yrs[0]}-{yrs[-1]}")
    last = max(data["count"])
    lines = ["year,count,rate,hidden,est_count,est_rate"]
    for y in range(YEAR_FROM, last + 1):
        def val(k: str, dec: int) -> str:
            v = data[k].get(y)
            return f"{v:.{dec}f}" if v is not None else ""
        lines.append(f"{y},{val('count', 0)},{val('rate', 2)},"
                     f"{val('hidden', 0)},{val('est_count', 0)},"
                     f"{val('est_rate', 2)}")
    out = DATA / "social" / "homicides.csv"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(lines) + "\n")
    print(f"  -> data/social/homicides.csv ({YEAR_FROM}-{last})")


def fetch_fbsp() -> None:
    path = CACHE / "fbsp" / "anuario-2026.xlsx"
    if not path.exists():
        print(f"  downloading {FBSP_URL}")
        r = requests.get(FBSP_URL, timeout=300,
                         headers={"User-Agent": "brasil-reports/1.0"})
        r.raise_for_status()
        path.write_bytes(r.content)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    # T02: Brasil row, counts in one column block and rates in the next, each
    # block headed by its own repeated run of years on row 7
    ws = wb["T02"]
    rows = list(ws.iter_rows(min_row=1, max_row=10, values_only=True))
    years_row = rows[6]
    brasil = next(r for r in rows if r[0] == "Brasil")
    pairs = [(i, y) for i, y in enumerate(years_row) if isinstance(y, int)]
    half = len(pairs) // 2
    mvi = {y: brasil[i] for i, y in pairs[:half]}
    mvi_rate = {y: brasil[i] for i, y in pairs[half:]}

    # P03 holds several tables side by side — the national series in columns
    # B/C and state rankings from column D on — so read by position; dropping
    # empty cells first would slide the rankings into the series.
    police = {}
    for row in wb["P03"].iter_rows(values_only=True):
        if (len(row) > 2 and isinstance(row[1], int) and 2000 < row[1] < 2100
                and isinstance(row[2], (int, float))):
            police[row[1]] = row[2]

    # P06: years on one row, then a labelled row per indicator
    ws = wb["P06"]
    prows = list(ws.iter_rows(min_row=1, max_row=20, values_only=True))
    yr_row = next(r for r in prows
                  if sum(isinstance(c, int) and 2000 < c < 2100 for c in r) > 5)
    cols = [(i, y) for i, y in enumerate(yr_row)
            if isinstance(y, int) and 2000 < y < 2100]

    def prow(label: str) -> dict[int, float]:
        # match the label exactly: "Feminicídios" is also a substring of
        # "Homicídios (incluindo feminicídios)", and a loose match silently
        # returns the women's-homicides row for both indicators
        row = next((r for r in prows
                    if any(isinstance(c, str) and c.strip().lower() == label
                           for c in r)), None)
        return {y: row[i] for i, y in cols} if row else {}

    women = prow("homicídios (incluindo feminicídios)")
    femi = prow("feminicídios")

    years = sorted(set(mvi) | set(police) | set(women))
    lines = ["year,mvi,mvi_rate,police_killings,women_homicides,femicides"]
    for y in years:
        def val(d: dict, dec: int) -> str:
            v = d.get(y)
            return f"{v:.{dec}f}" if isinstance(v, (int, float)) else ""
        lines.append(f"{y},{val(mvi, 0)},{val(mvi_rate, 2)},"
                     f"{val(police, 0)},{val(women, 0)},{val(femi, 0)}")
    (DATA / "social" / "violence_fbsp.csv").write_text("\n".join(lines) + "\n")
    print(f"  -> data/social/violence_fbsp.csv ({years[0]}-{years[-1]})")


def main() -> None:
    print("IPEA Atlas da Violência (SIM / health system)...")
    fetch_atlas()
    print("FBSP Anuário Brasileiro de Segurança Pública (police records)...")
    fetch_fbsp()
    print("Done.")


if __name__ == "__main__":
    main()
