#!/usr/bin/env python3
"""Refresh the cross-country comparables for Vietnam into sidecar CSVs.

Usage:  python3 -I countries/vietnam/fetch/fetch_comparables.py

These are the layer that makes Vietnam legible to a reader who does not
live in it: the same quantity, computed the same way, for everyone. Every
file here is `official: false` (a redistributor or a model), and none is
a substitute for a number the NSO publishes about Vietnam. Where both
exist (GDP, population, the Gini, trade) the NSO series is the curated one
and this is the cross-check, and the plate labels the difference rather
than choosing.

  World Bank WDI  GDP (level, per person, growth), openness, official FX,
                  reserves, Gini/poverty, population (UN WPP), homicide
                  (UNODC), private credit, net migration (UN WPP)
  ADB KIDB        DF_ER end-of-period VND/USD, the second column of the FX
                  sidecar (the manifest pairs it with the WB average)
  IMF DataMapper  WEO: debt and net lending; NGDPD, growth, inflation,
                  current account, unemployment (projections apart)
  UNDP            HDR25 composite indices (one vintage)
  Transparency    CPI 2025 workbook, comparable 2012-method series
  UN DESA         International Migrant Stock 2024, Table 1
  UN WPP          World Population Prospects 2024, medium variant
  WHO             GHO RS_198 (road deaths), VIOLENCE_HOMICIDERATE, and
                  GHE_FULL deaths by cause: all MODELLED estimates
  UN Comtrade     public preview API, Vietnam as reporter, HS TOTAL

Things these sources will mislead you about.

**The 2010 GDP revaluation is inside the World Bank series.** WDI copies
the NSO's levels: 2006-2009 are on the old scale and 2010+ on the revised
one, so US$106.0bn (2009) becomes US$147.2bn (2010), +39% in a year of
~6% real growth. The level and per-person columns are SPLIT at 2010 into
`*_old_scale` (2006-2009) and the current-scale column (2010+); they never
share a column. Ratios with GDP underneath (openness, private credit)
cannot be split without re-basing, so they stay one column and the seam is
flagged in the CSV comment. The IMF NGDPD is back-cast across the seam
and is the continuous level; it is labelled as the IMF's, never spliced
onto the World Bank's.

**IMF DataMapper.** A "Mozilla/5.0" User-Agent gets an Akamai 403; curl's
default UA passes, so that is what is sent. The /VNM path does not filter:
the reply carries every country and VNM is read out of it. Years after
last_full_year are WEO projections and live in `*_proj` columns only.

**The World Bank's personal remittances are EMPTY for Vietnam**
(BX.TRF.PWKR.CD.DT has no rows; the %GDP version stops in 2004). The
fetcher checks that and says so; it does not write a remittance file. The
manifest's labelled proxy is ADB secondary-income credit
(migration/remittances_proxy_from_adb.csv, refreshed by fetch_adb.py),
"secondary income received (includes remittances)", never "remittances".

**UN WPP.** The Data Portal API named in the manifest now answers 401
(Bearer token) to every data query; only its metadata is anonymous. The
anonymous route is the WPP bulk CSV, resolved from the site's own
assets/downloads.json rather than hardcoded.

**UN DESA's xlsx** answers a bare client with an HTML error page under
HTTP 200, so a browser UA is sent and the PK zip magic is checked before
the payload is cached.

**UN Comtrade** answers 429 after two quick calls and allows one period
per preview call, so the fetcher sleeps between years and backs off.

**WHO numbers are models.** Road deaths before 2022 are about three times
the police count; homicide is a Global Health Estimates model for a
country with no national homicide statistic; causes of death are the only
cause series Vietnam has, and they are estimates.
"""
from __future__ import annotations

import csv
import gzip
import io
import json
import re
import sys
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
CACHE = ROOT / "cache"
PREFIX = "cmp_"          # this fetcher's cache files, apart from fetch_adb's

ISO3, M49 = "VNM", 704
YEAR_FROM, YEAR_TO = 2006, 2026        # country.yaml years
LAST_FULL_YEAR = 2025                  # country.yaml last_full_year
SEAM = 2010                            # first year on the revised GDP scale

BROWSER = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
                         "(KHTML, like Gecko) Chrome/124.0 Safari/537.36",
           "Accept-Language": "en;q=0.8"}
# www.imf.org's Akamai edge 403s any UA starting "Mozilla/5.0" (verified
# 2026-10-22); curl's default UA is what DISCOVERY.md verified.
CURL = {"User-Agent": "curl/8.9.1", "Accept": "*/*"}

WB = ("https://api.worldbank.org/v2/country/VNM/indicator/{code}"
      "?format=json&per_page=200&date={span}")
ADB_ER = ("https://kidb.adb.org/api/v5/sdmx/data/ADB,DF_ER/A..VIE"
          "?format=sdmx-csv")
IMF = "https://www.imf.org/external/datamapper/api/v1/{code}/VNM"
HDR = ("https://hdr.undp.org/sites/default/files/2025_HDR/"
       "HDR25_Composite_indices_complete_time_series.csv")
DESA = ("https://www.un.org/development/desa/pd/sites/"
        "www.un.org.development.desa.pd/files/"
        "undesa_pd_2024_ims_stock_by_sex_destination_and_origin.xlsx")
WPP_SITE = "https://population.un.org/wpp/"
WPP_INDEX = WPP_SITE + "assets/downloads.json"
WPP_FILE = re.compile(r"WPP(\d{4})_Demographic_Indicators_Medium\.csv\.gz$")
GHO = "https://ghoapi.azureedge.net/api/{code}?"
GHE = "https://xmart-api-public.who.int/DEX_CMS/GHE_FULL?"
COMTRADE = ("https://comtradeapi.un.org/public/v1/preview/C/A/HS"
            "?reporterCode=704&period={y}&partnerCode=0&flowCode=X,M"
            "&cmdCode=TOTAL")
TI = "https://images.transparencycdn.org/images/CPI2025_Results.xlsx"

GAP = 1.5
_last: dict[str, float] = {}

OFFICIAL_FALSE = "official: false."


# ---------------------------------------------------------------- transport

def _get(url: str, headers: dict, gap: float = GAP) -> bytes:
    """One request, spaced per host and retried; raises on persistent failure.

    429 and 403 back off hard (Comtrade's 429 says "try again in 2
    seconds"; Akamai/CDN 403s sometimes clear); a 404 or 401 does not heal.
    """
    host = urllib.parse.urlsplit(url).netloc
    last = None
    for attempt in range(5):
        wait = _last.get(host, 0.0) + gap - time.monotonic()
        if wait > 0:
            time.sleep(wait)
        try:
            req = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(req, timeout=300) as r:
                blob = r.read()
            _last[host] = time.monotonic()
            return blob
        except urllib.error.HTTPError as e:
            _last[host] = time.monotonic()
            last = e
            if e.code in (400, 401, 404):
                break
            time.sleep((10 if e.code == 429 else 5) * (attempt + 1))
        except OSError as e:
            _last[host] = time.monotonic()
            last = e
            time.sleep(5 * (attempt + 1))
    raise RuntimeError(f"{url}: {last}")


def fetch(url: str, name: str, check, headers=BROWSER, gap: float = GAP) -> bytes:
    """Cached download; `check(blob)` raises before anything is cached."""
    CACHE.mkdir(exist_ok=True)
    local = CACHE / f"{PREFIX}{name}"
    if local.exists():
        return local.read_bytes()
    print(f"  downloading {url}")
    blob = _get(url, headers, gap)
    check(blob)
    local.write_bytes(blob)
    return blob


def check_json(pred, what: str):
    def check(blob: bytes) -> None:
        try:
            payload = json.loads(blob.decode("utf-8-sig"))
        except ValueError as e:
            raise RuntimeError(f"{what}: not JSON ({blob[:80]!r})") from e
        if not pred(payload):
            raise RuntimeError(f"{what}: unexpected payload shape")
    return check


def load_json(blob: bytes):
    return json.loads(blob.decode("utf-8-sig"))


def fail(msg: str) -> None:
    raise RuntimeError(msg)


def near(got, want: float, tol: float, what: str) -> None:
    if got is None or abs(got - want) > tol:
        fail(f"{what} = {got}, expected {want} ± {tol}")
    shown = f"{got:,.0f}" if abs(got) >= 1000 else f"{got:g}"
    print(f"  ok  {what} = {shown}")


# ---------------------------------------------------------------- writers

def fmt(v, nd: int = 1) -> str:
    """A cell: blank for a hole, never zero; no exponent, no separators."""
    if v is None:
        return ""
    if nd == 0:
        return str(int(round(v)))
    return f"{round(v, nd):.{nd}f}"


def write(rel: str, comment: str, header: list[str], rows: list[list]) -> int:
    if not rows:
        fail(f"{rel}: no rows to write")
    path = DATA / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        for line in comment.strip().splitlines():
            fh.write(f"# {line.strip()}".rstrip() + "\n")
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)
    span = f"{rows[0][0]}–{rows[-1][0]}"
    print(f"  data/{rel}  {len(rows)} rows ({span})")
    return len(rows)


def nfc(s: str) -> str:
    return unicodedata.normalize("NFC", s or "").strip()


# ---------------------------------------------------------------- World Bank

def wb(code: str, span: str = f"{YEAR_FROM}:2025") -> tuple[dict, str]:
    """One WDI indicator for Vietnam as ({year: value}, lastupdated).

    A bad request answers with a JSON object carrying a message, not the
    two-element array of a good one, so the shape is checked before
    caching. An indicator with no rows for Vietnam answers [meta, null].
    """
    blob = fetch(WB.format(code=code, span=span), f"wb_{code}.json",
                 check_json(lambda p: isinstance(p, list) and len(p) == 2
                            and isinstance(p[0], dict) and "page" in p[0],
                            f"WB {code}"))
    meta, rows = load_json(blob)
    if int(meta.get("pages", 1)) > 1:
        fail(f"WB {code}: more than one page; raise per_page")
    out = {}
    for row in rows or []:
        if row.get("countryiso3code") not in (ISO3, ""):
            fail(f"WB {code}: row for {row.get('countryiso3code')}")
        if row.get("value") is not None:
            out[int(row["date"])] = float(row["value"])
    return out, meta.get("lastupdated", "")


def wb_window(*series: dict) -> list[int]:
    years = set()
    for s in series:
        years |= set(s)
    return sorted(y for y in years if YEAR_FROM <= y <= YEAR_TO)


def world_bank() -> dict:
    s = {}
    upd = ""
    codes = ["NY.GDP.MKTP.CD", "NY.GDP.PCAP.CD", "NY.GDP.MKTP.KD.ZG",
             "NE.TRD.GNFS.ZS", "PA.NUS.FCRF", "FI.RES.TOTL.CD",
             "SI.POV.GINI", "SI.POV.DDAY", "SI.POV.NAHC",
             "SP.POP.TOTL", "SP.POP.65UP.TO.ZS", "SP.DYN.TFRT.IN",
             "SP.DYN.LE00.IN", "FS.AST.PRVT.GD.ZS", "SM.POP.NETM",
             "BX.TRF.PWKR.CD.DT", "BX.TRF.PWKR.DT.GD.ZS"]
    for code in codes:
        s[code], upd = wb(code)
    s["VC.IHR.PSRC.P5"], _ = wb("VC.IHR.PSRC.P5", "2000:2025")
    head = (f"Source: World Bank, World Development Indicators (API v2, "
            f"lastupdated {upd}). {OFFICIAL_FALSE} A redistributor.")

    # --- GDP: the seam is split into two columns, never joined.
    gdp, pc, gr = s["NY.GDP.MKTP.CD"], s["NY.GDP.PCAP.CD"], s["NY.GDP.MKTP.KD.ZG"]
    near(gdp.get(2024, 0) / 1e9, 476.3, 2.0, "WB GDP 2024 (US$bn)")
    near(gdp.get(2025, 0) / 1e9, 514.7, 3.0, "WB GDP 2025 (US$bn)")
    near(gdp.get(2009, 0) / 1e9, 106.0, 1.0, "WB GDP 2009 (US$bn, old scale)")
    near(gdp.get(2010, 0) / 1e9, 147.2, 1.0, "WB GDP 2010 (US$bn, new scale)")
    seam = gdp[2010] / gdp[2009] - 1
    near(seam * 100, 39.0, 1.5, "WB GDP 2009→2010 step (%)")
    if not (gr.get(2010) and gr[2010] < 10):
        fail(f"WB real growth 2010 = {gr.get(2010)}: the seam is a level "
             "break, growth should be single-digit")
    near(pc.get(2024), 4716.7, 30, "WB GDP per person 2024 (US$)")
    rows = [[y,
             fmt(gdp.get(y), 0) if y < SEAM else "",
             fmt(gdp.get(y), 0) if y >= SEAM else "",
             fmt(pc.get(y), 1) if y < SEAM else "",
             fmt(pc.get(y), 1) if y >= SEAM else "",
             fmt(gr.get(y), 2)] for y in wb_window(gdp, pc, gr)]
    write("international/gdp_usd_from_wb.csv", f"""
{head} NY.GDP.MKTP.CD, NY.GDP.PCAP.CD, NY.GDP.MKTP.KD.ZG.
Units: current US$ (level, per person); real growth in % per year.
BREAK: the NSO's 2010 GDP revaluation. 2006-2009 are on the OLD scale (*_old_scale), 2010+ on the revised scale; 2009->2010 reads +{seam * 100:.0f}% in a year of {gr[2010]:.1f}% real growth.
The two scales are separate columns and are never joined; for a continuous level use the IMF's back-cast (gdp_usd_from_imf.csv), labelled as the IMF's.
Real growth is computed within each scale by the source and is one column.
""", ["year", "gdp_usd_old_scale", "gdp_usd", "gdp_per_person_usd_old_scale",
      "gdp_per_person_usd", "real_growth_pct"], rows)

    # --- openness
    op = s["NE.TRD.GNFS.ZS"]
    near(op.get(2024), 173.9, 0.5, "WB openness 2024 (% GDP)")
    near(op.get(2006), 138.3, 0.5, "WB openness 2006 (% GDP)")
    write("trade/openness_from_wb.csv", f"""
{head} NE.TRD.GNFS.ZS: exports plus imports of goods and services, % of GDP.
SEAM FLAG: the denominator changes scale in 2010 (NSO GDP revaluation, +39% in WB current US$). Much of the 2006-2010 fall is the revaluation, not a fall in trade. Annotate it; do not read it as a trade collapse.
""", ["year", "trade_pct_gdp"],
          [[y, fmt(op[y], 1)] for y in wb_window(op)])

    # --- FX: WB annual average + ADB end of period, two columns
    fx = s["PA.NUS.FCRF"]
    near(fx.get(2020), 23208.37, 1.0, "WB VND/USD average 2020")
    near(fx.get(2024), 24164.89, 1.0, "WB VND/USD average 2024")
    ende = adb_end_period()
    near(ende.get(2025), 25121, 5, "ADB VND/USD end of period 2025")
    years = sorted(y for y in set(fx) | set(ende) if YEAR_FROM <= y <= YEAR_TO)
    write("markets/usdvnd_from_wb.csv", f"""
{head} PA.NUS.FCRF, official exchange rate, VND per US$, ANNUAL AVERAGE (from IMF IFS).
Column usdvnd_end_period: Asian Development Bank KIDB DF_ER ENDE_XDC_USD_RATE, VND per US$, END OF PERIOD ({OFFICIAL_FALSE} source field: IMF).
Average and end-of-period are two concepts in two columns, never spliced; a blank is a year that source has not published.
VND per USD: a RISING line is a WEAKER dong.
""", ["year", "usdvnd_avg", "usdvnd_end_period"],
          [[y, fmt(fx.get(y), 2), fmt(ende.get(y), 0)] for y in years])

    # --- reserves
    res = s["FI.RES.TOTL.CD"]
    near(res.get(2025, 0) / 1e9, 85.6, 0.5, "WB reserves 2025 (US$bn)")
    near(res.get(2006, 0) / 1e9, 13.4, 0.3, "WB reserves 2006 (US$bn)")
    write("international/reserves_from_wb.csv", f"""
{head} FI.RES.TOTL.CD: total reserves including gold, year-end stock, current US$ (from IMF IFS).
The SBV publishes no reserve stock series; the NSO's BoP carries only the annual change, which is never drawn as a level.
""", ["year", "reserves_usd"],
          [[y, fmt(res[y], 0)] for y in wb_window(res)])

    # --- inequality, survey years only
    gi, dd, na = s["SI.POV.GINI"], s["SI.POV.DDAY"], s["SI.POV.NAHC"]
    near(gi.get(2022), 36.1, 0.15, "WB Gini 2022")
    near(gi.get(2006), 35.8, 0.15, "WB Gini 2006")
    near(na.get(2022), 4.2, 0.15, "WB national poverty headcount 2022")
    near(na.get(2010), 20.7, 0.15, "WB national poverty headcount 2010")
    if any(y not in gi for y in dd):
        print("  note: poverty-line years not all Gini years", file=sys.stderr)
    write("social/inequality_from_wb.csv", f"""
{head} SI.POV.GINI (Gini, 0-100), SI.POV.DDAY (% of population below the international poverty line, current PIP line, 2021 PPP), SI.POV.NAHC (% of POPULATION below the national GSO-World Bank expenditure line).
Survey (VHLSS) years only; the years between are holes and stay blank. Nothing is interpolated.
NAHC is a population share on the expenditure line, NOT the NSO's household multidimensional share (V14.46). The 2010 jump in NAHC is a NEW POVERTY LINE, not a rise in poverty: treat 2006-2008 and 2010+ as separate segments.
The World Bank Gini is the harmonised layer; the NSO's Gini (V14.38) is a different computation with its own 2019->2020 step.
""", ["year", "gini", "poverty_intl_line_pct", "poverty_national_line_pct"],
          [[y, fmt(gi.get(y), 1), fmt(dd.get(y), 1), fmt(na.get(y), 1)]
           for y in wb_window(gi, dd, na)])

    # --- population (WPP via WB)
    pop, p65 = s["SP.POP.TOTL"], s["SP.POP.65UP.TO.ZS"]
    tfr, le = s["SP.DYN.TFRT.IN"], s["SP.DYN.LE00.IN"]
    near(pop.get(2024, 0) / 1e6, 101.0, 0.3, "WB population 2024 (mn)")
    near(tfr.get(2024), 1.90, 0.03, "WB TFR 2024")
    write("demography/population_from_wb.csv", f"""
{head} SP.POP.TOTL (persons, mid-year), SP.POP.65UP.TO.ZS (% aged 65+), SP.DYN.TFRT.IN (births per woman), SP.DYN.LE00.IN (years).
These are UN WPP 2024 ESTIMATES redistributed by the Bank, not the NSO's survey/census figures (NSO average population 2024: 101.34mn). Cross-check and ageing panel only.
""", ["year", "population", "pop_65plus_pct", "tfr", "life_expectancy"],
          [[y, fmt(pop.get(y), 0), fmt(p65.get(y), 2), fmt(tfr.get(y), 3),
            fmt(le.get(y), 2)] for y in wb_window(pop, p65, tfr, le)])

    # --- homicide (UNODC via WB): ends 2011, not extended
    ho = s["VC.IHR.PSRC.P5"]
    near(ho.get(2006), 1.28, 0.02, "WB/UNODC homicide 2006")
    near(ho.get(2010), 1.53, 0.02, "WB/UNODC homicide 2010")
    if max(ho) != 2011:
        fail(f"WB homicide now ends {max(ho)} (was 2011): re-check the note")
    write("security/homicide_from_wb.csv", f"""
{head} VC.IHR.PSRC.P5: intentional homicides per 100,000 people (UNODC CTS).
Covers {min(ho)}-{max(ho)} ONLY and is not extended by anything. The WHO homicide model (homicide_from_who.csv) is a different series and is never spliced onto it.
""", ["year", "homicide_per_100k"],
          [[y, fmt(ho[y], 2)] for y in sorted(ho)])

    # --- private credit
    pcr = s["FS.AST.PRVT.GD.ZS"]
    near(pcr.get(2022), 125.0, 0.5, "WB private credit 2022 (% GDP)")
    near(pcr.get(2006), 65.4, 0.5, "WB private credit 2006 (% GDP)")
    if max(pcr) != 2022:
        print(f"  note: WB private credit now ends {max(pcr)}", file=sys.stderr)
    write("credit/private_credit_from_wb.csv", f"""
{head} FS.AST.PRVT.GD.ZS: domestic credit to the private sector, % of GDP. Ends {max(pcr)}.
SEAM FLAG: part of the 2009->2010 step is the GDP denominator changing scale (NSO revaluation). Annotate it.
""", ["year", "private_credit_pct_gdp"],
          [[y, fmt(pcr[y], 1)] for y in wb_window(pcr)])

    # --- net migration (WPP via WB), and the empty remittance series
    nm = s["SM.POP.NETM"]
    near(nm.get(2006), 482399, 1, "WB net migration 2006 (the artefact)")
    near(nm.get(2025), -48171, 1, "WB net migration 2025")
    rem, rem_pct = s["BX.TRF.PWKR.CD.DT"], s["BX.TRF.PWKR.DT.GD.ZS"]
    if rem:
        print(f"  NOTE: BX.TRF.PWKR.CD.DT now has {len(rem)} rows for VNM "
              "(was empty): tell the lead; not written here.", file=sys.stderr)
    else:
        print("  ok  BX.TRF.PWKR.CD.DT is empty for VNM (as recorded); "
              "proxy is ADB secondary-income credit via fetch_adb.py")
    if rem_pct:
        print(f"  note: BX.TRF.PWKR.DT.GD.ZS has {sorted(rem_pct)} in window",
              file=sys.stderr)
    write("migration/net_migration_from_wb.csv", f"""
{head} SM.POP.NETM: net migration, persons, UN WPP 2024 modelled estimate. A residual estimate, not a count.
2006 READS +{nm[2006]:,.0f}: implausible as an inflow into a country of emigration; an artefact of the WPP estimate (cause unverified). Kept as published; the plate starts the line in 2007 or annotates the point.
Remittances: the World Bank's personal-remittance series (BX.TRF.PWKR.CD.DT) is EMPTY for Vietnam. The labelled proxy is migration/remittances_proxy_from_adb.csv, secondary income received (includes remittances), never called remittances.
""", ["year", "net_migration"],
          [[y, fmt(nm[y], 0)] for y in wb_window(nm)])
    return s


def adb_end_period() -> dict[int, float]:
    def check(blob: bytes) -> None:
        head = blob[:400].decode("utf-8-sig", "replace")
        if not head.startswith("DATAFLOW,") or "ENDE_XDC_USD_RATE" not in \
                blob.decode("utf-8", "replace"):
            fail("ADB DF_ER: not an SDMX-CSV with ENDE_XDC_USD_RATE")
    blob = fetch(ADB_ER, "adb_DF_ER.csv", check)
    out = {}
    for r in csv.DictReader(io.StringIO(blob.decode("utf-8-sig"))):
        if r["INDICATOR"] != "ENDE_XDC_USD_RATE" or r["ECONOMY_CODE"] != "VIE":
            continue
        if r["UNIT"] != "NCU_PER_USD":
            fail(f"ADB DF_ER unit is {r['UNIT']}")
        if r["OBS_VALUE"]:
            out[int(r["TIME_PERIOD"])] = float(r["OBS_VALUE"]) * \
                10 ** int(r["UNIT_MULT"] or 0)
    return out


# ---------------------------------------------------------------- IMF

def imf(code: str) -> dict[int, float]:
    blob = fetch(IMF.format(code=code), f"imf_{code}.json",
                 check_json(lambda p: isinstance(p.get("values", {})
                                                 .get(code, {}).get(ISO3),
                                                 dict)
                            and len(p["values"][code]) > 100,
                            f"IMF {code}"), headers=CURL, gap=3)
    vals = load_json(blob)["values"][code][ISO3]
    return {int(y): float(v) for y, v in vals.items() if v is not None}


def write_imf(rel: str, comment: str, wanted: list[tuple[str, str, int]],
              s: dict) -> None:
    header = ["year"]
    for _, col, _ in wanted:
        header += [col, f"{col}_proj"]
    years = sorted(set().union(*(s[c] for c, _, _ in wanted)))
    rows = []
    for y in years:
        row = [y]
        for code, _, nd in wanted:
            v = s[code].get(y)
            row += [fmt(v, nd), ""] if y <= LAST_FULL_YEAR else ["", fmt(v, nd)]
        rows.append(row)
    write(rel, comment, header, rows)


def imf_weo(wb_s: dict) -> None:
    codes = ["GGXWDG_NGDP", "GGXCNL_NGDP", "NGDPD", "NGDP_RPCH", "PCPIPCH",
             "BCA_NGDPD", "LUR"]
    s = {c: imf(c) for c in codes}
    near(s["GGXWDG_NGDP"].get(2024), 31.2, 0.15, "IMF GG debt 2024 (% GDP)")
    near(s["GGXWDG_NGDP"].get(2025), 30.5, 0.15, "IMF GG debt 2025 (% GDP)")
    near(s["GGXCNL_NGDP"].get(2020), -2.9, 0.15, "IMF net lending 2020")
    near(s["NGDPD"].get(2009), 129.0, 0.5, "IMF NGDPD 2009 (US$bn, back-cast)")
    near(s["NGDPD"].get(2024), 459.4, 1.0, "IMF NGDPD 2024 (US$bn)")
    # Second-source magnitude: IMF and WB levels agree within 5% after the
    # seam, and the IMF's 2009->2010 step is ordinary growth, not +39%.
    wbg = wb_s["NY.GDP.MKTP.CD"]
    for y in (2015, 2020, 2024):
        gap = s["NGDPD"][y] * 1e9 / wbg[y] - 1
        if abs(gap) > 0.05:
            fail(f"IMF vs WB GDP {y}: {gap:+.1%} apart (expected < 5%)")
    print(f"  ok  IMF NGDPD within 5% of WB GDP in 2015/2020/2024 "
          f"(2024 {s['NGDPD'][2024] * 1e9 / wbg[2024] - 1:+.1%})")
    step = s["NGDPD"][2010] / s["NGDPD"][2009] - 1
    if step > 0.2:
        fail(f"IMF NGDPD 2009->2010 = {step:+.0%}: no longer back-cast?")
    head = ("Source: IMF, World Economic Outlook via the DataMapper API. "
            f"{OFFICIAL_FALSE} Calendar years. Years after {LAST_FULL_YEAR} "
            "are WEO PROJECTIONS and sit only in the *_proj columns; an "
            "observed column and its _proj twin never share a row.")
    write_imf("fiscal/debt_balance_from_imf.csv", f"""
{head}
GGXWDG_NGDP: general-government gross debt, % of GDP. GGXCNL_NGDP: general-government net lending (+) / borrowing (-), % of GDP.
General government on the IMF's basis; NOT the MOF's public-debt definition and NOT the NSO state budget. Three concepts, three files, never spliced.
""", [("GGXWDG_NGDP", "gross_debt_pct_gdp", 1),
      ("GGXCNL_NGDP", "net_lending_pct_gdp", 1)], s)
    write_imf("international/gdp_usd_from_imf.csv", f"""
{head}
NGDPD: GDP, current US$ BILLIONS, back-cast by the IMF onto the revised (post-2010) scale, so continuous across 2010 (2009->2010 {step:+.1%}). It differs from the World Bank level by up to ~4% after 2010 (conversion); each is labelled with its own name.
NGDP_RPCH: real GDP growth, %. PCPIPCH: CPI inflation, annual average, %. BCA_NGDPD: current account, % of GDP. LUR: unemployment rate, %.
""", [("NGDPD", "gdp_usd_bn", 3), ("NGDP_RPCH", "real_growth_pct", 2),
      ("PCPIPCH", "inflation_pct", 2), ("BCA_NGDPD", "current_account_pct_gdp", 2),
      ("LUR", "unemployment_pct", 2)], s)


# ---------------------------------------------------------------- UNDP

def undp() -> None:
    def check(blob: bytes) -> None:
        if not blob.startswith(b"iso3,") or b"VNM" not in blob:
            fail("HDR CSV: not the composite-indices table")
    blob = fetch(HDR, "hdr25_series.csv", check)
    # The file is LATIN-1; decoding it as UTF-8 trips on a 0xf4 byte.
    reader = csv.DictReader(io.StringIO(blob.decode("latin-1")))
    row = next((r for r in reader if r.get("iso3") == ISO3), None)
    if row is None:
        fail("HDR CSV: no VNM row")
    series: dict[str, dict[int, float]] = {"hdi": {}, "le": {}, "gii": {}}
    rank = {}
    for key, value in row.items():
        if not key or value in (None, "", ".."):
            continue
        m = re.fullmatch(r"([a-z_]+)_(\d{4})", key.lower())
        if not m:
            continue
        name, year = m.group(1), int(m.group(2))
        if name in series:              # exact match: hdi_rank_* is not hdi
            series[name][year] = float(value)
        elif name in ("hdi_rank", "rank_hdi"):
            rank[year] = float(value)
    near(series["hdi"].get(2023), 0.766, 0.0015, "UNDP HDI 2023")
    near(series["hdi"].get(2006), 0.656, 0.0015, "UNDP HDI 2006")
    near(rank.get(2023), 93, 0.5, "UNDP HDI rank 2023")
    near(series["le"].get(2023), 74.6, 0.1, "UNDP life expectancy 2023")
    near(series["gii"].get(2023), 0.299, 0.002, "UNDP GII 2023")
    years = sorted(y for y in set().union(*series.values())
                   if YEAR_FROM <= y <= YEAR_TO)
    write("international/hdi_from_undp.csv", f"""
Source: UNDP, Human Development Report 2025, composite indices complete time series (one vintage throughout). {OFFICIAL_FALSE}
hdi: Human Development Index (0-1). hdi_rank: rank among countries, published for the latest year only. life_expectancy: years at birth (UNDP's input, from UN WPP). gii: Gender Inequality Index (0-1, lower is more equal).
The NSO computes its own HDI (V14.01: 2023 0.746); a different computation, never merged with this one. The file is Latin-1.
""", ["year", "hdi", "hdi_rank", "life_expectancy", "gii"],
          [[y, fmt(series["hdi"].get(y), 3), fmt(rank.get(y), 0),
            fmt(series["le"].get(y), 1), fmt(series["gii"].get(y), 3)]
           for y in years])


def ti_cpi() -> None:
    """Stage Vietnam's comparable CPI from TI's pinned 2025 workbook."""
    def parse(blob: bytes) -> list[dict[str, str]]:
        z = zipfile.ZipFile(io.BytesIO(blob))
        book = z.read("xl/workbook.xml").decode()
        rels = z.read("xl/_rels/workbook.xml.rels").decode()
        sheets = dict(re.findall(
            r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', book))
        targets = dict(re.findall(
            r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
        path = "xl/" + targets[sheets["CPI Historical"]].lstrip("/") \
            .removeprefix("xl/")
        shared = z.read("xl/sharedStrings.xml").decode()
        strings = ["".join(re.findall(r"<t[^>]*>(.*?)</t>", cell, re.S))
                   for cell in re.findall(r"<si>(.*?)</si>", shared, re.S)]
        rows = []
        for row in re.findall(r"<row[^>]*>(.*?)</row>",
                              z.read(path).decode(), re.S):
            cells = {}
            for match in re.finditer(
                    r'<c r="([A-Z]+)\d+"([^>]*?)(?:/>|>(.*?)</c>)',
                    row, re.S):
                ref, attrs, raw = match.groups()
                value = re.search(r"<v>(.*?)</v>", raw or "")
                text = value.group(1) if value else ""
                cells[ref] = strings[int(text)] if 't="s"' in attrs and text \
                    else text
            rows.append(cells)
        return rows

    def check_ti(blob: bytes) -> None:
        if blob[:2] != b"PK" or not parse(blob):
            fail("TI CPI: not the historical workbook")

    rows = parse(fetch(TI, "ti_CPI2025.xlsx", check_ti))
    head = next(row for row in rows
                if (row.get("A") or "").startswith("Country"))
    columns = {value: ref for ref, value in head.items() if value}
    scores = {int(float(row[columns["Year"]])):
              float(row[columns["CPI score"]])
              for row in rows if row.get(columns["ISO3"]) == ISO3
              and row.get(columns["CPI score"])}
    if sorted(scores) != list(range(2012, 2026)):
        fail(f"TI CPI years incomplete: {sorted(scores)}")
    near(scores.get(2025), 41, 0, "TI CPI 2025")
    write("international/cpi_from_ti.csv", """
Source: Transparency International, CPI 2025 workbook, CPI Historical sheet. official: false.
Comparable 2012-method series; scores are 0-100 and higher means perceived cleaner.
""", ["year", "cpi_score"],
          [[year, fmt(scores[year], 0)] for year in sorted(scores)])


# ---------------------------------------------------------------- UN DESA

def is_xlsx(blob: bytes) -> None:
    # An HTML error page under HTTP 200 is not a zip; a truncated one fails.
    if blob[:4] != b"PK\x03\x04":
        fail("UN DESA: payload is not an xlsx (HTML error page?)")
    bad = zipfile.ZipFile(io.BytesIO(blob)).testzip()
    if bad:
        fail(f"UN DESA: xlsx member {bad} is corrupt")


def undesa() -> None:
    import openpyxl
    blob = fetch(DESA, "undesa_ims_2024.xlsx", is_xlsx)
    wb_ = openpyxl.load_workbook(io.BytesIO(blob), read_only=True)
    if "Table 1" not in wb_.sheetnames:
        fail("UN DESA: no Table 1")
    grid = list(wb_["Table 1"].iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(grid) if r and r[0] == "Index"), None)
    if hi is None:
        fail("UN DESA: no header row starting 'Index'")
    head = grid[hi]
    if not (str(head[4]).startswith("Location code of destination")
            and str(head[6]).startswith("Location code of origin")):
        fail(f"UN DESA: header moved: {head[:7]}")
    # Both-sexes block: the first run of year columns after column 6.
    ycols = []
    for j in range(7, len(head)):
        if isinstance(head[j], int):
            if ycols and head[j] <= head[ycols[-1]]:
                break
            ycols.append(j)
    years = [head[j] for j in ycols]
    if years[0] != 1990 or years[-1] != 2024:
        fail(f"UN DESA: year columns {years}")
    rows = grid[hi + 1:]

    def pick(dest, orig):
        hit = [r for r in rows if r[4] == dest and r[6] == orig]
        if len(hit) != 1:
            fail(f"UN DESA: {len(hit)} rows for dest {dest} origin {orig}")
        return hit[0]

    em, im = pick(900, M49), pick(M49, 900)
    emig = dict(zip(years, (em[j] for j in ycols)))
    immig = dict(zip(years, (im[j] for j in ycols)))
    near(emig[2024], 3693537, 0, "UN DESA emigrants born in Vietnam 2024")
    near(emig[2010], 2339779, 0, "UN DESA emigrants 2010")
    near(immig[2024], 326418, 0, "UN DESA immigrants in Vietnam 2024")
    im_type = nfc(str(im[3] or ""))
    if im_type != "C R":
        fail(f"UN DESA: Vietnam's data type is {im_type!r}, expected 'C R'")
    # Destinations: country rows only (codes < 900 carry a data type;
    # region/income aggregates do not, or use codes >= 900).
    dests = [r for r in rows if r[6] == M49 and r[4] is not None
             and r[4] < 900 and r[3]]
    dests.sort(key=lambda r: -(r[ycols[-1]] or 0))
    top = dests[:6]
    # Sanity: destinations should not exceed the World total.
    if sum(r[ycols[-1]] or 0 for r in dests) > emig[2024] * 1.001:
        fail("UN DESA: destination rows sum above the World total")
    slug = lambda s: re.sub(r"[^a-z0-9]+", "_",
                            nfc(s).rstrip("*").lower()).strip("_")
    cols = [f"emigrants_in_{slug(r[1])}" for r in top]
    note = "; ".join(f"{nfc(r[1]).rstrip('*')} ({r[4]}): basis {r[3]}"
                     for r in top)
    out = []
    for k, y in enumerate(years):
        out.append([y, fmt(emig[y], 0), fmt(immig[y], 0)]
                   + [fmt(r[ycols[k]], 0) if r[ycols[k]] is not None else ""
                      for r in top])
    write("migration/migrant_stock_from_undesa.csv", f"""
Source: UN DESA Population Division, International Migrant Stock 2024 (POP/DB/MIG/Stock/Rev.2024), Table 1, both sexes, persons at mid-year. {OFFICIAL_FALSE}
emigrants_world: people from Viet Nam (origin 704) living anywhere abroad (destination World 900). immigrants_in_vietnam: destination 704, origin World.
BASIS for Vietnam's own (immigrant) row: data type '{im_type}' = CITIZENSHIP-based estimate INCLUDING REFUGEES, not country of birth. Each destination's figure is on that destination's basis (B = birth, C = citizenship, R = refugees added): {note}.
Top {len(top)} destinations by 2024 stock, ranked once in the fetcher. Five-yearly points plus 2024: draw as points, never interpolate into an annual line.
""", ["year", "emigrants_world", "immigrants_in_vietnam"] + cols, out)
    print("  top destinations 2024: " + ", ".join(
        f"{nfc(r[1]).rstrip('*')} {r[ycols[-1]]:,}" for r in top))


# ---------------------------------------------------------------- UN WPP

def wpp(wb_s: dict) -> None:
    idx = fetch(WPP_INDEX, "wpp_downloads.json",
                check_json(lambda p: isinstance(p, dict) and "Folders" in p,
                           "WPP downloads.json"))
    paths = sorted(set(re.findall(r'"Path"\s*:\s*"([^"]+)"',
                                  idx.decode("utf-8-sig"))))
    hits = [p for p in paths if WPP_FILE.search(p)]
    if len(hits) != 1:
        fail(f"WPP: {len(hits)} candidates for the medium indicators CSV")
    edition = int(WPP_FILE.search(hits[0]).group(1))
    url = WPP_SITE + urllib.parse.quote(hits[0])

    def check(blob: bytes) -> None:
        if blob[:2] != b"\x1f\x8b":
            fail("WPP: payload is not gzip")
        with gzip.open(io.BytesIO(blob)) as g:
            first = g.readline().decode("utf-8-sig")
        if not first.startswith("SortOrder,LocID"):
            fail(f"WPP: unexpected header {first[:60]!r}")
    blob = fetch(url, f"wpp{edition}_demographic_indicators_medium.csv.gz",
                 check, gap=2)
    # WPP 2024: estimates 1950-2023, medium projections 2024-2100.
    if edition != 2024:
        fail(f"WPP edition {edition}: re-pin the estimate/projection split")
    est_last = edition - 1
    rows = {}
    with gzip.open(io.BytesIO(blob), "rt", encoding="utf-8-sig") as g:
        for r in csv.DictReader(g):
            if r["LocID"] == str(M49):
                if r["Variant"] != "Medium":
                    fail(f"WPP: variant {r['Variant']}")
                rows[int(r["Time"])] = r
    if not rows:
        fail("WPP: no rows for LocID 704")
    loc = nfc(rows[2024]["Location"])
    if loc != "Viet Nam" or rows[2024]["ISO3_code"] != ISO3:
        fail(f"WPP: LocID 704 is {loc!r}")
    f = lambda y, c: float(rows[y][c]) if rows[y].get(c) not in (None, "") else None
    near(f(2024, "TPopulation1July"), 100987.7, 80, "WPP population 2024 (thousand)")
    near(f(2024, "TFR"), 1.90, 0.03, "WPP TFR 2024")
    # Second source: the World Bank's SP.POP.TOTL is WPP's July population.
    for y in (2010, 2024):
        gap = f(y, "TPopulation1July") * 1e3 / wb_s["SP.POP.TOTL"][y] - 1
        if abs(gap) > 0.005:
            fail(f"WPP vs WB population {y}: {gap:+.2%}")
    print("  ok  WPP July population = WB SP.POP.TOTL within 0.5% (2010, 2024)")
    cols = [("TPopulation1July", "population_thousand", 1),
            ("MedianAgePop", "median_age", 1), ("TFR", "tfr", 3),
            ("LEx", "life_expectancy", 2), ("Births", "births_thousand", 1),
            ("Deaths", "deaths_thousand", 1), ("CBR", "cbr_per_1000", 2),
            ("CDR", "cdr_per_1000", 2),
            ("NetMigrations", "net_migration_thousand", 3)]
    out = [[y, "estimate" if y <= est_last else "projection"]
           + [fmt(f(y, c), nd) for c, _, nd in cols]
           for y in sorted(rows)
           # 2101 carries only the 1 January population: no July row.
           if y >= YEAR_FROM and f(y, "TPopulation1July") is not None]
    write("demography/projection_from_unwpp.csv", f"""
Source: UN DESA Population Division, World Population Prospects {edition}, Demographic Indicators, MEDIUM variant (bulk CSV {hits[0].rsplit('/', 1)[-1]}, resolved from {WPP_INDEX}). {OFFICIAL_FALSE} Location 704 '{loc}'.
kind: 'estimate' for {YEAR_FROM}-{est_last}, 'projection' (medium variant) from {est_last + 1}. A projection is a scenario, not an outcome.
Units: population, births, deaths and net migration in THOUSANDS of persons (mid-year population at 1 July); median age in years; TFR births per woman; life expectancy at birth in years; CBR/CDR per 1,000.
The manifest's Data Portal API (dataportalapi/api/v1) answers 401 (Bearer token) to every data query since registration was required; this bulk file is the anonymous route to the same estimates.
The NSO's 2019-census projection (Dự báo dân số 2019–2069) is a PDF publication, not this. WPP net migration 2006 carries the same artefact as WB SM.POP.NETM.
""", ["year", "kind"] + [c for _, c, _ in cols], out)


# ---------------------------------------------------------------- WHO

def gho(code: str, flt: str) -> list[dict]:
    url = GHO.format(code=code) + urllib.parse.urlencode({"$filter": flt})
    blob = fetch(url, f"who_{code}.json",
                 check_json(lambda p: isinstance(p.get("value"), list)
                            and p["value"], f"WHO {code}"))
    rows = load_json(blob)["value"]
    for r in rows:
        if r["SpatialDim"] != ISO3 or r["IndicatorCode"] != code:
            fail(f"WHO {code}: row for {r['SpatialDim']}/{r['IndicatorCode']}")
    return rows


def who() -> None:
    rs = [r for r in gho("RS_198", f"SpatialDim eq '{ISO3}'")
          if r["Dim1"] == "SEX_BTSX"]
    rs = {r["TimeDim"]: r for r in rs}
    if len(rs) < 20:
        fail(f"WHO RS_198: only {len(rs)} both-sexes years")
    near(rs[2019]["NumericValue"], 23.5, 0.05, "WHO road deaths 2019 (/100k)")
    near(rs[2023]["NumericValue"], 12.9, 0.05, "WHO road deaths 2023 (/100k)")
    write("security/road_deaths_from_who.csv", f"""
Source: WHO Global Health Observatory, RS_198, estimated road traffic death rate per 100,000 population, both sexes (Dim1 SEX_BTSX), with the WHO uncertainty interval. {OFFICIAL_FALSE} A MODELLED ESTIMATE, not a count.
Before 2022 the model sits about three times above the police count (NSO V14.86); after the 2022 police reclassification the two converge. Two publishers, one idea: the gap is the finding, and the two are never spliced.
""", ["year", "road_deaths_per_100k", "low", "high"],
          [[y, fmt(rs[y]["NumericValue"], 1), fmt(rs[y]["Low"], 1),
            fmt(rs[y]["High"], 1)] for y in sorted(rs)])

    hm = gho("VIOLENCE_HOMICIDERATE",
             f"SpatialDim eq '{ISO3}' and Dim1 eq 'SEX_BTSX'")
    hm = {r["TimeDim"]: r for r in hm}
    near(hm[2021]["NumericValue"], 1.91, 0.01, "WHO homicide 2021 (/100k)")
    near(hm[2014]["NumericValue"], 1.79, 0.01, "WHO homicide 2014 (/100k)")
    write("security/homicide_from_who.csv", f"""
Source: WHO Global Health Observatory, VIOLENCE_HOMICIDERATE (Global Health Estimates), estimated homicide rate per 100,000, both sexes, with uncertainty interval. {OFFICIAL_FALSE}
A MODEL, not a count: Vietnam publishes no national homicide statistic. Never spliced onto the UNODC series (homicide_from_wb.csv, 2001-2011).
""", ["year", "homicide_per_100k", "low", "high"],
          [[y, fmt(hm[y]["NumericValue"], 2), fmt(hm[y]["Low"], 2),
            fmt(hm[y]["High"], 2)] for y in sorted(hm)])

    flt = (f"FLAG_LEVEL le 3 and DIM_COUNTRY_CODE eq '{ISO3}' and "
           "DIM_SEX_CODE eq 'TOTAL' and DIM_AGEGROUP_CODE eq 'TOTAL'")
    sel = ("DIM_YEAR_CODE,DIM_GHECAUSE_CODE,DIM_GHECAUSE_TITLE,FLAG_LEVEL,"
           "VAL_DTHS_COUNT_NUMERIC,VAL_DTHS_COUNT_LOW,VAL_DTHS_COUNT_HIGH,"
           "VAL_DTHS_RATE100K_NUMERIC,VAL_PROP_DTHS_PERCENT")
    url = GHE + urllib.parse.urlencode({"$filter": flt, "$select": sel})
    blob = fetch(url, "who_GHE_FULL_L0-3.json",
                 check_json(lambda p: isinstance(p.get("value"), list)
                            and len(p["value"]) > 1000, "WHO GHE_FULL"),
                 gap=3)
    payload = load_json(blob)
    if "@odata.nextLink" in payload:
        fail("WHO GHE_FULL: paged reply; follow nextLink")
    ghe = payload["value"]
    years = sorted({r["DIM_YEAR_CODE"] for r in ghe})
    per_year = {y: sum(1 for r in ghe if r["DIM_YEAR_CODE"] == y) for y in years}
    if len(set(per_year.values())) != 1:
        fail(f"WHO GHE: uneven causes per year {per_year}")
    l1 = {(r["DIM_YEAR_CODE"], r["DIM_GHECAUSE_CODE"]): r for r in ghe
          if r["FLAG_LEVEL"] == 1}
    near(l1[(2021, 10)]["VAL_DTHS_RATE100K_NUMERIC"], 145.3, 0.1,
         "WHO GHE 2021 communicable etc. (/100k)")
    near(l1[(2021, 600)]["VAL_DTHS_RATE100K_NUMERIC"], 506.7, 0.1,
         "WHO GHE 2021 noncommunicable (/100k)")
    near(l1[(2021, 1510)]["VAL_DTHS_RATE100K_NUMERIC"], 64.4, 0.1,
         "WHO GHE 2021 injuries (/100k)")
    # Internal consistency: the three level-1 groups sum to all causes.
    allc = {r["DIM_YEAR_CODE"]: r["VAL_DTHS_COUNT_NUMERIC"] for r in ghe
            if r["FLAG_LEVEL"] == 0}
    for y in years:
        s = sum(l1[(y, c)]["VAL_DTHS_COUNT_NUMERIC"] for c in (10, 600, 1510))
        if abs(s / allc[y] - 1) > 0.001:
            fail(f"WHO GHE {y}: level-1 groups sum {s:.0f} vs all {allc[y]:.0f}")
    print("  ok  WHO GHE level-1 groups sum to all causes, every year")
    ghe.sort(key=lambda r: (r["DIM_YEAR_CODE"], r["DIM_GHECAUSE_CODE"]))
    keys = [(r["DIM_YEAR_CODE"], nfc(r["DIM_GHECAUSE_TITLE"]),
             r["DIM_GHECAUSE_CODE"]) for r in ghe]
    if len(set(k[:2] for k in keys)) != len(keys):
        fail("WHO GHE: a cause title repeats within a year")
    out = [[r["DIM_YEAR_CODE"], r["FLAG_LEVEL"], r["DIM_GHECAUSE_CODE"],
            nfc(r["DIM_GHECAUSE_TITLE"]),
            fmt(r["VAL_DTHS_COUNT_NUMERIC"], 0),
            fmt(r["VAL_DTHS_COUNT_LOW"], 0), fmt(r["VAL_DTHS_COUNT_HIGH"], 0),
            fmt(r["VAL_DTHS_RATE100K_NUMERIC"], 2),
            fmt(r["VAL_PROP_DTHS_PERCENT"], 1)] for r in ghe]
    write("mortality/causes_from_who.csv", f"""
Source: WHO Global Health Estimates (xmart DEX_CMS/GHE_FULL), deaths by cause, Viet Nam, both sexes, all ages, {years[0]}-{years[-1]}. {OFFICIAL_FALSE}
MODELLED deaths, not registrations: the only cause-of-death source for Vietnam, and an estimate. The chapter says so in its title.
Long format. level: 0 all causes, 1 the three broad groups (sum to level 0), 2 cause groups, 3 causes; a level-n row is inside its level-(n-1) parent, so never sum across levels.
deaths, deaths_low, deaths_high: persons (uncertainty interval). rate_per_100k: deaths per 100,000 population. share_pct: % of all deaths.
""", ["year", "level", "cause_code", "cause", "deaths", "deaths_low",
      "deaths_high", "rate_per_100k", "share_pct"], out)


# ---------------------------------------------------------------- Comtrade

def comtrade() -> None:
    """Vietnam's own report to the UN, one period per call (preview cap)."""
    got = {}
    empty = []
    for y in range(YEAR_FROM, LAST_FULL_YEAR + 1):
        def check(blob: bytes, y=y) -> None:
            p = json.loads(blob)
            if p.get("error") or not isinstance(p.get("data"), list):
                fail(f"Comtrade {y}: {p.get('error') or 'no data list'}")
            if not p["data"]:
                fail(f"Comtrade {y}: count 0 (not loaded yet)")
        try:
            blob = fetch(COMTRADE.format(y=y), f"comtrade_{y}.json", check,
                         headers=BROWSER, gap=4)
        except RuntimeError as e:
            if "count 0" in str(e):
                empty.append(y)
                continue
            raise
        data = load_json(blob)["data"]
        flows = {}
        for r in data:
            if (r["reporterCode"] != M49 or r["partnerCode"] != 0
                    or r["cmdCode"] != "TOTAL" or r["customsCode"] != "C00"
                    or r["motCode"] != 0 or r["partner2Code"] != 0):
                continue
            if r["flowCode"] in flows:
                fail(f"Comtrade {y}: two {r['flowCode']} totals")
            flows[r["flowCode"]] = (r["primaryValue"], r["classificationCode"])
        got[y] = flows
    if empty and empty != list(range(min(empty), LAST_FULL_YEAR + 1)):
        fail(f"Comtrade: holes inside the series at {empty}")
    near(got[2023]["X"][0] / 1e9, 353.08, 0.05, "Comtrade exports 2023 (US$bn)")
    # Second source: the NSO's own 2023 exports (V09.01, 354,721 US$mn).
    gap = got[2023]["X"][0] / 354_721e6 - 1
    if abs(gap) > 0.02:
        fail(f"Comtrade vs NSO exports 2023: {gap:+.2%}")
    print(f"  ok  Comtrade exports 2023 within {abs(gap):.1%} of NSO V09.01")
    out = [[y, fmt(got[y].get("X", (None,))[0], 0),
            fmt(got[y].get("M", (None,))[0], 0),
            "/".join(sorted({c for _, c in got[y].values()}))]
           for y in sorted(got)]
    write("trade/trade_from_comtrade.csv", f"""
Source: UN Comtrade public preview API, Viet Nam (704) as REPORTER, partner World, HS TOTAL, goods only. {OFFICIAL_FALSE} Cross-check only.
Units: current US$. exports FOB; imports CIF (as reported). hs_edition: the HS classification Vietnam reported in that year (H2...H6).
{"Not yet loaded (count 0), left out, not zero: " + ", ".join(map(str, empty)) + "." if empty else ""}
The NSO's customs-based trade (V09.01) is the curated series; this is Vietnam's report to the UN of the same flows.
""", ["year", "exports_usd", "imports_usd", "hs_edition"], out)
    if empty:
        print(f"  Comtrade not loaded yet: {empty}")


# ---------------------------------------------------------------- main

def main() -> None:
    print("Comparables for Vietnam: World Bank, ADB, IMF, UNDP, TI, UN DESA, "
          "WPP, WHO, Comtrade")
    s = world_bank()
    imf_weo(s)
    undp()
    ti_cpi()
    undesa()
    wpp(s)
    who()
    comtrade()
    print("Done. Diff the sidecars before adopting.")


if __name__ == "__main__":
    main()
