#!/usr/bin/env python3
"""Stage Bank of Israel credit and market gaps as review sidecars.

Usage: python3 countries/israel/fetch/fetch_boi_gaps.py

The BoI topic pages can return a Radware loader to non-browser clients, but
their linked ``boi_files`` workbooks are stable downloads.  This fetcher uses
those exact publication files, validates their signatures and schemas before
caching them, and writes review sidecars only.  It never overwrites a curated
publication CSV.

Outputs:
  data/markets/policy_rate_from_boi.csv
  data/markets/government_bond_10y_from_boi.csv
  data/credit/household_debt_from_boi.csv
  data/credit/mortgage_rate_from_boi.csv
  data/credit/npl_ratio_from_boi.csv

``xlrd`` is needed for the BoI's genuine Excel-97 files.  ``openpyxl`` reads
the OOXML files.  Both formats are checked before a workbook reader sees them;
an HTTP-200 HTML challenge therefore cannot become an empty dataset.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import math
import sys
import time
import urllib.request
from collections import defaultdict
from pathlib import Path

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - refresh environment only
    raise SystemExit("openpyxl is required: pip install openpyxl") from exc
try:
    import xlrd
except ImportError as exc:  # pragma: no cover - refresh environment only
    raise SystemExit("xlrd is required for BoI .xls files: pip install xlrd") from exc


ROOT = Path(__file__).resolve().parent.parent
CACHE = ROOT / "cache" / "boi_gaps"
MARKETS = ROOT / "data" / "markets"
CREDIT = ROOT / "data" / "credit"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (X11; Linux x86_64) worldreports/1.0",
    "Accept": "application/vnd.ms-excel,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/json",
}

URLS = {
    "policy_events": "https://www.boi.org.il/boi_files/Statistics/bointcre_m.xls",
    "policy_history": "https://www.boi.org.il/boi_files/Statistics/bointcre_1.xls",
    "bond_10y": "https://www.boi.org.il/boi_files/Statistics/shcd05_e.xls",
    "household_debt": "https://www.boi.org.il/boi_files/Statistics/itrchove.xlsx",
    "mortgage_rate": "https://www.boi.org.il/boi_files/Pikuah/mashfix.xls",
    "fsi": "https://www.boi.org.il/boi_files/Statistics/FSI_ISR.xlsx",
    "gdp": "https://api.worldbank.org/v2/country/ISR/indicator/NY.GDP.MKTP.CN?format=json&per_page=100",
}

OLE = bytes.fromhex("D0CF11E0A1B11AE1")
OOXML = b"PK\x03\x04"


class FetchError(RuntimeError):
    pass


def validate_blob(raw: bytes, kind: str, url: str) -> None:
    """Reject HTML/error payloads and require the advertised file family."""
    probe = raw[:2048].lower()
    if b"<html" in probe or b"radware" in probe or b"loader page" in probe:
        raise FetchError(f"{url}: HTML/challenge payload, not {kind}")
    signature = OLE if kind == "xls" else OOXML
    if not raw.startswith(signature):
        raise FetchError(f"{url}: wrong {kind} signature {raw[:12]!r}")
    if len(raw) < 10_000:
        raise FetchError(f"{url}: implausibly short workbook ({len(raw)} bytes)")


def download(name: str, kind: str) -> bytes:
    CACHE.mkdir(parents=True, exist_ok=True)
    suffix = ".xls" if kind == "xls" else ".xlsx"
    target = CACHE / f"{name}{suffix}"
    if target.exists():
        raw = target.read_bytes()
        try:
            validate_blob(raw, kind, URLS[name])
            return raw
        except FetchError:
            target.unlink()

    last: Exception | None = None
    for attempt in range(4):
        try:
            request = urllib.request.Request(URLS[name], headers=HEADERS)
            with urllib.request.urlopen(request, timeout=90) as response:
                raw = response.read()
                content_type = response.headers.get("Content-Type", "").lower()
            if "html" in content_type:
                raise FetchError(f"HTML content type {content_type!r}")
            validate_blob(raw, kind, URLS[name])
            target.write_bytes(raw)
            return raw
        except Exception as exc:  # noqa: BLE001 - retry remote failures
            last = exc
            print(f"  {name} attempt {attempt + 1}: {exc}", file=sys.stderr)
            time.sleep(2 * (attempt + 1))
    raise FetchError(f"{URLS[name]}: {last}")


def download_json(name: str) -> object:
    request = urllib.request.Request(URLS[name], headers=HEADERS)
    with urllib.request.urlopen(request, timeout=90) as response:
        raw = response.read()
        content_type = response.headers.get("Content-Type", "").lower()
    if "html" in content_type or b"<html" in raw[:1024].lower():
        raise FetchError(f"{URLS[name]}: HTML instead of JSON")
    try:
        return json.loads(raw)
    except json.JSONDecodeError as exc:
        raise FetchError(f"{URLS[name]}: invalid JSON") from exc


def xls_rows(raw: bytes, sheet: int | str = 0) -> tuple[xlrd.book.Book, list[list[object]]]:
    book = xlrd.open_workbook(file_contents=raw)
    ws = book.sheet_by_index(sheet) if isinstance(sheet, int) else book.sheet_by_name(sheet)
    return book, [ws.row_values(i) for i in range(ws.nrows)]


def xldate(book: xlrd.book.Book, value: object) -> dt.date | None:
    if isinstance(value, (int, float)):
        try:
            return xlrd.xldate_as_datetime(value, book.datemode).date()
        except (ValueError, OverflowError):
            return None
    if isinstance(value, str):
        text = value.strip()
        for pattern in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return dt.datetime.strptime(text, pattern).date()
            except ValueError:
                pass
    return None


def number(value: object) -> float | None:
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    return None


def write_sidecar(path: Path, header: list[str], rows: list[list[object]], notes: list[str]) -> None:
    """Write a review file, refusing to replace a differently-shaped file."""
    path.parent.mkdir(parents=True, exist_ok=True)
    expected = ",".join(header)
    if path.exists():
        first = next((line for line in path.read_text().splitlines()
                      if line and not line.startswith("#")), "")
        if first != expected:
            raise FetchError(f"refusing to overwrite {path.name}: header is {first!r}")
    with path.open("w", newline="") as handle:
        for note in notes:
            handle.write(f"# {note}\n")
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(header)
        writer.writerows(rows)
    print(f"  {path.relative_to(ROOT)}: {len(rows)} rows")


def policy_rate() -> None:
    """Annual calendar-year mean of the declared nominal policy rate.

    For 2006--16 the BoI history file publishes monthly averages for the
    liquidity period; twelve observations are averaged.  From 2017 the event
    file supplies the effective start date, so the annual value is weighted by
    the exact number of calendar days each declared rate was in force.
    """
    old_book, old_rows = xls_rows(download("policy_history", "xls"))
    monthly: dict[int, dict[int, float]] = defaultdict(dict)
    for row in old_rows:
        if len(row) < 2:
            continue
        day, value = xldate(old_book, row[0]), number(row[1])
        if day and value is not None and 2006 <= day.year <= 2016:
            if day.month in monthly[day.year]:
                raise FetchError(f"duplicate policy month {day:%Y-%m}")
            monthly[day.year][day.month] = value
    for year in range(2006, 2017):
        if len(monthly[year]) != 12:
            raise FetchError(f"policy history {year}: expected 12 months, got {len(monthly[year])}")

    event_book, event_rows = xls_rows(download("policy_events", "xls"))
    events: list[tuple[dt.date, float]] = []
    for row in event_rows:
        if len(row) < 4:
            continue
        start, value = xldate(event_book, row[1]), number(row[3])
        if start and value is not None:
            events.append((start, value))
    events.sort()
    if not events or events[0][0].year != 2017 or events[0][1] != 0.1:
        raise FetchError("policy event history no longer begins with 0.1% in 2017")

    result: list[list[object]] = []
    for year in range(2006, 2017):
        result.append([year, f"{sum(monthly[year].values()) / 12:.6f}",
                       "mean_of_12_monthly_liquidity_period_averages"])
    for year in range(2017, 2026):
        start, end = dt.date(year, 1, 1), dt.date(year + 1, 1, 1)
        prior = [(day, rate) for day, rate in events if day <= start]
        rate = prior[-1][1] if prior else events[0][1]
        cursor, weighted = start, 0.0
        for day, new_rate in events:
            if day <= start:
                continue
            if day >= end:
                break
            weighted += (day - cursor).days * rate
            cursor, rate = day, new_rate
        weighted += (end - cursor).days * rate
        result.append([year, f"{weighted / (end - start).days:.6f}",
                       "calendar_day_weighted_declared_nominal_rate"])

    known = float(result[-1][1])
    if abs(known - 4.476027) > 0.00001:
        raise FetchError(f"policy known check failed: 2025={known}")
    write_sidecar(
        MARKETS / "policy_rate_from_boi.csv",
        ["year", "policy_rate_pct", "annual_aggregation"], result,
        [URLS["policy_history"], URLS["policy_events"],
         "Percent, nominal. Definition/aggregation seam is explicit in annual_aggregation."],
    )


def bond_yield() -> None:
    book, rows = xls_rows(download("bond_10y", "xls"))
    if not any("longest government bond" in str(cell).lower()
               for row in rows for cell in row):
        raise FetchError("10-year workbook footnote/definition missing")
    monthly: dict[int, dict[int, float]] = defaultdict(dict)
    for row in rows:
        if len(row) < 3:
            continue
        day, value = xldate(book, row[1]), number(row[2])
        if day and value is not None and 2006 <= day.year <= 2025:
            if day.month in monthly[day.year]:
                raise FetchError(f"duplicate 10-year yield month {day:%Y-%m}")
            monthly[day.year][day.month] = value
    result: list[list[object]] = []
    for year in range(2006, 2026):
        values = list(monthly[year].values())
        if len(values) != 12:
            raise FetchError(f"10-year yield {year}: expected 12 monthly means, got {len(values)}")
        result.append([year, f"{sum(values) / 12:.6f}", 12])
    if abs(float(result[-1][1]) - 4.244167) > 0.00001:
        raise FetchError(f"10-year yield known check failed: 2025={result[-1][1]}")
    write_sidecar(
        MARKETS / "government_bond_10y_from_boi.csv",
        ["year", "government_bond_10y_yield_pct", "monthly_observations"], result,
        [URLS["bond_10y"],
         "Annual arithmetic mean of 12 BoI monthly observations; percent.",
         "Publisher footnote: longest government bond with maturity between 8 and 11 years."],
    )


def household_debt() -> None:
    raw = download("household_debt", "xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet_name = "Credit Outstanding - Quarterly"
    if sheet_name not in wb.sheetnames:
        raise FetchError(f"household debt sheet missing: {wb.sheetnames}")
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    dates = rows[7][2:]
    target = next((row for row in rows if row[0] == "CRA_OUT_0204"), None)
    if target is None or target[1] != "Total debt to all lenders":
        raise FetchError("CRA_OUT_0204 household total row missing or relabelled")
    debt: dict[int, float] = {}
    for day, value in zip(dates, target[2:]):
        if isinstance(day, dt.datetime) and day.month == 12 and number(value) is not None:
            debt[day.year] = float(value)

    payload = download_json("gdp")
    if not (isinstance(payload, list) and len(payload) == 2 and isinstance(payload[1], list)):
        raise FetchError("World Bank GDP response shape changed")
    gdp: dict[int, float] = {}
    for item in payload[1]:
        if item.get("indicator", {}).get("id") != "NY.GDP.MKTP.CN":
            raise FetchError("World Bank response contains wrong indicator")
        if item.get("countryiso3code") != "ISR":
            raise FetchError("World Bank response contains wrong geography")
        if item.get("value") is not None:
            gdp[int(item["date"])] = float(item["value"]) / 1_000_000_000

    result: list[list[object]] = []
    for year in range(2006, 2026):
        if year not in debt or year not in gdp:
            raise FetchError(f"household debt/GDP missing {year}")
        ratio = debt[year] / gdp[year] * 100
        result.append([year, f"{debt[year]:.6f}", f"{gdp[year]:.6f}", f"{ratio:.6f}"])
    if abs(float(result[-2][1]) - 845.849452) > 0.001:
        raise FetchError(f"household debt known check failed: 2024={result[-2][1]}")
    if abs(float(result[-2][3]) - 42.159657) > 0.001:
        raise FetchError(f"household debt/GDP known check failed: 2024={result[-2][3]}")
    write_sidecar(
        CREDIT / "household_debt_from_boi.csv",
        ["year", "household_debt_nis_bn", "nominal_gdp_nis_bn", "household_debt_pct_gdp"], result,
        [URLS["household_debt"], URLS["gdp"],
         "BoI CRA_OUT_0204 total household debt to all lenders, Q4/end-year, current NIS billion.",
         "Ratio = BoI Q4 stock / World Bank WDI annual nominal GDP (NY.GDP.MKTP.CN)."],
    )


def mortgage_rate() -> None:
    book, rows = xls_rows(download("mortgage_rate", "xls"), "הלוואות שניתנו מ-01.03.2013")
    title = "Average Rate of Interest on Unindexed Mortgages"
    if len(rows) <= 7 or not rows[1] or rows[1][0] != title:
        raise FetchError("mortgage workbook title moved or changed")

    def header_cell(label: str) -> tuple[int, int]:
        hits = [(row_index, column_index)
                for row_index, row in enumerate(rows[:10])
                for column_index, cell in enumerate(row)
                if str(cell).strip() == label]
        if len(hits) != 1:
            raise FetchError(f"mortgage header {label!r}: expected one cell, got {hits}")
        return hits[0]

    # These are deliberately asserted as well as discovered. If BoI inserts a
    # column or changes the multi-row header, do not silently read a maturity
    # bucket as the overall rate or an effective date as the reporting month.
    average_header = header_cell("Average")
    date_header = header_cell("month of return")
    if average_header != (4, 0) or date_header != (6, 9):
        raise FetchError("mortgage Average/reporting-month columns moved: "
                         f"{average_header=}, {date_header=}")
    average_column = average_header[1]
    date_column = date_header[1]
    first_data_row = max(average_header[0], date_header[0]) + 1

    monthly: dict[int, dict[int, float]] = defaultdict(dict)
    for row in rows[first_data_row:]:
        if len(row) <= max(average_column, date_column):
            continue
        day = xldate(book, row[date_column])
        value = number(row[average_column])
        if day and value is not None and 2014 <= day.year <= 2025:
            if day.month in monthly[day.year]:
                raise FetchError(f"duplicate mortgage-rate month {day:%Y-%m}")
            monthly[day.year][day.month] = value
    result: list[list[object]] = []
    for year in range(2014, 2026):
        values = list(monthly[year].values())
        if len(values) != 12:
            raise FetchError(f"mortgage rate {year}: expected 12 months, got {len(values)}")
        result.append([year, f"{sum(values) / 12:.6f}", 12])
    if abs(float(result[-1][1]) - 5.054880) > 0.00001:
        raise FetchError(f"mortgage rate known check failed: 2025={result[-1][1]}")
    write_sidecar(
        CREDIT / "mortgage_rate_from_boi.csv",
        ["year", "new_unindexed_mortgage_rate_pct", "monthly_observations"], result,
        [URLS["mortgage_rate"],
         "Annual arithmetic mean of 12 monthly weighted-average effective rates.",
         "New unindexed-shekel housing loans, fixed and variable, for dwelling purchase or secured on a dwelling."],
    )


def npl_ratio() -> None:
    raw = download("fsi", "xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    if "FSD" not in wb.sheetnames:
        raise FetchError(f"FSI sheet missing: {wb.sheetnames}")
    rows = list(wb["FSD"].iter_rows(values_only=True))
    periods = rows[8][5:]
    target = next((row for row in rows if row[1] == "FSANL"), None)
    if target is None or "Nonperforming loans to total gross loans" not in str(target[0]):
        raise FetchError("FSANL NPL ratio row missing or relabelled")
    q4: dict[int, float] = {}
    for period, value in zip(periods, target[5:]):
        if isinstance(period, str) and period.endswith("Q4") and number(value) is not None:
            q4[int(period[:4])] = float(value)
    result: list[list[object]] = []
    for year in range(2012, 2026):
        if year not in q4:
            raise FetchError(f"FSANL Q4 missing {year}")
        vintage = "IMF_FSI_2006_guide" if year < 2022 else "IMF_FSI_2019_guide_transition_2022"
        result.append([year, f"{q4[year]:.6f}", "Q4", vintage])
    if abs(float(result[-1][1]) - 0.638792) > 0.00001:
        raise FetchError(f"NPL known check failed: 2025={result[-1][1]}")
    write_sidecar(
        CREDIT / "npl_ratio_from_boi.csv",
        ["year", "npl_to_total_gross_loans_pct", "observation", "method_vintage"], result,
        [URLS["fsi"],
         "BoI report to IMF FSI, code FSANL; Q4/end-year percent.",
         "BoI warns of a November 2022 transition from the IMF 2006 to 2019 FSI guide; do not hide this seam."],
    )


def main() -> None:
    print("Bank of Israel credit and markets gap sidecars")
    policy_rate()
    bond_yield()
    household_debt()
    mortgage_rate()
    npl_ratio()
    print("Done. Review diffs and definitions before promoting any sidecar.")


if __name__ == "__main__":
    main()
