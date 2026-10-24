#!/usr/bin/env python3
"""Recorded crime from the Office of the Prosecutor General of Ukraine (OGP).

Usage:  python3 -I fetch/fetch_gp.py   (run from the country directory)

Output (review sidecar, never curated data):
  data/security/security_gp.csv

Source: the OGP "Єдиний звіт про кримінальні правопорушення" (Unified report
on criminal offences, Form № 1-monthly, approved by OGP Order № 299 of
30.06.2020). The statistics index on new.gp.gov.ua links the current-year
post, which links the 2013-2024 archive post; both carry monthly cumulative
xlsx attachments served by old.gp.gov.ua's file downloader as
application/octet-stream. The fetcher resolves the attachment URLs from the
post pages on every run (file_ids are never hardcoded), keeps only the
January-December (full-year) edition of each year, and caches the workbooks
under cache/gp/.

What this is and is not (SOURCING_A_COUNTRY.md §4a):
  - These are PROCEEDINGS recorded in the Unified Register of Pre-trial
    Investigations and counted by the prosecutor's offices — a
    police/prosecutor-recorded offence count, NOT the vital-statistics
    homicide count (death certificates) and NOT a civilian war-casualty
    count. Never splice with the UNODC line in security.csv.
  - From 2022 the Article 115 row mixes conflict deaths into recorded
    crime: deaths from Russian strikes are opened as proceedings under
    Art. 115 (usually with the Art. 438 war-crime qualifier). The OGP's
    own sub-row "учинених у ході бойових дій" (committed in the course of
    hostilities) quantifies part of that admixture from 2023 only; it is
    exported as a separate column, blank where the publisher does not
    report it (never zero-filled).

Verified 2026-10-24 against the publisher's workbooks and the media
headline that republishes them (skilky-skilky.info, 19.03.2026): the
Art. 115 column 2020-2025 (3809, 3230, 22083, 36667, 52197, 59497) and
the hostilities sub-row 2023-2025 (14991, 15113, 10090) match byte-for-
byte; the 2022 aggregate (362,636) matches the figure quoted in academic
literature citing the OGP.
"""
from __future__ import annotations

import html as html_mod
import re
import sys
import urllib.request
import zipfile
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("fetch_gp: openpyxl is required to read the OGP xlsx "
             "workbooks (pip install openpyxl); refusing to skip the "
             "parse silently.")

ROOT = Path(__file__).resolve().parent.parent
CACHE = ROOT / "cache" / "gp"
DATA = ROOT / "data" / "security"
CACHE.mkdir(parents=True, exist_ok=True)
DATA.mkdir(parents=True, exist_ok=True)

UA = ("Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, "
      "like Gecko) Chrome/126.0.0.0 Safari/537.36")
INDEX_URL = "https://new.gp.gov.ua/ua/posts/statistika"
CURRENT_SLUG = ("pro-zareyestrovani-kriminalni-pravoporushennya"
                "-ta-rezultati-yih-dosudovogo-rozsliduvannya")
ARCHIVE_SLUG = "2013-2024-roki-" + CURRENT_SLUG

ATTACH_RE = re.compile(
    r'<a[^>]*href="([^"]*(?:file_downloader|stst2011)\.html[^"]*)"[^>]*>'
    r"(.*?)</a>", re.S)
FULL_YEAR_RE = re.compile(r"січень\s*[-–]\s*грудень\s+(\d{4})")
ART115_RE = re.compile(r"^Умисне вбивство,\s*ст\.\s*115$")
HOSTILITIES_RE = re.compile(r"(учинених|вчинених) у ході бойових дій")


def get(url: str, binary: bool = False):
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=120) as r:
        blob = r.read()
    if not binary:
        text = blob.decode("utf-8", errors="replace")
        # Official gateways answer block/challenge pages with HTTP 200;
        # refuse to parse those as content.
        if "Just a moment" in text or "The request is blocked" in text:
            sys.exit(f"fetch_gp: block/challenge page from {url}")
        return text
    return blob


def post_url(page: str, slug: str) -> str:
    m = re.search(r'href="(https://new\.gp\.gov\.ua/ua/posts/'
                  + re.escape(slug) + r'[^"]*)"', page)
    if not m:
        sys.exit(f"fetch_gp: post '{slug}' not linked from the page that "
                 "should carry it — the OGP statistics section changed; "
                 "resolve the new URL by hand.")
    return m.group(1)


def full_year_attachments(*pages: str) -> dict[int, str]:
    """Map year -> download URL of the January-December workbook."""
    found: dict[int, str] = {}
    for page in pages:
        for url, text in ATTACH_RE.findall(page):
            label = html_mod.unescape(re.sub(r"<[^>]+>|\s+", " ", text))
            m = FULL_YEAR_RE.search(label)
            if not m:
                continue
            year = int(m.group(1))
            url = html_mod.unescape(url)
            if year in found and found[year] != url:
                sys.exit(f"fetch_gp: two different full-year attachments "
                         f"for {year}; review by hand.")
            found[year] = url
    return found


def download(year: int, url: str) -> Path:
    path = CACHE / f"gp_unified_report_{year}.xlsx"
    if not path.exists():
        print(f"  downloading {year}: {url}")
        blob = get(url, binary=True)
        # Served as application/octet-stream; the zip signature is the
        # contract, not the content type.
        if not zipfile.is_zipfile(__import__("io").BytesIO(blob)):
            sys.exit(f"fetch_gp: {year} attachment is not an xlsx zip "
                     f"({len(blob)} bytes from {url}) — refusing to cache "
                     "an error page.")
        path.write_bytes(blob)
    return path


def num(cell) -> int | None:
    if isinstance(cell, (int, float)):
        return int(cell)
    if isinstance(cell, str) and cell.strip().replace(" ", "").isdigit():
        return int(cell.replace(" ", ""))
    return None


def value_column(rows: list, upto: int) -> int:
    """Column of the '1' header cell in the nearest 'А/Б/1/2' header row
    at or above row index `upto`. Each table repeats its own header, so
    the column is located per table, never assumed."""
    for i in range(upto, -1, -1):
        cells = [str(c).strip() if c is not None else "" for c in rows[i]]
        for j in range(len(cells) - 1):
            if cells[j] == "Б" and cells[j + 1] == "1":
                return j + 1
    sys.exit("fetch_gp: no 'Б'/'1' column header found — workbook layout "
             "changed; re-pin the table structure by hand.")


def parse(path: Path, year: int) -> tuple[int, int, int | None]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # Cover sheet must confirm this is the January-December edition of
    # the expected year.
    cover = " ".join(
        str(c) for row in wb["0"].iter_rows(values_only=True)
        for c in row if isinstance(c, str))
    if not (FULL_YEAR_RE.search(cover) and str(year) in cover):
        sys.exit(f"fetch_gp: {path.name} cover does not read "
                 f"'січень-грудень {year}' — wrong attachment?")

    rows = [list(r) for r in wb["1"].iter_rows(values_only=True)]

    # Aggregate: Таблиця 1.1, рядок 1 — the row right after the first
    # 'А/Б/1/2' header (its label cell is unreliable: 2015 ships a
    # literal '0' there). Sanity-bounded below.
    col = value_column(rows, 20)
    total = None
    for i, row in enumerate(rows[:40]):
        if any(str(c).strip() == "Б" for c in row if c is not None):
            total = num(rows[i + 1][col])
            break
    if total is None or not 100_000 < total < 2_000_000:
        sys.exit(f"fetch_gp: {year} aggregate missing or implausible "
                 f"({total}) — workbook layout changed.")

    # Article 115 row inside the offences-against-life table of Розділ 1.
    art115 = hostilities = None
    for i, row in enumerate(rows):
        label = next((re.sub(r"\s+", " ", str(c)).strip()
                      for c in row if isinstance(c, str)
                      and "мисне вбив" in c.lower()), "")
        if ART115_RE.match(label):
            art115 = num(row[value_column(rows, i)])
        elif hostilities is None and any(
                isinstance(c, str) and HOSTILITIES_RE.search(c)
                for c in row):
            hostilities = num(row[value_column(rows, i)])
    if art115 is None or not 0 < art115 < total:
        sys.exit(f"fetch_gp: {year} Art. 115 row missing or implausible "
                 f"({art115} of {total}) — the row structure the manifest "
                 "flags as unverified has changed; re-pin by hand.")

    # Cross-check: the same row reappears in Розділ 2 (sheet '2', 2020+)
    # with a slightly narrower perimeter; the two must not drift apart.
    if "2" in wb.sheetnames:
        rows2 = [list(r) for r in wb["2"].iter_rows(values_only=True)]
        for i, row in enumerate(rows2):
            label = next((re.sub(r"\s+", " ", str(c)).strip()
                          for c in row if isinstance(c, str)
                          and "мисне вбив" in c.lower()), "")
            if ART115_RE.match(label):
                alt = num(row[value_column(rows2, i)])
                if alt and abs(alt - art115) / art115 > 0.05:
                    print(f"  WARNING {year}: Розділ 1 = {art115} vs "
                          f"Розділ 2 = {alt} (>5% apart)")
                break
    return total, art115, hostilities


def main() -> None:
    print("Office of the Prosecutor General — Unified report on criminal "
          "offences...")
    index = get(INDEX_URL)
    current = get(post_url(index, CURRENT_SLUG))
    archive = get(post_url(current, ARCHIVE_SLUG))
    years = full_year_attachments(current, archive)
    if len(years) < 10 or min(years) > 2013:
        sys.exit(f"fetch_gp: only {len(years)} full-year editions found "
                 f"({sorted(years)}) — expected 2013 onward; the OGP "
                 "archive page changed.")

    out = {}
    for year in sorted(years):
        out[year] = parse(download(year, years[year]), year)
        print(f"  {year}: total={out[year][0]} art115={out[year][1]} "
              f"of_which_hostilities={out[year][2]}")

    # Known-value checks (verified 2026-10-24, see module docstring).
    if not 3_000 <= out[2021][1] <= 4_000:
        sys.exit(f"fetch_gp: 2021 Art. 115 = {out[2021][1]}, expected "
                 "~3,230 — historical revision or a parse drift; "
                 "investigate before staging.")
    if not 300_000 <= out[2021][0] <= 350_000:
        sys.exit(f"fetch_gp: 2021 aggregate = {out[2021][0]}, expected "
                 "~321,443.")
    if list(out) != list(range(min(out), max(out) + 1)):
        sys.exit(f"fetch_gp: non-contiguous years {sorted(out)}")

    lines = [
        "# Office of the Prosecutor General of Ukraine — Unified report on",
        "# criminal offences (Єдиний звіт, Form № 1-monthly, OGP Order № 299",
        "# of 30.06.2020), January-December editions, Розділ 1 (sheet '1'),",
        "# column 'Обліковано кримінальних правопорушень у звітному періоді'.",
        "# POLICE/PROSECUTOR-RECORDED proceedings in the Unified Register of",
        "# Pre-trial Investigations — NOT the vital-statistics homicide count",
        "# and NOT a war-casualty count; never splice with the UNODC line.",
        "# From 2022 the Art. 115 row includes deaths from Russian attacks",
        "# opened as homicide proceedings; the publisher's own sub-row",
        "# 'committed in the course of hostilities' (reported from 2023 only)",
        "# quantifies part of that admixture — blank means NOT REPORTED, not",
        "# zero. Coverage is the territory where Ukraine's pre-trial",
        "# investigation bodies operate: the form is still nominally filed by",
        "# the ARC/Sevastopol prosecutor's offices, but they have not worked",
        "# on the occupied territories since 2014, and the 2022+ occupations",
        "# shrink coverage further — the publisher prints no explicit",
        "# territorial note. First full-year edition is 2013; earlier years",
        "# are not published in this series. Partial-year editions (2026) are",
        "# deliberately excluded. Source: new.gp.gov.ua statistics posts,",
        "# attachments on old.gp.gov.ua (file ids resolved at fetch time).",
        "year,recorded_offences_total,homicide_art115_recorded,"
        "homicide_art115_of_which_hostilities",
    ]
    for year in sorted(out):
        total, art115, host = out[year]
        lines.append(f"{year},{total},{art115},"
                     f"{host if host is not None else ''}")
    dest = DATA / "security_gp.csv"
    dest.write_text("\n".join(lines) + "\n")
    print(f"  -> data/security/security_gp.csv "
          f"({min(out)}-{max(out)}, {len(out)} years)")
    print("Done. Diff against the curated CSV before promoting anything.")


if __name__ == "__main__":
    main()
