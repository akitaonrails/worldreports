#!/usr/bin/env python3
"""Refresh the Home Office migration and police-homicide series.

Usage:  python3 -I countries/united-kingdom/fetch/fetch_homeoffice.py

The Home Office publishes the immigration system statistics as one landing
page with ~50 xlsx/ods attachments on assets.publishing.service.gov.uk.
Those media URLs die with every edition, so they are never hardcoded: the
fetcher reads the gov.uk content API
(/api/content/government/statistical-data-sets/immigration-system-
statistics-data-tables and .../police-recorded-crime-and-outcomes-open-
data-tables), picks the latest edition of each dataset family by its
title, and downloads from whatever URL the API is serving today.

Things this source will mislead you about.

**Applications, decisions and the backlog are three different numbers.**
An asylum claim arrives in Asy_D01, is decided in Asy_D02 (a year or more
later, on average), and sits in the Asy_D03 backlog in between. They are
NEVER summed into one "asylum" line (SOURCING §4a). All three count PEOPLE
(main applicants plus dependants), the Home Office headline convention.

**"Decisions" in Asy_D02 is not "initial decisions".** The sheet's
Decisions column also carries withdrawals and administrative outcomes.
The Home Office's headline "people who received an initial decision" is
grants (protection + other leave) plus refusals only; withdrawals and
administrative outcomes are a separate column here, never folded in.

**The asylum series is revised.** The year-ending-December-2024 edition
headlined 108,138 people claiming asylum in 2024; the current vintage of
Asy_D01 reads lower for the same year (the Atlas caseworking migration
and data-quality work, user guide §2.8). The fetcher asserts against the
CURRENT edition's headlines, not the old one.

**A visa grant is a permission, not an arrival.** A grant may never be
used and one person can hold several in a year; the grants series is the
demand side of the ONS net-migration story and is drawn beside flows,
never summed with them. The applicant-type rows ("All", "Main
Applicant", "Dependant") are DISJOINT — a visa type is published either
as "All" or split — so every row is summed; verified against the
headline per-route totals. Transit visas sit inside the "Other" group in
the dataset but the Home Office headline "Other" excludes them, so they
are a column of their own.

**A detection is an event, not a person.** The illegal-entry-routes
dataset counts DETECTED arrivals by method (small boats, inadequately
documented air, port and in-country detections) and only the detected
fraction. It is drawn as its own labelled series, never as "illegal
immigration". The Home Office warns the in-country and port detections
read lower than in earlier releases while data collection is reworked.

**Settlement is two regimes in one file family.** The detailed Se_D01/
Se_D02 only cover 2020 onwards (the CID → Atlas caseworking transition);
the summary tables (ods) carry 2010–2025 and state that pre-2020 data
live there. The fetcher draws the summary's calendar years and asserts
they equal the detailed dataset for every overlapping year. Up to May
2019 the series includes settlement ON ARRIVAL (immediate-settlement
visas); afterwards it does not — that component is its own column so the
break is visible. 2006–2009 is not published in the current edition and
stays blank.

**Homicide is England & Wales, year ending March.** No UK-wide police
homicide count exists (Scotland and NI are meta.pending in
SOURCES.yaml). The series is the Home Office Homicide Index, published
in the ONS homicide appendix tables (Table 1): offences INITIALLY
recorded as homicide and offences CURRENTLY recorded as homicide are two
columns, never one — the gap is cases later reclassified. One offence
per victim. The fiscal_year column is the starting year (Apr 2006 to Mar
2007 is fiscal_year 2006), the atlas's April–March convention. Years
carrying a multi-victim event (Hillsborough recorded in YE March 2017,
the 2017 terror attacks, the 2019 Essex lorry deaths) keep the table's
own note text in a `table_note` column, so a spike is never read as a
trend. The police-recorded-crime open-data set is also enumerated, but
its offence tables are large ods workbooks of rolling quarters by force
area that duplicate the Homicide Index; its recrime-geo-pfa.csv is a
geographical lookup (not data) and is used for exactly one thing:
asserting the 43 territorial forces of England & Wales, the scope of the
homicide series.

Every spot value below is asserted (the run fails before a sidecar is
written) against the headline the Home Office / ONS published for the
same reference period.

Outputs are `_from_homeoffice.csv` sidecars. Nothing here overwrites a
curated file.
"""
from __future__ import annotations

import csv
import io
import json
import re
import time
import urllib.request
import warnings
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from pathlib import Path

try:
    import openpyxl
except ImportError:
    raise SystemExit(
        "openpyxl is required to read the Home Office xlsx datasets "
        "(pip install -r requirements-fetch.txt). Refusing to skip the "
        "xlsx sources silently."
    )

# The appendix tables trip a harmless openpyxl defined-name warning.
warnings.filterwarnings("ignore", message="Print area cannot be set")

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
CACHE = ROOT / "cache"

Y0, Y1 = 2006, 2026          # atlas window; 2026 flows are partial and blank
FULL = {"Q1", "Q2", "Q3", "Q4"}

CONTENT_API = "https://www.gov.uk/api/content/government/statistical-data-sets/{slug}"
IMMIGRATION_SLUG = "immigration-system-statistics-data-tables"
CRIME_SLUG = "police-recorded-crime-and-outcomes-open-data-tables"

# Homicide Index, hosted by the ONS: the dataset page links its current
# edition; the media URL is resolved from the page, never hardcoded.
HOMICIDE_PAGE = ("https://www.ons.gov.uk/peoplepopulationandcommunity/"
                 "crimeandjustice/datasets/"
                 "appendixtableshomicideinenglandandwales")

MONTHS = {m: i for i, m in enumerate(
    ["January", "February", "March", "April", "May", "June", "July",
     "August", "September", "October", "November", "December"], start=1)}

# title prefix -> (cache stem, extension), for the families this draws.
IMMIGRATION_FILES = {
    "Asylum claims and initial decisions detailed datasets":
        ("asylum-claims-datasets", ".xlsx"),
    "Asylum claims awaiting a decision detailed datasets":
        ("asylum-awaiting-datasets", ".xlsx"),
    "Entry clearance visa applications and outcomes detailed datasets":
        ("visa-outcomes-datasets", ".xlsx"),
    "Settlement detailed datasets": ("settlement-datasets", ".xlsx"),
    "Settlement summary tables": ("settlement-summary", ".ods"),
    "Citizenship detailed datasets": ("citizenship-datasets", ".xlsx"),
    "Illegal entry routes to the UK detailed dataset":
        ("illegal-entry-routes-dataset", ".xlsx"),
}

# Published headlines, each copied from the release text for the period
# named, asserted against what this run computes. Rolling year-ending
# figures are revised a little each quarter, hence the tolerance; a check
# is skipped (and says so) only when the data no longer cover the period.
#   asylum/visas/settlement/citizenship/irregular: "Immigration system
#   statistics, year ending June 2026" (gov.uk, 27 Aug 2026), and the
#   year ending December 2024 edition for the two calendar-2024 figures;
#   homicide: ONS "Homicide in England and Wales: year ending March 2025".
YE_JUN_2026 = ("2025 Q3", "2025 Q4", "2026 Q1", "2026 Q2")
TOLERANCE = 0.02


def check(label: str, got, want: int, tol: float = TOLERANCE) -> None:
    if got is None:
        raise SystemExit(f"  !! {label}: period missing from the data")
    if abs(got - want) > tol * want:
        raise SystemExit(f"  !! {label} = {got:,}, published {want:,} — "
                         f"outside {tol:.0%}; not writing")
    print(f"  ok {label}: {got:,} (published {want:,})")


def fetch(url: str, cache_file: Path, binary_check: bytes) -> bytes:
    """One payload, shape-checked BEFORE caching, with retries.

    A lying endpoint (an HTML error page under HTTP 200, a truncated
    download) must never land in the cache: xlsx/ods files are zip
    archives and must start with PK, the content API with '{'.
    """
    CACHE.mkdir(exist_ok=True)
    if cache_file.exists():
        return cache_file.read_bytes()
    last = None
    for attempt in range(4):
        try:
            req = urllib.request.Request(url, headers={"User-Agent":
                                                       "worldreports-fetch"})
            with urllib.request.urlopen(req, timeout=300) as r:
                blob = r.read()
            if not blob.lstrip()[:len(binary_check)].startswith(binary_check):
                raise ValueError(f"payload does not start with "
                                 f"{binary_check!r} (error page?)")
            if binary_check == b"PK":
                zipfile.ZipFile(io.BytesIO(blob)).testzip()
            cache_file.write_bytes(blob)
            return blob
        except Exception as e:  # noqa: BLE001
            last = e
            # ONS answers 429 when a second request follows the dataset-page
            # fetch too closely; give rate limits room to reset.
            wait = 15 * (attempt + 1) if "429" in str(e) else 2 * (attempt + 1)
            time.sleep(wait)
    raise RuntimeError(f"{url}: {last}")


def edition(title: str) -> tuple[int, int]:
    """(year, month) of 'year ending June 2026'; undated titles sort first."""
    m = re.search(r"year ending (\w+) (\d{4})", title)
    return (int(m.group(2)), MONTHS[m.group(1)]) if m else (0, 0)


def attachments(slug: str) -> list[dict]:
    blob = fetch(CONTENT_API.format(slug=slug),
                 CACHE / f"homeoffice_content-{slug}.json", b"{")
    payload = json.loads(blob.decode("utf-8"))
    try:
        return payload["details"]["attachments"]
    except KeyError:
        raise RuntimeError(f"content API {slug}: no details.attachments")


def immigration_files() -> dict[str, Path]:
    """Latest file of each needed family, enumerated from the content API."""
    found = {}
    listing = attachments(IMMIGRATION_SLUG)
    for prefix, (stem, ext) in IMMIGRATION_FILES.items():
        cands = [a for a in listing
                 if a.get("title", "").startswith(prefix)
                 and a.get("url", "").endswith(ext)]
        if not cands:
            raise RuntimeError(f"content API: '{prefix}...' {ext} not found")
        best = max(cands, key=lambda a: edition(a["title"]))
        found[stem] = CACHE / f"homeoffice_{stem}{ext}"
        fetch(best["url"], found[stem], b"PK")
        print(f"  {stem}: {best['title']}")
    return found


def rows(path: Path, sheet: str, header_row: int = 2):
    """(header, data-row iterator) for a Data_* sheet, columns by name."""
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[sheet]
    it = ws.iter_rows(min_row=header_row, values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    return header, it


def ods_table(path: Path, name: str) -> list[list[str | None]]:
    """One sheet of an ods workbook as rows of cell strings (stdlib only).

    An ods file is a zip holding content.xml; numeric cells carry their
    value in office:value, text cells in text:p. Repeated empty columns
    are capped so a sheet styled to column XFD stays small.
    """
    t_ = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
    x_ = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
    o_ = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
    root = ET.fromstring(zipfile.ZipFile(path).read("content.xml"))
    for table in root.iter(t_ + "table"):
        if table.get(t_ + "name") != name:
            continue
        out = []
        for row in table.iter(t_ + "table-row"):
            cells: list[str | None] = []
            for cell in row:
                n = min(int(cell.get(t_ + "number-columns-repeated", "1")), 64)
                v = cell.get(o_ + "value")
                if v is None:
                    v = " ".join("".join(p.itertext())
                                 for p in cell.iter(x_ + "p")) or None
                cells += [v] * n
            while cells and cells[-1] is None:
                cells.pop()
            out.append(cells)
        return out
    raise RuntimeError(f"{path.name}: no sheet {name}")


def number(raw) -> int | None:
    """One cell as a count, or None; 'z'/'[x]'/':' markers stay None."""
    if raw is None:
        return None
    text = str(raw).replace(",", "").strip()
    if re.fullmatch(r"\d+", text):
        return int(text)
    try:
        return int(float(text))
    except ValueError:
        return None


def quarter(raw) -> str:
    return str(raw).split()[-1] if raw else ""


def write(rel: str, header: list[str], columns: list[dict], label: str,
          ) -> int:
    path = DATA / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    rows_out = 0
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        for y in range(Y0, Y1 + 1):
            cells = [col.get(y) for col in columns]
            if all(v is None for v in cells):
                continue
            w.writerow([y] + ["" if v is None else v for v in cells])
            rows_out += 1
    print(f"  -> data/{rel}  {rows_out} rows ({label})")
    return rows_out


def quarterly_flows(header, it, value_col: str, keep=lambda r: True,
                    ) -> tuple[Counter, Counter]:
    """({year: total} for complete calendar years, YE-June-2026 total).

    A year short of its four quarters is a partial year and is dropped so
    the CSV leaves it blank rather than drawing a collapse that never
    happened. `keep(row_dict)` filters rows before summing.
    """
    i_y = header.index("Year")
    i_q = header.index("Quarter")
    i_v = header.index(value_col)
    totals: Counter[int] = Counter()
    quarters: dict[int, set] = {}
    ye = 0
    seen_ye = set()
    for r in it:
        y = number(r[i_y])
        v = number(r[i_v])
        if y is None:
            continue
        quarters.setdefault(y, set()).add(quarter(r[i_q]))
        if v is None or not keep(dict(zip(header, r))):
            continue
        if Y0 <= y <= Y1:
            totals[y] += v
        if r[i_q] in YE_JUN_2026:
            ye += v
            seen_ye.add(r[i_q])
    complete = {y for y, qs in quarters.items() if qs >= FULL}
    return (Counter({y: v for y, v in totals.items() if y in complete}),
            ye if seen_ye == set(YE_JUN_2026) else None)


# ---------------------------------------------------------------- asylum

GRANTS = {"Grant of Protection", "Grant of Other Leave"}
DECIDED = GRANTS | {"Refused"}
OTHER_OUTCOMES = {"Withdrawn", "Administrative Outcome"}


def asylum(files: dict[str, Path]) -> int:
    header, it = rows(files["asylum-claims-datasets"], "Data_Asy_D01")
    applications, app_ye = quarterly_flows(header, it, "Claims")

    groups = set()

    def group(name):
        def keep(d):
            groups.add(d["Case outcome group"])
            return d["Case outcome group"] in name
        return keep

    header, it = rows(files["asylum-claims-datasets"], "Data_Asy_D02")
    decisions, dec_ye = quarterly_flows(header, it, "Decisions",
                                        group(DECIDED))
    header, it = rows(files["asylum-claims-datasets"], "Data_Asy_D02")
    grants, _ = quarterly_flows(header, it, "Decisions", group(GRANTS))
    header, it = rows(files["asylum-claims-datasets"], "Data_Asy_D02")
    other, _ = quarterly_flows(header, it, "Decisions", group(OTHER_OUTCOMES))
    if groups - DECIDED - OTHER_OUTCOMES:
        raise SystemExit(f"  !! Asy_D02: unknown outcome groups "
                         f"{groups - DECIDED - OTHER_OUTCOMES}")

    # The backlog is a STOCK at a snapshot date ("31 Dec 2010"); the
    # year-end reading is the 31 December snapshot, not a sum. Only the
    # "Pending initial decision" stage is the headline backlog.
    header, it = rows(files["asylum-awaiting-datasets"], "Data_Asy_D03")
    i_d, i_v = header.index("Date (as at...)"), header.index("Claims")
    i_s = header.index("Application stage")
    backlog: Counter[int] = Counter()
    snapshot: Counter[str] = Counter()
    for r in it:
        v = number(r[i_v])
        if r[i_d] is None or v is None:
            continue
        if r[i_s] != "Pending initial decision":
            continue
        snapshot[str(r[i_d])] += v
        parts = str(r[i_d]).split()
        if len(parts) == 3 and parts[:2] == ["31", "Dec"]:
            y = int(parts[2])
            if Y0 <= y <= Y1:
                backlog[y] += v

    check("asylum claims, people, YE Jun 2026", app_ye, 85_891)
    check("asylum initial decisions, people, YE Jun 2026", dec_ye, 115_872)
    check("awaiting initial decision, people, 30 Jun 2026",
          snapshot.get("30 Jun 2026"), 40_168)
    check("awaiting initial decision, people, 31 Dec 2024",
          backlog.get(2024), 124_802)

    return write("migration/asylum_from_homeoffice.csv",
                 ["year", "applications_persons", "initial_decisions_persons",
                  "initial_grants_persons", "withdrawn_or_admin_outcome_persons",
                  "awaiting_initial_decision_31dec_persons"],
                 [applications, decisions, grants, other, backlog],
                 "people; three series, never summed; backlog is the "
                 "31 Dec stock")


# ----------------------------------------------------------------- visas

VISA_GROUPS = {"Work": "issued_work", "Study": "issued_study",
               "Visitor": "issued_visitor", "Family": "issued_family",
               "Other": "issued_other_excl_transit"}


def visa_totals(path: Path, sheet: str, value_col: str,
                outcome: str | None):
    """{(year, column): total} over every applicant-type row (disjoint),
    complete years, and the YE-June-2026 totals by column."""
    header, it = rows(path, sheet, 4)
    i_y = header.index("Year")
    i_g = header.index("Visa type group")
    i_s = header.index("Visa type subgroup")
    i_v = header.index(value_col)
    i_c = header.index("Case outcome") if outcome else None
    i_q = header.index("Quarter")
    out: Counter = Counter()
    ye: Counter = Counter()
    quarters: dict[int, set] = {}
    for r in it:
        y = number(r[i_y])
        if y is None:
            continue
        quarters.setdefault(y, set()).add(quarter(r[i_q]))
        v = number(r[i_v])
        if v is None:
            continue
        if i_c is not None and str(r[i_c]) != outcome:
            continue
        g = str(r[i_g])
        if g not in VISA_GROUPS:
            raise SystemExit(f"  !! {sheet}: unknown visa type group {g!r}")
        col = "issued_transit" if r[i_s] == "Transit" else VISA_GROUPS[g]
        if Y0 <= y <= Y1:
            out[(y, col)] += v
        if r[i_q] in YE_JUN_2026:
            ye[col] += v
    return out, {y for y, qs in quarters.items() if qs >= FULL}, ye


def visas(files: dict[str, Path]) -> int:
    issued, ok_i, ye = visa_totals(files["visa-outcomes-datasets"],
                                   "Data_Vis_D02", "Decisions", "Issued")
    applied, ok_a, _ = visa_totals(files["visa-outcomes-datasets"],
                                   "Data_Vis_D01", "Applications", None)
    ok = ok_i & ok_a
    cols = [*VISA_GROUPS.values(), "issued_transit"]
    by_col: dict[str, Counter] = {c: Counter() for c in cols}
    total: Counter[int] = Counter()
    apps: Counter[int] = Counter()
    for (y, c), v in issued.items():
        if y in ok:
            by_col[c][y] += v
            total[y] += v
    for (y, _), v in applied.items():
        if y in ok:
            apps[y] += v

    check("study visas granted, YE Jun 2026", ye["issued_study"], 385_974)
    check("work visas granted, YE Jun 2026", ye["issued_work"], 234_841)
    check("family visas granted, YE Jun 2026", ye["issued_family"], 57_590)
    check("other visas granted (excl. transit), YE Jun 2026",
          ye["issued_other_excl_transit"], 49_082)
    check("transit visas granted, YE Jun 2026", ye["issued_transit"], 19_684)
    check("non-visitor/transit visas granted, YE Jun 2026",
          sum(ye[c] for c in ("issued_study", "issued_work", "issued_family",
                              "issued_other_excl_transit")), 727_487)
    # "2.2 million Visitor visas" — the headline is rounded.
    check("visitor visas granted, YE Jun 2026", ye["issued_visitor"],
          2_200_000, tol=0.025)

    return write("migration/visas_from_homeoffice.csv",
                 ["year", "applications", "visas_issued", *cols],
                 [apps, total, *[by_col[c] for c in cols]],
                 "grants are permissions, not arrivals")


# ------------------------------------------------ settlement, citizenship

def settlement(files: dict[str, Path]) -> int:
    # Detailed, 2020+ (Atlas era): annual Se_D01, quarterly Se_D02.
    header, it = rows(files["settlement-datasets"], "Data_Se_D01")
    i_y, i_o, i_v = (header.index("Year"), header.index("Case Outcome"),
                     header.index("Decisions"))
    d_grants: Counter[int] = Counter()
    d_refusals: Counter[int] = Counter()
    for r in it:
        y, v = number(r[i_y]), number(r[i_v])
        if y is None or v is None:
            continue
        outcome = str(r[i_o])
        if outcome not in ("Grant", "Refusal"):
            raise SystemExit(f"  !! Se_D01: unknown outcome {outcome!r}")
        (d_grants if outcome == "Grant" else d_refusals)[y] += v

    header, it = rows(files["settlement-datasets"], "Data_Se_D02")
    i_q, i_o2 = header.index("Quarter"), header.index("Outcome")
    i_v2 = header.index("Decisions")
    seen, ye = set(), 0
    for r in it:
        v = number(r[i_v2])
        if v is not None and r[i_o2] == "Grant" and r[i_q] in YE_JUN_2026:
            ye += v
            seen.add(r[i_q])
    grants_ye = ye if seen == set(YE_JUN_2026) else None

    # Summary tables, 2010+: Se_02 grants and refusals, Se_06 on-arrival.
    def by_label(table, label):
        head = next(r for r in table if r and r[0] == "Year")
        row = next((r for r in table if r and r[0] == label), None)
        if row is None:
            raise SystemExit(f"  !! settlement summary: no row {label!r}")
        out = {}
        for h, v in zip(head[1:], row[1:]):
            if h and re.fullmatch(r"\d{4}", h) and Y0 <= int(h) <= Y1:
                out[int(h)] = number(v)
        return out

    se02 = ods_table(files["settlement-summary"], "Se_02")
    se06 = ods_table(files["settlement-summary"], "Se_06")
    grants = by_label(se02, "Total Grants of settlement")
    refusals = by_label(se02, "Total Refusals")
    total06 = by_label(se06, "Grand Total")
    arr_c = by_label(se06, "Settlement on arrival Commonwealth citizens")
    arr_f = by_label(se06, "Settlement on arrival Foreign nationals")
    # 'z' (not applicable) from 2020 stays blank, not zero.
    on_arrival = {y: arr_c[y] + arr_f[y] for y in total06
                  if arr_c.get(y) is not None and arr_f.get(y) is not None}

    for y in sorted(set(d_grants) & set(range(Y0, Y1 + 1))):
        if grants.get(y) != d_grants[y] or refusals.get(y) != d_refusals[y]:
            raise SystemExit(
                f"  !! settlement {y}: summary {grants.get(y)}/"
                f"{refusals.get(y)} vs detailed {d_grants[y]}/"
                f"{d_refusals[y]} (grants/refusals)")
    for y, v in grants.items():
        if total06.get(y) != v:
            raise SystemExit(f"  !! settlement {y}: Se_02 {v} vs Se_06 "
                             f"{total06.get(y)}")
    print(f"  ok settlement summary = detailed for "
          f"{min(d_grants)}–{max(d_grants)}; Se_02 = Se_06 for every year")
    check("settlement grants, YE Jun 2026", grants_ye, 199_628)

    return write("migration/settlement_from_homeoffice.csv",
                 ["year", "settlement_grants", "settlement_refusals",
                  "of_which_settlement_on_arrival"],
                 [grants, refusals, on_arrival],
                 "grants, not people; 2010+, on-arrival ends May 2019")


def citizenship(files: dict[str, Path]) -> int:
    header, it = rows(files["citizenship-datasets"], "Data_Cit_D02")
    grants, ye = quarterly_flows(header, it, "Grants")
    check("citizenship grants, YE Jun 2026", ye, 245_520)
    return write("migration/citizenship_from_homeoffice.csv",
                 ["year", "citizenship_grants"], [grants],
                 "grants of British citizenship; the migr_acq continuation")


# ------------------------------------------------------------- irregular

def irregular(files: dict[str, Path]) -> int:
    path = files["illegal-entry-routes-dataset"]
    header, it = rows(path, "Data_IER_D01")
    total, total_ye = quarterly_flows(header, it, "Number of detections")
    header, it = rows(path, "Data_IER_D01")
    methods = set()

    def boat(d):
        methods.add(d["Method of entry"])
        return d["Method of entry"] == "Small boat arrivals"
    boats, _ = quarterly_flows(header, it, "Number of detections", boat)
    if "Small boat arrivals" not in methods or any(
            "total" in str(m).lower() for m in methods):
        raise SystemExit(f"  !! IER_D01 methods look wrong: {methods}")

    check("small-boat arrivals, 2024", boats.get(2024), 36_816)
    check("detected arrivals via illegal routes, YE Jun 2026", total_ye,
          38_086)
    return write("migration/irregular_from_homeoffice.csv",
                 ["year", "irregular_detections", "small_boat_detections"],
                 [total, boats],
                 "detections are events, not people; series starts 2018")


# -------------------------------------------------------------- homicide

def homicide() -> int:
    """Homicide Index, E&W, year ending March — from the ONS appendix."""
    blob = fetch(HOMICIDE_PAGE, CACHE / "homeoffice_homicide-page.html",
                 b"<")
    page = blob.decode("utf-8", errors="replace")
    m = re.search(r"file\?uri=([^\"']*appendixtableshomicideinenglandandwales"
                  r"[^\"']*\.xlsx)", page)
    if not m:
        raise RuntimeError("ONS homicide dataset page: no current xlsx link")
    time.sleep(3)  # the ONS rate-limits a download that follows the page
    xlsx = fetch("https://www.ons.gov.uk/file?uri=/" + m.group(1).lstrip("/"),
                 CACHE / "homeoffice_homicide-appendix-tables.xlsx", b"PK")

    wb = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True)
    notes = {}
    for r in wb["Notes"].iter_rows(values_only=True):
        if isinstance(r[0], int) and r[1]:
            notes[r[0]] = str(r[1]).strip()
    it = wb["Table 1"].iter_rows(values_only=True)
    for header in it:
        if header and header[0] == "Year":
            break
    else:
        raise RuntimeError("homicide Table 1: no header row")
    header = [str(h or "").replace("\n", " ") for h in header]
    i_init = header.index("Number of offences initially recorded as homicide")
    i_curr = header.index("Number of offences currently recorded as homicide")
    initially: dict[int, int] = {}
    currently: dict[int, int] = {}
    note: dict[int, str] = {}
    for r in it:
        label = str(r[0] or "")
        mm = re.fullmatch(r"Apr (\d{4}) to Mar (\d{4})\s*(?:\[note (\d+)\])?",
                          label.strip())
        if not mm:
            continue  # the pre-1998 rows are calendar years
        start, end = int(mm.group(1)), int(mm.group(2))
        if end != start + 1:
            raise RuntimeError(f"homicide Table 1: bad YE-March row {label}")
        if Y0 <= start <= Y1:
            initially[start] = number(r[i_init])
            currently[start] = number(r[i_curr])
            if mm.group(3):
                note[start] = notes[int(mm.group(3))]

    # The police-recorded-crime open data set is enumerated as the manifest
    # requires; its geographical lookup serves as the scope assertion: 43
    # territorial forces is England & Wales, and only England & Wales.
    geo = None
    for a in attachments(CRIME_SLUG):
        if a.get("title") == "Recorded crime data geographical reference table":
            geo = a["url"]
    if geo is None:
        raise RuntimeError("content API: recrime-geo-pfa.csv not found")
    blob = fetch(geo, CACHE / "homeoffice_recrime-geo-pfa.csv", b"CSP")
    forces = {row["Force Name"] for row in
              csv.DictReader(io.StringIO(blob.decode("utf-8-sig")))}
    if len(forces) != 43:
        raise RuntimeError(f"recrime-geo-pfa.csv: {len(forces)} forces, "
                           f"expected the 43 of England & Wales")
    print("  ok recrime-geo-pfa.csv: 43 territorial forces (E&W scope)")

    # "522 victims of homicide recorded in YE March 2025 ... previous year
    # (566 victims)"; "the 524 homicides initially recorded in YE March 2025".
    check("homicides currently recorded, YE Mar 2025", currently.get(2024),
          522, tol=0)
    check("homicides currently recorded, YE Mar 2024", currently.get(2023),
          566, tol=0.01)
    check("homicides initially recorded, YE Mar 2025", initially.get(2024),
          524, tol=0)

    return write("security/homicide_police_from_homeoffice.csv",
                 ["fiscal_year", "offences_initially_recorded",
                  "offences_currently_recorded", "table_note"],
                 [initially, currently, note],
                 "Homicide Index, E&W, YE March; fiscal_year = starting year")


def main() -> None:
    print("Home Office — gov.uk content API")
    files = immigration_files()
    asylum(files)
    visas(files)
    settlement(files)
    citizenship(files)
    irregular(files)
    homicide()
    print("Done. Diff the sidecars before adopting.")


if __name__ == "__main__":
    main()
