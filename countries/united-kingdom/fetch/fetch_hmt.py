#!/usr/bin/env python3
"""Refresh the public spending series from HM Treasury's PESA.

Usage:  make C=united-kingdom fetch

PESA (Public Expenditure Statistical Analyses) is the Treasury's annual
compendium of public spending, published each July as a main document plus
one workbook per chapter. This fetcher takes Chapter 4, Table 4.2 — public
sector expenditure on services by UN COFOG function, with Total Managed
Expenditure — and writes the single sidecar the manifest declares:
fiscal/spending_by_function_from_hmt.csv.

Three things about this source that a chart will get wrong if nobody
writes them down:

**The PESA year is not the calendar year.** Rows are labelled `2024-25`
and run 1 April to 31 March, so the `fiscal_year` column here holds the
starting year and must never be plotted against a calendar series without
saying so.

**These are the Treasury's own COFOG assignments, not Eurostat's.** The UK
no longer makes the Eurostat COFOG transmission, so there is no harmonised
twin to draw beside this series. The `sub_*` columns are the table's
published "of which" sub-function detail: components of their function,
never additive across.

**TME is not the sum of the functions.** Between the last function and TME
sit EU transactions (signed as published: negative was a net receipt) and
the accounting adjustments that reconcile the services aggregate to the
national-accounts total. TME excludes the temporary effects of banks being
classified to the public sector (the table's own note).

The edition is never hardcoded: the collection page's change history names
the latest PESA year, the gov.uk search API finds the publication, and the
content API enumerates its xlsx attachments — assets.publishing.service.gov.uk
media URLs are per-edition and must not be baked in.

The workbook is cached under cache/ on the first run; delete it to force a
refresh. Nothing here overwrites a curated file.
"""
from __future__ import annotations

import csv
import json
import re
import sys
import urllib.parse
import urllib.request
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("fetch_hmt.py needs openpyxl to read the PESA xlsx "
             "(pip install -r requirements-fetch.txt) — not installed, "
             "nothing fetched")

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
CACHE = ROOT / "cache"
OUT = DATA / "fiscal/spending_by_function_from_hmt.csv"

COLLECTION = ("https://www.gov.uk/api/content/government/collections/"
              "public-expenditure-statistical-analyses-pesa")
SEARCH = "https://www.gov.uk/api/search.json?q={}&count=10"

FY = re.compile(r"^(\d{4})-\d{2}$")

# Normalised row label (footnote markers stripped) -> column name. The
# functions keep their COFOG numbering in the column name so a later edit
# cannot silently reorder them; sub_* are the published "of which" lines.
ROWS = {
    "total managed expenditure": "tme_bn",
    "public sector expenditure on services": "expenditure_on_services_bn",
    "accounting adjustments": "accounting_adjustments_bn",
    "eu transactions": "eu_transactions_bn",
    "1. general public services": "general_public_services_bn",
    "2. defence": "defence_bn",
    "3. public order and safety": "public_order_and_safety_bn",
    "4. economic affairs": "economic_affairs_bn",
    "5. environment protection": "environment_protection_bn",
    "6. housing and community amenities": "housing_and_community_amenities_bn",
    "7. health": "health_bn",
    "8. recreation, culture and religion": "recreation_culture_and_religion_bn",
    "9. education": "education_bn",
    "10. social protection": "social_protection_bn",
    "of which: public and common services": "sub_public_and_common_services_bn",
    "of which: international services": "sub_international_services_bn",
    "of which: public sector debt interest": "sub_public_sector_debt_interest_bn",
    "of which: enterprise and economic development":
        "sub_enterprise_and_economic_development_bn",
    "of which: science and technology": "sub_science_and_technology_bn",
    "of which: employment policies": "sub_employment_policies_bn",
    "of which: agriculture, fisheries and forestry":
        "sub_agriculture_fisheries_and_forestry_bn",
    "of which: transport": "sub_transport_bn",
}
ORDER = ["tme_bn", "expenditure_on_services_bn", "accounting_adjustments_bn",
         "eu_transactions_bn", "general_public_services_bn", "defence_bn",
         "public_order_and_safety_bn", "economic_affairs_bn",
         "environment_protection_bn", "housing_and_community_amenities_bn",
         "health_bn", "recreation_culture_and_religion_bn", "education_bn",
         "social_protection_bn", "sub_public_and_common_services_bn",
         "sub_international_services_bn", "sub_public_sector_debt_interest_bn",
         "sub_enterprise_and_economic_development_bn",
         "sub_science_and_technology_bn", "sub_employment_policies_bn",
         "sub_agriculture_fisheries_and_forestry_bn", "sub_transport_bn"]


def get(url: str, timeout: int = 300) -> bytes:
    with urllib.request.urlopen(url, timeout=timeout) as r:
        return r.read()


def latest_publication() -> tuple[str, str]:
    """(year, base_path) of the newest PESA, from the collection page."""
    col = json.loads(get(COLLECTION, timeout=60))
    year = None
    for entry in col["details"]["change_history"]:
        m = re.search(r"Public Expenditure Statistical Analyses (\d{4})",
                      entry["note"])
        if m:
            year = m.group(1)
            break
    if year is None:
        raise RuntimeError("no PESA edition named in the collection's "
                           "change history")
    query = urllib.parse.quote(f'"Public Expenditure Statistical Analyses '
                               f'{year}"')
    results = json.loads(get(SEARCH.format(query), timeout=60))["results"]
    for r in results:
        if r["link"].endswith(f"-{year}"):
            return year, r["link"]
    raise RuntimeError(f"the search API cannot find the PESA {year} "
                       f"publication page")


def workbook(year: str, base_path: str) -> openpyxl.Workbook:
    pub = json.loads(get(f"https://www.gov.uk/api/content{base_path}",
                         timeout=60))
    url = None
    for a in pub["details"].get("attachments", []):
        if (re.search(rf"PESA {year} Chapter 4 Tables", a.get("title", ""))
                and "spreadsheetml" in a.get("content_type", "")):
            url = a["url"]
            break
    if url is None:
        raise RuntimeError(f"no Chapter 4 xlsx among the PESA {year} "
                           f"attachments")
    CACHE.mkdir(exist_ok=True)
    local = CACHE / f"pesa{year}_chapter4_tables.xlsx"
    if not local.exists():
        print(f"  downloading {url}")
        blob = get(url)
        # An error page under HTTP 200 is not a workbook. Every xlsx is a
        # zip, so the magic number is the cheapest possible check.
        if blob[:2] != b"PK":
            raise RuntimeError("the PESA download is not a workbook")
        local.write_bytes(blob)
    return openpyxl.load_workbook(local, data_only=True)


def label(value) -> str:
    """A row label normalised: lowercase, footnote markers removed."""
    text = re.sub(r"\(\d+\)", "", str(value))
    return " ".join(text.lower().split()).rstrip(" ,")


def cell(value) -> str:
    """A spreadsheet cell as it should appear in a CSV.

    PESA writes a negligible figure as a dash, which a chart would
    otherwise draw as zero.
    """
    if value is None or value == "" or value == "-":
        return ""
    if isinstance(value, (int, float)):
        return f"{value:g}"
    text = str(value).strip().replace(",", "")
    return text if text not in ("-", "–", "..", "n/a") else ""


def extract(wb: openpyxl.Workbook) -> tuple[list[str], dict[str, list[str]]]:
    """(fiscal years, {column name: values}) from Table 4.2."""
    sheet = next((s for s in wb.sheetnames if s.startswith("4_2")), None)
    if sheet is None:
        raise RuntimeError("no Table 4.2 sheet in the Chapter 4 workbook")
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = None
    for i, row in enumerate(rows):
        hits = [j for j, v in enumerate(row)
                if v is not None and FY.match(str(v).strip())]
        if len(hits) >= 5:
            header = i
            break
    if header is None:
        raise RuntimeError("no fiscal-year header row in Table 4.2")
    years = {j: FY.match(str(v).strip()).group(1)
             for j, v in enumerate(rows[header])
             if v is not None and FY.match(str(v).strip())}
    cols: dict[str, list[str]] = {}
    for row in rows[header + 1:]:
        if not row or row[0] is None:
            continue
        name = ROWS.get(label(row[0]))
        if name is None:
            continue
        cols[name] = [cell(row[j]) if j < len(row) else "" for j in years]
    missing = [n for n in ROWS.values() if n not in cols]
    if missing:
        raise RuntimeError(f"Table 4.2 rows not found: {missing}")
    return list(years.values()), cols


def verify(years: list[str], cols: dict[str, list[str]]) -> None:
    """Shape checks against figures PESA has already published."""
    fy2020 = years.index("2020")
    tme2020 = float(cols["tme_bn"][fy2020])
    # The pandemic year's TME is a widely published number (~£1,103bn);
    # a wrong table or unit lands nowhere near it.
    if not 1000 < tme2020 < 1200:
        raise RuntimeError(f"TME 2020-21 reads {tme2020} £bn, expected "
                           f"~1,103")
    # TME = expenditure on services + accounting adjustments, exactly,
    # in every year the table publishes.
    for k, y in enumerate(years):
        tme, tes, adj = (float(cols[n][k]) for n in
                         ("tme_bn", "expenditure_on_services_bn",
                          "accounting_adjustments_bn"))
        if abs(tme - (tes + adj)) > 0.15:
            raise RuntimeError(f"{y}: TME {tme} != services {tes} + "
                               f"adjustments {adj}")


def main() -> None:
    year, base_path = latest_publication()
    print(f"PESA {year} (HM Treasury) — {base_path}")
    years, cols = extract(workbook(year, base_path))
    verify(years, cols)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", newline="", encoding="utf-8") as fh:
        fh.write(
            "# Public sector expenditure by function and Total Managed\n"
            "# Expenditure, £ billion, nominal. HM Treasury, Public\n"
            f"# Expenditure Statistical Analyses (PESA) {year}, Table 4.2.\n"
            "# FINANCIAL YEARS, April to March: fiscal_year is the STARTING\n"
            "# year (2020 = 2020-21) — never plot against a calendar-year\n"
            "# series without saying so.\n"
            "# Functions are UN COFOG level 1 as assigned by the Treasury,\n"
            "# NOT Eurostat's COFOG transmission (which the UK no longer\n"
            "# makes). sub_* columns are the published \"of which\"\n"
            "# sub-function detail: components of their function, never\n"
            "# additive across. EU transactions keep the published sign\n"
            "# (negative was a net receipt from the EU). TME excludes the\n"
            "# temporary effects of banks classified to the public sector.\n"
            f"# Source: https://www.gov.uk{base_path} (attachment resolved\n"
            "# via the gov.uk content API; do not hardcode media URLs).\n")
        w = csv.writer(fh)
        w.writerow(["fiscal_year"] + ORDER)
        for k, y in enumerate(years):
            w.writerow([y] + [cols[n][k] for n in ORDER])
    print(f"  data/fiscal/spending_by_function_from_hmt.csv  "
          f"{len(years)} rows ({years[0]}–{years[-1]})")
    print("Done. Diff the sidecar before adopting.")


if __name__ == "__main__":
    main()
