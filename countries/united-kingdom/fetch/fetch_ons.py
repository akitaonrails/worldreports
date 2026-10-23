#!/usr/bin/env python3
"""Refresh the Office for National Statistics series for the United Kingdom.

Usage:  python3 -I countries/united-kingdom/fetch/fetch_ons.py

The ONS is the workhorse of this country: national accounts, prices,
the labour market, public sector finances, the balance of payments,
population, vital statistics, life tables, causes of death, income
inequality and migration all come from it, over two machine channels.

**The timeseries generator** —
`/generator?format=csv&uri=/<theme>/timeseries/<cdid>/<dataset>` — serves
a header block (Title, CDID, Source dataset ID, PreUnit, Unit, Release
date, Next release, Important notes) and then one row per period, with
annual (`2024`), quarterly (`2024 Q1`) and monthly (`2024 JAN`) rows
stacked in one file. The annual row is the ONS's own figure: nothing
here sums quarters into years except where the manifest says so (the
financial-year PSF totals, built from the monthly rows). The theme path
matters: a CDID asked for under the wrong theme answers a 404 page.

  YBHA ABMI IHXW   GDP, current prices / chained volume / per head (UKEA)
  L55O D7G7 L522   CPIH and CPI annual rates, CPIH index (MM23)
  MGSX LF24 KAB9   LFS unemployment and employment rates, AWE (LMS)
  J5II JW2O KX5Q   PSNB ex, current receipts ex, TME ex (PUSF)
  HF6W HF6X        PSND ex, £bn and % of GDP (PUSF)
  IKBH IKBI        exports and imports of goods and services, BoP (MRET)
  HBOP             current account balance (PNBP)
  CVZH CVZI        household debt-to-income ratio, annual / quarterly
  NIWK HABN        S.14 loans liability and gross disposable income
  NNRE QWND        the same two for S.14 + S.15 — only to prove which
                   sector CVZH/CVZI cover

**Bulletin dataset workbooks** — per-edition xlsx files whose paths are
resolved from each dataset's landing page on every run. The `/current/`
alias 404s for some datasets (mid-year population, LTIM), so it is never
written into this file; where a landing page itself links `/current/`
(life tables, 21st Century Mortality) that link is what gets used.

Things these sources will do to you, verified 2026-10-22:

**Cloudflare error 1015 under rapid fire.** ons.gov.uk answered HTTP 429
"error code: 1015" to requests 1.2 s apart. Requests here are spaced
3 s and a 429 backs off 30 s, 60 s, 90 s.

**JW2V is not current receipts.** The research pass recorded JW2V as
"current receipts"; the generator's own title says it is "PS: Current
Budget Deficit (exc PS Banks) as a % of GDP" and its rows stop in 2024.
Receipts ex public sector banks are JW2O and total managed expenditure
ex banks is KX5Q (titles read from the PUSF dataset); the fetcher
asserts the identity J5II = JW2O − KX5Q on every month, so TME is the
published series, not a computed column.

**HF6W is £ billions**, not £ millions (its title says so); the column
name carries the unit.

**CVZH and CVZI are the same S.14 ratio at two frequencies.** Both
carry the title "Households (S.14): Debt to Income ratio: percentage:
NSA" in UKEA; the manifest feared one was S.14 + S.15. It is not: CVZH
has only annual rows, CVZI only quarterly rows, every CVZH year equals
that year's CVZI Q4, and both reproduce S.14 loans ÷ S.14 gross
disposable income (NIWK ÷ HABN) to rounding, while the households-plus-
NPISH quotient (NNRE ÷ QWND) sits about 3 points lower. Asserted below.

**PSNB financial years move with revisions.** Calendar 2020 is
−£269,927m (asserted to the pound-million); FY 2020-21 summed from the
current vintage's months is about −£311bn, not the ≈£322bn the
discovery note carried from an earlier vintage. Asserted as a range.

**The birth summary tables stop at 2022.** Their successor dataset,
"Births in England and Wales: birth registrations", carries a 1938–
latest time series in Table 6; that is the source used. Some of its
years store counts as text with thousands separators ("594,677").

**Series DR has no all-ages homicide row.** Its Table 4 lists
"Homicide and probable homicide" only where it ranks among the leading
causes of an age band. The all-ages ICD-10 X85–Y09 count comes from the
21st Century Mortality dataset (deaths by 4-character ICD-10 code,
2001–2022, released December 2023); the years after 2022 are not
published at that grain yet and stay absent, never estimated.

**Causes of death: two tables, one quantity.** DR 2025 Table 5 carries
ICD-10 chapters for 2016–2025; 21st Century Mortality, aggregated to the
same chapters, carries 2001–2022. On the 2016–2022 overlap every chapter
matches DR to the death (asserted), so 2006–2015 come from the latter
and 2016 onward from DR, with a `source` column naming which.

Outputs are `_from_ons` sidecars under data/. Nothing here overwrites a
curated file.
"""
from __future__ import annotations

import csv
import io
import re
import sys
import time
import urllib.error
import urllib.request
import zipfile
from collections import Counter
from pathlib import Path

try:
    import openpyxl
except ImportError:  # the bulletin workbooks cannot be read without it
    sys.exit("openpyxl is required: pip install -r requirements-fetch.txt")

import warnings
# ONS workbooks carry no default style; openpyxl warns on every open.
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
CACHE = ROOT / "cache"

ONS = "https://www.ons.gov.uk"
UA = {"User-Agent": "worldreports-fetch"}
GAP = 3.0                 # seconds between ONS requests (Cloudflare 1015)
Y0 = 2006                 # the atlas window opens here
LAST_FULL = 2025          # country.yaml last_full_year

MONTHS = "JAN FEB MAR APR MAY JUN JUL AUG SEP OCT NOV DEC".split()
MON_NUM = {m.title(): i + 1 for i, m in enumerate(MONTHS)}

# CDID → theme path and dataset, as the generator wants them.
SERIES = {
    "YBHA": "economy/grossdomesticproductgdp/timeseries/ybha/ukea",
    "ABMI": "economy/grossdomesticproductgdp/timeseries/abmi/ukea",
    "IHXW": "economy/grossdomesticproductgdp/timeseries/ihxw/ukea",
    "L55O": "economy/inflationandpriceindices/timeseries/l55o/mm23",
    "D7G7": "economy/inflationandpriceindices/timeseries/d7g7/mm23",
    "L522": "economy/inflationandpriceindices/timeseries/l522/mm23",
    "MGSX": "employmentandlabourmarket/peoplenotinwork/unemployment/"
            "timeseries/mgsx/lms",
    "LF24": "employmentandlabourmarket/peopleinwork/employmentandemployeetypes/"
            "timeseries/lf24/lms",
    "KAB9": "employmentandlabourmarket/peopleinwork/earningsandworkinghours/"
            "timeseries/kab9/lms",
    "J5II": "economy/governmentpublicsectorandtaxes/publicsectorfinance/"
            "timeseries/j5ii/pusf",
    "JW2O": "economy/governmentpublicsectorandtaxes/publicsectorfinance/"
            "timeseries/jw2o/pusf",
    "KX5Q": "economy/governmentpublicsectorandtaxes/publicsectorfinance/"
            "timeseries/kx5q/pusf",
    "HF6W": "economy/governmentpublicsectorandtaxes/publicsectorfinance/"
            "timeseries/hf6w/pusf",
    "HF6X": "economy/governmentpublicsectorandtaxes/publicsectorfinance/"
            "timeseries/hf6x/pusf",
    "IKBH": "economy/nationalaccounts/balanceofpayments/timeseries/ikbh/mret",
    "IKBI": "economy/nationalaccounts/balanceofpayments/timeseries/ikbi/mret",
    "HBOP": "economy/nationalaccounts/balanceofpayments/timeseries/hbop/pnbp",
    "CVZH": "economy/grossdomesticproductgdp/timeseries/cvzh/ukea",
    "CVZI": "economy/grossdomesticproductgdp/timeseries/cvzi/ukea",
    "NIWK": "economy/grossdomesticproductgdp/timeseries/niwk/ukea",
    "HABN": "economy/grossdomesticproductgdp/timeseries/habn/ukea",
    "NNRE": "economy/nationalaccounts/uksectoraccounts/timeseries/nnre/ukea",
    "QWND": "economy/grossdomesticproductgdp/timeseries/qwnd/ukea",
}

POP = "peoplepopulationandcommunity"
PAGES = {
    "inequality": f"{POP}/personalandhouseholdfinances/incomeandwealth/"
                  "datasets/householddisposableincomeandinequality",
    "population": f"{POP}/populationandmigration/populationestimates/datasets/"
                  "populationestimatesforukenglandandwalesscotlandandnorthernireland",
    "births": f"{POP}/birthsdeathsandmarriages/livebirths/datasets/"
              "birthsinenglandandwalesbirthregistrations",
    "deaths": f"{POP}/birthsdeathsandmarriages/deaths/datasets/"
              "deathsregisteredinenglandandwalesseriesdrreferencetables",
    "c21": f"{POP}/birthsdeathsandmarriages/deaths/datasets/"
           "the21stcenturymortalityfilesdeathsdataset",
    "lifetables": f"{POP}/birthsdeathsandmarriages/lifeexpectancies/datasets/"
                  "nationallifetablesunitedkingdomreferencetables",
    "ltim": f"{POP}/populationandmigration/internationalmigration/datasets/"
            "longterminternationalimmigrationemigrationandnetmigrationflowsprovisional",
}


# ---------------------------------------------------------------- transport

_last = [0.0]


def _get(url: str) -> bytes:
    """One ONS request, spaced and retried. Raises on persistent failure."""
    last = None
    for attempt in range(4):
        wait = _last[0] + GAP - time.monotonic()
        if wait > 0:
            time.sleep(wait)
        try:
            req = urllib.request.Request(url, headers=UA)
            with urllib.request.urlopen(req, timeout=300) as r:
                blob = r.read()
            _last[0] = time.monotonic()
            return blob
        except urllib.error.HTTPError as e:
            _last[0] = time.monotonic()
            last = e
            if e.code == 404:
                break  # a wrong path does not heal on retry
            time.sleep(30 * (attempt + 1) if e.code == 429 else 3)
        except OSError as e:
            _last[0] = time.monotonic()
            last = e
            time.sleep(3 * (attempt + 1))
    raise RuntimeError(f"{url}: {last}")


def fetch(url: str, name: str, check) -> bytes:
    """Cached download; `check(blob)` raises before anything is cached."""
    CACHE.mkdir(exist_ok=True)
    local = CACHE / name
    if local.exists():
        return local.read_bytes()
    print(f"  downloading {url}")
    blob = _get(url)
    check(blob)
    local.write_bytes(blob)
    return blob


def is_xlsx(blob: bytes) -> None:
    # An error page under HTTP 200 is not a zip; a truncated one fails testzip.
    if blob[:2] != b"PK":
        raise RuntimeError("payload is not an xlsx (error page?)")
    bad = zipfile.ZipFile(io.BytesIO(blob)).testzip()
    if bad:
        raise RuntimeError(f"xlsx member {bad} is corrupt")


def is_page(blob: bytes) -> None:
    if b"/file?uri=" not in blob:
        raise RuntimeError("dataset page lists no files (wall or error page?)")


# ---------------------------------------------------------------- generator

def series(cdid: str) -> dict:
    """{'title', 'unit', 'release', 'A': {y: v}, 'Q': {(y, q): v},
    'M': {(y, m): v}} for one CDID from the generator."""
    url = f"{ONS}/generator?format=csv&uri=/{SERIES[cdid]}"

    def check(blob: bytes) -> None:
        text = blob.decode("utf-8-sig", "replace")
        rows = list(csv.reader(io.StringIO(text)))
        if not rows or rows[0][:1] != ["Title"]:
            raise RuntimeError(f"{cdid}: generator payload has no header "
                               "block (error page?)")
        got = {r[0]: r[1] for r in rows[:8] if len(r) > 1}
        if got.get("CDID") != cdid:
            raise RuntimeError(f"{cdid}: generator answered CDID "
                               f"{got.get('CDID')!r}")
        if sum(1 for r in rows[8:] if r and r[0][:1].isdigit()) < 20:
            raise RuntimeError(f"{cdid}: fewer than 20 period rows")

    blob = fetch(url, f"ons_{cdid.lower()}.csv", check)
    rows = list(csv.reader(io.StringIO(blob.decode("utf-8-sig"))))
    head = {r[0]: (r[1] if len(r) > 1 else "") for r in rows[:8]}
    out = {"title": head["Title"], "unit": head.get("Unit", ""),
           "release": head.get("Release date", ""),
           "A": {}, "Q": {}, "M": {}}
    for r in rows:
        if len(r) < 2 or not r[0][:1].isdigit() or r[1].strip() == "":
            continue
        p, v = r[0].strip(), float(r[1])
        if re.fullmatch(r"\d{4}", p):
            out["A"][int(p)] = v
        elif m := re.fullmatch(r"(\d{4}) Q([1-4])", p):
            out["Q"][(int(m[1]), int(m[2]))] = v
        elif m := re.fullmatch(r"(\d{4}) ([A-Z]{3})", p):
            out["M"][(int(m[1]), MONTHS.index(m[2]) + 1)] = v
        else:
            raise RuntimeError(f"{cdid}: unrecognised period {p!r}")
    return out


def need_title(s: dict, cdid: str, *words: str) -> None:
    """The title is the series' identity card; a re-pointed CDID fails here."""
    for w in words:
        if w.lower() not in s["title"].lower():
            raise RuntimeError(f"{cdid}: title lost {w!r}: {s['title']!r}")


def full_years(s: dict, cdid: str, first: int = Y0,
               last: int = LAST_FULL) -> None:
    gap = [y for y in range(first, last + 1) if y not in s["A"]]
    if gap:
        raise RuntimeError(f"{cdid}: annual rows missing for {gap}")


def fy_sum(s: dict, y: int) -> float | None:
    """April y to March y+1 summed from monthly flows; None if incomplete."""
    months = [(y, m) for m in range(4, 13)] + [(y + 1, m) for m in (1, 2, 3)]
    if not all(k in s["M"] for k in months):
        return None
    return sum(s["M"][k] for k in months)


def fy_gdp(gdp: dict, y: int) -> float | None:
    """Nominal GDP for financial year y: Q2..Q4 of y plus Q1 of y+1."""
    qs = [(y, 2), (y, 3), (y, 4), (y + 1, 1)]
    if not all(q in gdp["Q"] for q in qs):
        return None
    return sum(gdp["Q"][q] for q in qs)


# ---------------------------------------------------------------- workbooks

def resolve(key: str, pattern: str, rank=None) -> str:
    """The workbook link on a dataset landing page, matched by `pattern`.

    `rank(match)` orders candidate editions; the highest wins. Without
    it the page must list exactly one candidate.
    """
    page = fetch(f"{ONS}/{PAGES[key]}", f"ons_page_{key}.html", is_page)
    links = re.findall(r'href="(/file\?uri=[^"]+)"', page.decode("utf-8",
                                                                  "replace"))
    found = [(m, l) for l in links if (m := re.search(pattern, l))]
    if not found:
        raise RuntimeError(f"{key}: no link matching {pattern!r} on the page")
    if rank is None:
        if len({l for _, l in found}) != 1:
            raise RuntimeError(f"{key}: {len(found)} candidates for "
                               f"{pattern!r}; pick rule needed")
        return ONS + found[0][1]
    return ONS + max(found, key=lambda ml: rank(ml[0]))[1]


def workbook(key: str, pattern: str, rank=None):
    url = resolve(key, pattern, rank)
    name = "ons_" + url.rsplit("/", 1)[-1]
    return openpyxl.load_workbook(io.BytesIO(fetch(url, name, is_xlsx)),
                                  read_only=True, data_only=True), url


def grid(wb, sheet: str) -> list[tuple]:
    if sheet not in wb.sheetnames:
        raise RuntimeError(f"sheet {sheet!r} not in workbook")
    ws = wb[sheet]
    ws.reset_dimensions()  # ONS files under-declare their used range
    return [r for r in ws.iter_rows(values_only=True)]


def num(v) -> float | None:
    """A published cell as a number: thousands separators stripped,
    ONS shorthand ('[x]', '-', 'u', ':') treated as no value."""
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        t = v.replace(",", "").strip()
        if re.fullmatch(r"-?\d+(\.\d+)?", t):
            return float(t)
    return None


def header_index(rows: list[tuple], first: str) -> int:
    for i, r in enumerate(rows):
        if r and isinstance(r[0], str) and r[0].strip().startswith(first):
            return i
    raise RuntimeError(f"no header row starting {first!r}")


# ---------------------------------------------------------------- writers

def write(rel: str, comment: str, header: list[str], rows: list[list]) -> None:
    path = DATA / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        for line in comment.strip().splitlines():
            fh.write(f"# {line}".rstrip() + "\n")
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(rows)
    span = f"{rows[0][0]}–{rows[-1][0]}" if rows else "EMPTY"
    print(f"  data/{rel}  {len(rows)} rows ({span})")


def fmt(v, nd: int = 1):
    if v is None:
        return ""
    return int(round(v)) if nd == 0 else round(v, nd)


# ---------------------------------------------------------------- macro

def macro() -> dict:
    ybha, abmi, ihxw = series("YBHA"), series("ABMI"), series("IHXW")
    need_title(ybha, "YBHA", "Gross Domestic Product", "Current price", "£m")
    need_title(abmi, "ABMI", "chained volume", "£m")
    need_title(ihxw, "IHXW", "per head", "CVM")
    for s, c in ((ybha, "YBHA"), (abmi, "ABMI"), (ihxw, "IHXW")):
        full_years(s, c)
    # Must-verify spot value: nominal GDP 2024 = £2,905.0bn.
    # £2,905.0bn to one decimal, half-up: 2,904,950 <= v < 2,905,050 £m.
    if not 2_904_950 <= ybha["A"][2024] < 2_905_050:
        raise RuntimeError(f"YBHA 2024 = {ybha['A'][2024]} £m, "
                           "expected £2,905.0bn")
    # The annual row is the ONS's own; for current prices it must equal
    # the sum of the quarters (a level series), as a shape check.
    if abs(sum(ybha["Q"][(2024, q)] for q in range(1, 5))
           - ybha["A"][2024]) > 2:
        raise RuntimeError("YBHA 2024 annual ≠ sum of quarters")
    ref = re.search(r"reference year = (\d{4})", ihxw["title"])
    years = sorted(y for y in ybha["A"] if y >= Y0)
    write("macro/gdp_from_ons.csv", f"""
ONS UK Economic Accounts (UKEA), annual rows of the generator, release
{ybha['release']}. gdp_nominal = YBHA, GDP at market prices, current
prices, SA, £ million. gdp_real_cvm = ABMI, chained volume measure, SA,
£ million (annual CVM is its own chain-linked figure, never re-summed).
gdp_per_head_cvm = IHXW, GDP per head, CVM market prices, £ (reference
year {ref[1] if ref else '?'}). Verified: 2024 nominal = £2,905.0bn.
""", ["year", "gdp_nominal_gbp_mn", "gdp_real_cvm_gbp_mn",
      "gdp_per_head_cvm_gbp"],
          [[y, fmt(ybha["A"][y], 0), fmt(abmi["A"].get(y), 0),
            fmt(ihxw["A"].get(y), 0)] for y in years])

    cpih, cpi, idx = series("L55O"), series("D7G7"), series("L522")
    need_title(cpih, "L55O", "CPIH ANNUAL RATE")
    need_title(cpi, "D7G7", "CPI ANNUAL RATE")
    need_title(idx, "L522", "CPIH INDEX")
    for s, c in ((cpih, "L55O"), (cpi, "D7G7"), (idx, "L522")):
        full_years(s, c)  # must-verify: both present for the full window
    years = sorted(y for y in cpih["A"] if y >= Y0)
    write("macro/prices_from_ons.csv", f"""
ONS consumer price inflation (MM23), annual rows of the generator,
release {cpih['release']}. cpih = L55O, CPIH 12-month rate, all items,
% (the headline since March 2017; includes owner-occupiers' housing
costs). cpi = D7G7, CPI 12-month rate, all items, %. Two columns, never
one. The annual row is the ONS's annual average rate.
""", ["year", "cpih_annual_rate_pct", "cpi_annual_rate_pct"],
          [[y, fmt(cpih["A"][y]), fmt(cpi["A"].get(y))] for y in years])

    unemp, emp, awe = series("MGSX"), series("LF24"), series("KAB9")
    need_title(unemp, "MGSX", "Unemployment rate", "16 and over")
    need_title(emp, "LF24", "Employment rate", "16 to 64")
    need_title(awe, "KAB9", "AWE", "Whole Economy", "Total Pay")
    for s, c in ((unemp, "MGSX"), (emp, "LF24"), (awe, "KAB9")):
        full_years(s, c)
        # The generator's annual row is the ONS's annual figure; it must
        # sit on the mean of that year's monthly rows (rolling three-month
        # LFS readings), or the series has been re-pointed.
        for y in range(Y0, LAST_FULL + 1):
            mean = sum(s["M"][(y, m)] for m in range(1, 13)) / 12
            tol = 2.0 if c == "KAB9" else 0.15
            if abs(mean - s["A"][y]) > tol:
                raise RuntimeError(f"{c} {y}: annual {s['A'][y]} vs "
                                   f"monthly mean {mean:.2f}")
    years = sorted(y for y in unemp["A"] if y >= Y0)
    write("macro/labour_from_ons.csv", f"""
ONS Labour Force Survey (LMS), annual rows of the generator (the mean of
the year), release {unemp['release']}. unemployment = MGSX, rate aged
16+, SA, %. employment = LF24, rate aged 16-64, SA, %. LFS QUALITY BREAK:
response rates collapsed through 2023-24; ONS suspended and reweighted
outputs in 2024 and the survey is moving to the Transformed LFS -
annotate 2023 onward, do not smooth.
""", ["year", "unemployment_rate_pct", "employment_rate_16_64_pct"],
          [[y, fmt(unemp["A"][y]), fmt(emp["A"].get(y))] for y in years])

    base = max(y for y in idx["A"] if y <= LAST_FULL)
    years = sorted(y for y in awe["A"] if y >= Y0)
    write("macro/earnings_from_ons.csv", f"""
ONS Average Weekly Earnings (LMS), annual rows of the generator, release
{awe['release']}. awe_nominal = KAB9, whole economy, total pay excluding
arrears, SA, £ per week. awe_real_{base}prices is COMPUTED here: KAB9 x
CPIH index (L522) in {base} / CPIH index in that year - label it as
computed on the plate. Same LFS-adjacent caveats for 2023-24.
""", ["year", "awe_nominal_gbp_week", f"awe_real_{base}prices_gbp_week"],
          [[y, fmt(awe["A"][y], 0),
            fmt(awe["A"][y] * idx["A"][base] / idx["A"][y], 0)]
           for y in years])
    return {"ybha": ybha}


# ---------------------------------------------------------------- fiscal

def fiscal(ybha: dict) -> None:
    psnb, rec, tme = series("J5II"), series("JW2O"), series("KX5Q")
    need_title(psnb, "J5II", "Net Borrowing", "excluding public sector banks")
    need_title(rec, "JW2O", "exc PS Banks", "Total current receipts")
    need_title(tme, "KX5Q", "exc PS Banks", "Total managed expenditure")
    # Must-verify spot value: calendar 2020 PSNB ex = −£269.9bn.
    if psnb["A"].get(2020) != -269927:
        raise RuntimeError(f"J5II 2020 = {psnb['A'].get(2020)}, "
                           "expected −269,927 £m")
    # Sign: negative = borrowing. A flipped sign fails here, not on a chart.
    if not all(psnb["A"][y] < 0 for y in range(2008, LAST_FULL + 1)):
        raise RuntimeError("J5II: a year 2008-2025 is not negative — "
                           "sign convention changed?")
    # PSNB is exactly receipts minus TME in the PUSF accounts: assert on
    # every month the three share, so TME below is the published series.
    for k in psnb["M"]:
        if k in rec["M"] and k in tme["M"]:
            if abs(rec["M"][k] - tme["M"][k] - psnb["M"][k]) > 1:
                raise RuntimeError(f"PUSF identity broken at {k}")
    for s, c in ((psnb, "J5II"), (rec, "JW2O"), (tme, "KX5Q")):
        full_years(s, c)
        for y in range(Y0, LAST_FULL + 1):
            if abs(sum(s["M"][(y, m)] for m in range(1, 13))
                   - s["A"][y]) > 2:
                raise RuntimeError(f"{c} {y}: annual ≠ sum of months")
    fy2020 = fy_sum(psnb, 2020)
    if not -340_000 < fy2020 < -290_000:
        raise RuntimeError(f"J5II FY2020-21 = {fy2020}: outside the "
                           "published ~£300-330bn band")

    years = sorted(y for y in psnb["A"] if y >= Y0)
    write("fiscal/psnb_from_ons.csv", f"""
ONS public sector finances (PUSF), release {psnb['release']}. J5II, public
sector net borrowing excluding public sector banks, £ million, current
prices, NSA. SIGN: NEGATIVE = BORROWING (as published). Two keys on one
row: psnb_calendar is the generator's calendar-year row for `year`;
psnb_fiscal is the financial year STARTING in April of `year` (2020 =
FY 2020-21), summed here from the twelve monthly rows and blank until
all twelve are published. Verified: calendar 2020 = -269,927; FY 2020-21
in this vintage = {fmt(fy2020, 0)}.
""", ["year", "psnb_ex_calendar_gbp_mn", "psnb_ex_fiscal_gbp_mn"],
          [[y, fmt(psnb["A"][y], 0), fmt(fy_sum(psnb, y), 0)]
           for y in years])

    rows = []
    for y in years:
        g_c, g_f = ybha["A"].get(y), fy_gdp(ybha, y)
        r_f, t_f = fy_sum(rec, y), fy_sum(tme, y)
        rows.append([y, fmt(rec["A"][y], 0), fmt(tme["A"][y], 0),
                     fmt(100 * rec["A"][y] / g_c) if g_c else "",
                     fmt(100 * tme["A"][y] / g_c) if g_c else "",
                     fmt(r_f, 0), fmt(t_f, 0),
                     fmt(100 * r_f / g_f) if r_f and g_f else "",
                     fmt(100 * t_f / g_f) if t_f and g_f else ""])
    write("fiscal/revenue_spending_from_ons.csv", f"""
ONS public sector finances (PUSF), release {rec['release']}. receipts =
JW2O, public sector excluding PS banks, total current receipts; tme =
KX5Q, total managed expenditure excluding PS banks; both £ million,
current prices, NSA. (JW2V, named in the research pass, is the current
budget deficit as % of GDP - not receipts.) Identity asserted on every
month: J5II = JW2O - KX5Q. *_calendar columns are calendar years (the
generator's annual rows); *_fiscal columns are the financial year
STARTING in April of `year`, summed from monthly rows. *_pct_gdp columns
are COMPUTED against YBHA nominal GDP: calendar-year YBHA, or for fiscal
years the sum of YBHA quarters Q2..Q4 of `year` and Q1 of year+1.
""", ["year", "receipts_calendar_gbp_mn", "tme_calendar_gbp_mn",
      "receipts_calendar_pct_gdp", "tme_calendar_pct_gdp",
      "receipts_fiscal_gbp_mn", "tme_fiscal_gbp_mn",
      "receipts_fiscal_pct_gdp", "tme_fiscal_pct_gdp"], rows)

    debt, ratio = series("HF6W"), series("HF6X")
    need_title(debt, "HF6W", "Net Debt", "excluding public sector banks", "£bn")
    need_title(ratio, "HF6X", "Net Debt", "% of GDP")
    rows = []
    for y in range(Y0, LAST_FULL + 1):
        dec = (y, 12)
        mar = (y + 1, 3)
        if dec not in debt["M"]:
            raise RuntimeError(f"HF6W: no December {y} reading")
        rows.append([y, fmt(debt["M"][dec]), fmt(ratio["M"].get(dec)),
                     fmt(debt["M"].get(mar)), fmt(ratio["M"].get(mar))])
    # The generator's annual row for a stock must be the December level.
    if abs(debt["A"][2024] - debt["M"][(2024, 12)]) > 0.05:
        raise RuntimeError("HF6W annual row is not the December stock")
    write("fiscal/psnd_from_ons.csv", f"""
ONS public sector finances (PUSF), release {debt['release']}. PSND ex =
public sector net debt excluding public sector banks: HF6W in £ BILLION
and HF6X as % of GDP, NSA. MONTH-END STOCKS, never summed. *_dec columns
are the December reading of `year` (= the generator's annual row);
*_fy_end columns are the end-March reading closing the financial year
that STARTS in `year` (2024 = end of FY 2024-25, March 2025), blank
until published. PSND ex is not Maastricht gross debt; the harmonised
comparison is the IMF series in international/.
""", ["year", "psnd_ex_dec_gbp_bn", "psnd_ex_dec_pct_gdp",
      "psnd_ex_fy_end_gbp_bn", "psnd_ex_fy_end_pct_gdp"], rows)


# ---------------------------------------------------------------- trade

def trade() -> None:
    x, m, ca = series("IKBH"), series("IKBI"), series("HBOP")
    need_title(x, "IKBH", "Total Trade", "Exports", "BOP")
    need_title(m, "IKBI", "Total Trade", "Imports", "BOP")
    need_title(ca, "HBOP", "Current Account Balance")
    for s, c in ((x, "IKBH"), (m, "IKBI"), (ca, "HBOP")):
        full_years(s, c)
    # Magnitude: UK goods + services exports run around £0.8-1.0tn.
    if not 700_000 < x["A"][2024] < 1_100_000:
        raise RuntimeError(f"IKBH 2024 = {x['A'][2024]}: implausible")
    years = sorted(y for y in x["A"] if y >= Y0)
    write("trade/trade_totals_from_ons.csv", f"""
ONS UK trade (MRET) and balance of payments (PNBP), annual rows of the
generator, release {x['release']}. exports = IKBH, imports = IKBI: total
trade in goods AND services, world, balance-of-payments basis, current
prices, SA, £ million. current_account = HBOP, BoP current account
balance, SA, £ million. NOT the customs-value goods-only figures in
goods_by_partner_from_uktradeinfo.csv; never spliced.
""", ["year", "exports_gs_gbp_mn", "imports_gs_gbp_mn",
      "current_account_gbp_mn"],
          [[y, fmt(x["A"][y], 0), fmt(m["A"].get(y), 0),
            fmt(ca["A"].get(y), 0)] for y in years])


# ---------------------------------------------------------------- credit

def credit() -> None:
    a, q = series("CVZH"), series("CVZI")
    loans, gdi = series("NIWK"), series("HABN")
    loans_n, gdi_n = series("NNRE"), series("QWND")
    for s, c in ((a, "CVZH"), (q, "CVZI")):
        need_title(s, c, "Households (S.14)", "Debt to Income")
    need_title(loans, "NIWK", "Households (S.14)", "Loans", "Liability")
    need_title(gdi, "HABN", "Households (S.14)", "Disposable income, gross")
    need_title(loans_n, "NNRE", "S.14 + S.15", "Loans", "Liability")
    need_title(gdi_n, "QWND", "S.14 + S.15", "Disposable income, gross")
    # Disambiguation, asserted: CVZH is the annual frequency only, CVZI
    # the quarterly only, and they agree at every Q4 ...
    if a["Q"] or a["M"] or q["A"] or q["M"]:
        raise RuntimeError("CVZH/CVZI frequencies changed: re-disambiguate")
    for y, v in a["A"].items():
        if q["Q"].get((y, 4)) != v:
            raise RuntimeError(f"CVZH {y} = {v} ≠ CVZI {y} Q4 "
                               f"= {q['Q'].get((y, 4))}")
    # ... and both are S.14 alone: loans (year-end stock) / GDI (flow).
    for y in range(Y0, LAST_FULL + 1):
        s14 = 100 * loans["A"][y] / gdi["A"][y]
        s1415 = 100 * loans_n["A"][y] / gdi_n["A"][y]
        if abs(s14 - a["A"][y]) > 0.1:
            raise RuntimeError(f"CVZH {y} = {a['A'][y]} is not NIWK/HABN "
                               f"= {s14:.2f}")
        if abs(s1415 - a["A"][y]) < 1.0:
            raise RuntimeError(f"CVZH {y} is indistinguishable from the "
                               f"S.14+S.15 quotient {s1415:.2f}")
    full_years(a, "CVZH")
    years = sorted(y for y in a["A"] if y >= Y0)
    write("credit/household_debt_from_ons.csv", f"""
ONS UK Economic Accounts (UKEA), release {a['release']}. Household
debt-to-income ratio, %, NSA: households sector S.14 ONLY (NOT S.14 +
S.15 NPISH). CVZH and CVZI share the title "Households (S.14): Debt to
Income ratio" and are the SAME ratio at two frequencies: CVZH = annual
rows only, CVZI = quarterly rows only; every CVZH year equals CVZI Q4,
and both reproduce NIWK (S.14 loans liability, stock) / HABN (S.14 gross
disposable income) within 0.1 pt, while NNRE/QWND (S.14+S.15) runs ~3
pt lower - all asserted by the fetcher. Annual reading = Q4 (a stock
over the year's income), labelled as such.
""", ["year", "hh_debt_to_income_pct_q4"],
          [[y, fmt(a["A"][y])] for y in years])


# ---------------------------------------------------------------- social

def social() -> None:
    wb, url = workbook(
        "inequality", r"/financialyearending(\d{4})/[^/]+\.xlsx",
        rank=lambda m: int(m[1]))
    rows = grid(wb, "Table 10")
    if "Gini" not in str(rows[0][0]) or "UK" not in str(rows[0][0]):
        raise RuntimeError(f"Table 10 title changed: {rows[0][0]!r}")
    h = header_index(rows, "Year")
    head = [str(c).strip() if c else "" for c in rows[h]]
    want = ["S80/S20", "P90/P10", "Palma", "Gini", "Top 1%"]
    if head[1:6] != want:
        raise RuntimeError(f"Table 10 columns changed: {head}")
    out = []
    for r in rows[h + 1:]:
        lab = str(r[0]).strip() if r and r[0] is not None else ""
        m = re.fullmatch(r"(\d{4})/(\d{2})", lab)
        if not m:
            if re.fullmatch(r"\d{4}", lab):
                continue  # 1977-1993 are calendar years; left in the book
            if out:
                break
            continue
        start = int(m[1])
        if start < Y0:
            continue
        if (start + 1) % 100 != int(m[2]):
            raise RuntimeError(f"Table 10: odd year label {lab!r}")
        vals = [num(c) for c in r[1:6]]
        if vals[3] is None or not 20 < vals[3] < 50:
            raise RuntimeError(f"Table 10 {lab}: implausible Gini {vals[3]}")
        out.append([start] + [fmt(v) for v in vals])
    if out[-1][0] < LAST_FULL - 1:
        raise RuntimeError(f"inequality edition ends at FY {out[-1][0]}")
    write("social/gini_from_ons.csv", f"""
ONS, The effects of taxes and benefits on household income ({url.rsplit('/', 1)[-1]}),
Table 10, UK. Gini coefficient of equivalised disposable household
income (%, 0-100 scale), with the S80/S20, P90/P10 and Palma ratios and
the top 1% income share (from 2001/02 only). Incomes from 2001/02 are
adjusted for top-earner under-coverage (release note 11). FINANCIAL
YEARS (April-March): fiscal_year is the STARTING year, 2024 = 2024/25 =
FYE 2025. Rows before the atlas window (incl. 1977-1993 calendar years)
are left in the workbook.
""", ["fiscal_year", "s80_s20_ratio", "p90_p10_ratio", "palma_ratio",
      "gini_pct", "top1_share_pct"], out)


# ---------------------------------------------------------------- demography

def demography() -> None:
    wb, url = workbook(
        "population", r"/ukpopulationestimates1838to(\d{4})/[^/]+\.xlsx",
        rank=lambda m: int(m[1]))
    tables = {"uk": ("Table 1", "UK"), "england": ("Table 10", "England"),
              "wales": ("Table 12", "Wales"),
              "scotland": ("Table 14", "Scotland"),
              "northern_ireland": ("Table 17", "Northern Ireland")}
    pop = {}
    for key, (sheet, place) in tables.items():
        rows = grid(wb, sheet)
        if place not in str(rows[0][0]):
            raise RuntimeError(f"{sheet}: title lost {place!r}")
        h = header_index(rows, "Year")
        col = 1  # 'Total Population' (Table 1) or 'Persons'
        if str(rows[h][col]).strip() not in ("Total Population", "Persons"):
            raise RuntimeError(f"{sheet}: column 2 is {rows[h][col]!r}")
        pop[key] = {}
        for r in rows[h + 1:]:
            m = re.fullmatch(r"Mid-(\d{4})", str(r[0]).strip()) if r else None
            if m and num(r[col]) is not None:
                pop[key][int(m[1])] = num(r[col])
    years = sorted(y for y in pop["uk"] if y >= Y0)
    if years[-1] < LAST_FULL:
        raise RuntimeError(f"population ends mid-{years[-1]}")
    for y in years:
        parts = sum(pop[k][y] for k in tables if k != "uk")
        if abs(parts - pop["uk"][y]) > 5:
            raise RuntimeError(f"mid-{y}: nations sum {parts} ≠ UK "
                               f"{pop['uk'][y]}")
    write("demography/population_from_ons.csv", f"""
ONS mid-year population estimates ({url.rsplit('/', 1)[-1]}), Tables 1, 10,
12, 14, 17: usual residents at 30 JUNE (mid-year), persons. Not the
Eurostat 1-January stock. One vintage throughout (this edition's
series, census-rebased by the ONS itself); the four nations are asserted
to sum to the UK figure.
""", ["year", "population_uk", "population_england", "population_wales",
      "population_scotland", "population_northern_ireland"],
          [[y] + [fmt(pop[k][y], 0) for k in tables] for y in years])

    wbb, burl = workbook("births", r"/(\d{4})/\d{4}birthregistrations\.xlsx",
                         rank=lambda m: int(m[1]))
    rows = grid(wbb, "Table_6")
    if "1938 to" not in str(rows[0][0]):
        raise RuntimeError(f"births Table_6 title changed: {rows[0][0]!r}")
    h = header_index(rows, "Year")
    if str(rows[h][6]).strip() != "Number of Live births":
        raise RuntimeError(f"births Table_6 column 7 is {rows[h][6]!r}")
    births = {}
    for r in rows[h + 1:]:
        if (r and isinstance(r[0], int)
                and r[1] == "England, Wales and Elsewhere"
                and str(r[2]).startswith("All registration")
                and r[3] == "All ages" and r[4] == "All ages"):
            if r[0] in births:
                raise RuntimeError(f"births {r[0]}: duplicate total row")
            births[r[0]] = num(r[6])

    wbd, durl = workbook("deaths", r"/(\d{4})/[^/]+\.xlsx",
                         rank=lambda m: int(m[1]))
    rows = grid(wbd, "Table_1")
    h = header_index(rows, "Year of registration")
    head = [str(c).strip() if c else "" for c in rows[h]]
    if head[3] != "Number of deaths" or not head[9].startswith("Infant mortality"):
        raise RuntimeError(f"deaths Table_1 columns changed: {head}")
    deaths = {}
    for r in rows[h + 1:]:
        if (r and isinstance(r[0], int)
                and r[1] == "England, Wales and Non-residents"
                and r[2] == "All People"):
            deaths[r[0]] = [num(r[3]), num(r[4]), num(r[5]), num(r[9])]
    years = list(range(Y0, LAST_FULL + 1))
    gap = [y for y in years if not births.get(y) or not deaths.get(y)]
    if gap:
        raise RuntimeError(f"vital statistics missing years {gap}")
    for y in years:
        if not 500_000 < births[y] < 800_000 or \
                not 400_000 < deaths[y][0] < 750_000:
            raise RuntimeError(f"vital {y}: implausible {births[y]} / "
                               f"{deaths[y][0]}")
    write("demography/vital_from_ons.csv", f"""
ONS, ENGLAND AND WALES scope, calendar years of REGISTRATION. live_births
= Births in England and Wales: birth registrations ({burl.rsplit('/', 1)[-1]}),
Table 6, England, Wales and Elsewhere, all registration types, all ages.
deaths, crude_death_rate (per 1,000), asmr (age-standardised, per
100,000, 2013 European Standard Population) and infant_mortality_rate
(per 1,000 live births) = Deaths registered, Series DR
({durl.rsplit('/', 1)[-1]}), Table 1, England, Wales and Non-residents, all
people. The latest year can still move with late registrations.
""", ["year", "live_births", "deaths", "crude_death_rate_per_1000",
      "asmr_per_100k", "infant_mortality_per_1000"],
          [[y, fmt(births[y], 0), fmt(deaths[y][0], 0), fmt(deaths[y][1]),
            fmt(deaths[y][2]), fmt(deaths[y][3])] for y in years])

    wbl, lurl = workbook("lifetables", r"/nltuk[^/]*\.xlsx")
    out = []
    for sheet in wbl.sheetnames:
        m = re.fullmatch(r"(\d{4})-(\d{4})", sheet)
        if not m or int(m[2]) < Y0:
            continue
        rows = grid(wbl, sheet)
        if "United Kingdom" not in str(rows[0][0]):
            raise RuntimeError(f"life table {sheet}: scope changed")
        h = header_index(rows, "age")
        head = [str(c).strip() if c is not None else "" for c in rows[h]]
        if head[:6] != ["age", "mx", "qx", "lx", "dx", "ex"] or \
                head[7:13] != ["age", "mx", "qx", "lx", "dx", "ex"]:
            raise RuntimeError(f"life table {sheet}: columns {head}")
        by_age = {r[0]: r for r in rows[h + 1:] if r and isinstance(r[0], int)}
        e = [num(by_age[a][5]) for a in (0, 65)] + \
            [num(by_age[a][12]) for a in (0, 65)]
        if not (70 < e[0] < 85 and 75 < e[2] < 90):
            raise RuntimeError(f"life table {sheet}: e0 {e[0]}/{e[2]}")
        out.append([int(m[2]), sheet] + [fmt(v, 2) for v in e])
    out.sort()
    write("demography/life_expectancy_from_ons.csv", f"""
ONS National life tables, United Kingdom ({lurl.rsplit('/', 1)[-1]}): period
expectation of life (ex) at birth and at 65, males and females. UK scope.
THREE-YEAR POOLED WINDOWS: `year` is the window's last year and
`window` names all three (2024 = 2022-2024); a one-year COVID dip is
spread across three rows.
""", ["year", "window", "e0_male", "e65_male", "e0_female", "e65_female"],
          out)


# ---------------------------------------------------------------- mortality

# ICD-10 chapters as Series DR Table 5 groups them (chapter label, ranges).
# U50.9 is folded into Chapter 20 and the rest of U00-U85 is Chapter 22,
# exactly as DR labels them ('U509, V01 to Y89' / 'U00 to U85').
CHAPTERS = [
    ("Chapter 1", [("A00", "B99")]), ("Chapter 2", [("C00", "D48")]),
    ("Chapter 3", [("D50", "D89")]), ("Chapter 4", [("E00", "E90")]),
    ("Chapter 5", [("F00", "F99")]), ("Chapter 6", [("G00", "G99")]),
    ("Chapter 7", [("H00", "H59")]), ("Chapter 8", [("H60", "H95")]),
    ("Chapter 9", [("I00", "I99")]), ("Chapter 10", [("J00", "J99")]),
    ("Chapter 11", [("K00", "K93")]), ("Chapter 12", [("L00", "L99")]),
    ("Chapter 13", [("M00", "M99")]), ("Chapter 14", [("N00", "N99")]),
    ("Chapter 15", [("O00", "O99")]), ("Chapter 16", [("P00", "P96")]),
    ("Chapter 17", [("Q00", "Q99")]), ("Chapter 18", [("R00", "R99")]),
    ("Chapter 20", [("V01", "Y89")]), ("Chapter 22", [("U00", "U85")]),
]


def chapter_of(code: str) -> str:
    if code.startswith("U509"):
        return "Chapter 20"
    three = code[:3]
    for name, spans in CHAPTERS:
        if any(lo <= three <= hi for lo, hi in spans):
            return name
    raise RuntimeError(f"ICD-10 code {code!r} maps to no DR chapter")


def mortality() -> None:
    # 21st Century Mortality: one sheet per registration year, deaths by
    # 4-character ICD-10 underlying cause, sex and age band.
    wb, curl = workbook("c21", r"/21stcmortality[^/]*\.xlsx")
    by_year = {}
    for sheet in wb.sheetnames:
        if not sheet.isdigit():
            continue
        rows = grid(wb, sheet)
        h = header_index(rows, "ICD-10")
        if [str(c).strip() for c in rows[h][:5]] != \
                ["ICD-10", "Year", "Sex", "Age", "Number of deaths"]:
            raise RuntimeError(f"21C sheet {sheet}: columns {rows[h][:5]}")
        c = Counter()
        years = set()
        for r in rows[h + 1:]:
            if not r or r[0] is None or num(r[4]) is None:
                continue
            years.add(r[1])
            # 'NULL' rows are the neonatal deaths the ONS codes outside
            # the underlying-cause frame; DR's chapter totals exclude them.
            if r[0] != "NULL":
                c[str(r[0]).strip()] += num(r[4])
        if len(years) != 1:
            raise RuntimeError(f"21C sheet {sheet}: years {years}")
        by_year[years.pop()] = c

    def assault(c):
        return sum(v for k, v in c.items()
                   if "X85" <= k[:3] <= "X99" or "Y00" <= k[:3] <= "Y09")

    def onsdef(c):
        return assault(c) + sum(v for k, v in c.items()
                                if k.startswith(("U509", "Y871")))

    years = sorted(y for y in by_year if y >= Y0)
    if years[0] != Y0:
        raise RuntimeError(f"21C starts at {years[0]}")
    for y in years:
        if not 100 < assault(by_year[y]) < 1000:
            raise RuntimeError(f"21C {y}: implausible assault count "
                               f"{assault(by_year[y])}")
    write("security/homicide_deaths_from_ons.csv", f"""
ONS 21st Century Mortality dataset ({curl.rsplit('/', 1)[-1]}), deaths
REGISTERED in ENGLAND AND WALES (incl. non-residents), calendar years,
UNDERLYING cause, summed here over 4-character ICD-10 codes.
assault_x85_y09 = ICD-10 X85-Y09 (assault). homicide_ons_definition =
X85-Y09 + U50.9 + Y87.1, the ONS "homicide and probable homicide"
grouping (U50.9 = homicide/probable homicide pending inquest; it carries
250-570 deaths a year from 2007). CODING BREAK 2006/2007: before 2007
the ONS coded deaths pending inquest to Y33.9 (event of undetermined
intent; ~435 in 2006), which this grouping does not include, so the
{Y0} value is NOT comparable with later years. The dataset ends at {years[-1]};
later years are not yet published at code level and are absent, not
estimated. Series DR's Table 4 shows homicide only where it ranks in an
age band, never as an all-ages total. Death-certificate counts - the
pair of the Home Office Homicide Index, never spliced with it.
""", ["year", "assault_x85_y09", "homicide_ons_definition"],
          [[y, fmt(assault(by_year[y]), 0), fmt(onsdef(by_year[y]), 0)]
           for y in years])

    chap21 = {}
    for y, c in by_year.items():
        agg = Counter()
        for k, v in c.items():
            agg[chapter_of(k)] += v
        agg["All causes"] = sum(c.values())
        chap21[y] = agg

    wbd, durl = workbook("deaths", r"/(\d{4})/[^/]+\.xlsx",
                         rank=lambda m: int(m[1]))
    rows = grid(wbd, "Table_5")
    h = header_index(rows, "Year of registration")
    head = [str(c).strip() for c in rows[h][:7]]
    if head != ["Year of registration", "Sex", "Age group",
                "ICD-10 chapter codes", "ICD-10 chapter name",
                "Place of death", "Number of deaths"]:
        raise RuntimeError(f"DR Table_5 columns changed: {head}")
    dr, names = {}, {}
    for r in rows[h + 1:]:
        if (r and isinstance(r[0], int) and r[1] == "All people"
                and r[5] == "All places"
                # 'All ages' rows (incl. neonates) exist for All causes
                # only; the chapters are excl. neonatal, so all rows are.
                and r[2] == "All ages (excluding neonatal deaths)"):
            key = str(r[4]).split(" - ")[0].strip()
            names[key] = (str(r[3]).replace("‚", ","),
                          str(r[4]).split(" - ", 1)[-1].strip())
            if (r[0], key) in dr:
                raise RuntimeError(f"DR Table_5 {r[0]} {key}: duplicate row")
            dr[(r[0], key)] = num(r[6])
    dr_years = sorted({y for y, _ in dr})
    keys = ["All causes"] + [n for n, _ in CHAPTERS]
    missing = [k for k in keys if (dr_years[-1], k) not in dr]
    if missing:
        raise RuntimeError(f"DR Table_5 lacks {missing}")
    # The splice test: on every overlapping year, every chapter must
    # match Series DR to the death, or the two tables are not one series.
    overlap = [y for y in dr_years if y in chap21]
    if len(overlap) < 3:
        raise RuntimeError(f"21C/DR overlap only {overlap}")
    for y in overlap:
        for k in keys:
            if chap21[y][k] != dr.get((y, k), 0):
                raise RuntimeError(f"{y} {k}: 21C {chap21[y][k]} ≠ DR "
                                   f"{dr.get((y, k))}")
    first_dr = dr_years[0]
    out = []
    for y in range(Y0, dr_years[-1] + 1):
        for k in keys:
            if y >= first_dr:
                v, src = dr[(y, k)], "series_dr_table5"
            else:
                v, src = chap21[y][k], "21st_century_mortality"
            code, name = names[k]
            out.append([y, k.replace("Chapter ", "ch") if k != "All causes"
                        else "all", code, name, fmt(v, 0), src])
    write("mortality/causes_of_death_from_ons.csv", f"""
ONS deaths REGISTERED in ENGLAND AND WALES (incl. non-residents) by
UNDERLYING cause, ICD-10 chapter, all people, all ages excluding
neonatal deaths (DR's convention), all places, calendar years. Long
format. {first_dr} onward = Series DR ({durl.rsplit('/', 1)[-1]}) Table 5;
{Y0}-{first_dr - 1} = 21st Century Mortality dataset aggregated to the same
chapters (U50.9 in Chapter 20). Every chapter of the two sources matches
to the death on {overlap[0]}-{overlap[-1]} (asserted). Underlying cause
only - not 'deaths involving'. Coding breaks to annotate, not smooth:
the 2011 ICD-10 v2010 update and the 2014 MUSE-to-IRIS software change.
""", ["year", "chapter", "icd10_codes", "chapter_name", "deaths", "source"],
          out)


# ---------------------------------------------------------------- migration

MONTH_RANK = {m: i for i, m in enumerate(
    ["january", "february", "march", "april", "may", "june", "july",
     "august", "september", "october", "november", "december"])}


def migration() -> None:
    wb, url = workbook(
        "ltim", r"/yearending([a-z]+)(\d{4})/[^/]+\.xlsx",
        rank=lambda m: (int(m[2]), MONTH_RANK.get(m[1], -1)))
    rows = grid(wb, "1")
    if "Long-term international migration" not in str(rows[0][0]):
        raise RuntimeError(f"LTIM Table 1 title changed: {rows[0][0]!r}")
    h = header_index(rows, "Flow")
    head = [str(c).split("\n")[0].strip() for c in rows[h][:6]]
    if head != ["Flow", "Period", "All Nationalities", "British", "EU+",
                "Non-EU+"]:
        raise RuntimeError(f"LTIM Table 1 columns changed: {head}")
    flows = {"Immigration": "immigration", "Emigration": "emigration",
             "Net migration": "net"}
    out, seen = [], {}
    for r in rows[h + 1:]:
        if not r or r[0] not in flows:
            continue
        m = re.fullmatch(r"YE (\w{3}) (\d{2})((?: [PR])*)", str(r[1]).strip())
        if not m or m[1] not in MON_NUM:
            raise RuntimeError(f"LTIM: unrecognised period {r[1]!r}")
        end = f"20{m[2]}-{MON_NUM[m[1]]:02d}"
        status = "".join(sorted(m[3].split())).replace("PR", "P,R")
        vals = [num(c) for c in r[2:6]]
        if vals[0] is None:
            raise RuntimeError(f"LTIM {r[0]} {r[1]}: empty total")
        seen[(flows[r[0]], end)] = vals[0]
        out.append([end, str(r[1]).strip(), status, flows[r[0]]]
                   + [fmt(v, 0) for v in vals])
    for (f, end), v in seen.items():
        if f == "net":
            calc = seen[("immigration", end)] - seen[("emigration", end)]
            # Net is computed by the ONS from unrounded flows (note 2).
            if abs(calc - v) > 2000:
                raise RuntimeError(f"LTIM {end}: net {v} vs {calc}")
    out.sort(key=lambda r: (["immigration", "emigration", "net"].index(r[3]),
                            r[0]))
    write("migration/flows_from_ons.csv", f"""
ONS Long-term international migration, provisional ({url.rsplit('/', 1)[-1]}),
Table 1: immigration, emigration and net migration, UK, by nationality,
"new approach" admin-based estimates (Home Office visas + DWP RAPID),
rounded to the nearest thousand by the ONS. ONE VINTAGE: this edition's
whole back-series (it carries the late-2025 revisions). REFERENCE PERIODS
ARE ROLLING YEARS ENDING (YE) the month in period_end, published
quarterly - YE Jun 2023 is NOT calendar 2023; only the YE Dec rows are
calendar years. status: P = provisional, R = revised (as published).
British-national estimates before YE June 2021 are not comparable with
later ones (release note 16). Long format; never joined to the
Eurostat IPS-era segment.
""", ["period_end", "period_label", "status", "flow", "all_nationalities",
      "british", "eu_plus", "non_eu_plus"], out)


def main() -> None:
    print("ONS — generator timeseries and bulletin workbooks")
    got = macro()
    fiscal(got["ybha"])
    trade()
    credit()
    social()
    demography()
    mortality()
    migration()
    print("Done. Diff the sidecars before adopting.")


if __name__ == "__main__":
    main()
