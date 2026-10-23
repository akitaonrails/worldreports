#!/usr/bin/env python3
"""Refresh the Federal Statistical Office (BFS / OFS) series for Switzerland.

Usage:  python3 -I countries/switzerland/fetch/fetch_bfs.py

The BFS is the workhorse of this country: national accounts, prices,
the labour market, wages, income distribution, poverty, payment arrears,
the demographic balance, fertility, life tables, projections, police
crime statistics, causes of death and the population by permit all
come from it, over two machine channels.

**The DAM asset API** — `dam-api.bfs.admin.ch/hub/api/dam/assets?orderNr=`
answers a JSON list of assets, and `/assets/{damId}/master` serves the
file (csv or xlsx). The damId changes with every edition, so it is never
written into this file: each run resolves the asset number afresh, and
the cached master is keyed by damId, so a new edition is a new download.
Two things the lookup will do to you (verified 2026-10-22):

  * `orderNr` is a PREFIX match. `je-d-03.03.01.14` also returns the
    `.14.01` and `.14.03` sub-tables; `je-d-01.04.01.01` returns six
    siblings. The fetcher keeps only the exact `shop.orderNr`.
  * Superseded editions stay listed with `bfs.lifecycleGroup`
    NON_CURRENT (four for the CPI, two for life expectancy). The fetcher
    keeps CURRENT only, demands exactly one, and checks that its embargo
    (publication) date is recent enough for the series' rhythm.

**STAT-TAB PxWeb v1** — `www.pxweb.bfs.admin.ch/api/v1/de/{db}/{db}.px`.
A GET returns metadata JSON; a POST with a JSON query and
`"response":{"format":"csv"}` returns CSV in WINDOWS-1252 (decoded as
cp1252 here, never UTF-8). The server answers HTTP 429 after a burst of
about 30 calls in 10 s, so PxWeb calls are spaced 5 s and a 429 backs
off 20 s, 40 s, 60 s. The STATPOP-by-permit table takes about 90 s per
query whatever its size, so it is asked once, with a long timeout.

Things these sources will do to you, verified 2026-10-22:

**Holes published as zeros.** The demographic balance publishes 0 for
immigration and emigration before 1981 (net migration is non-zero in
those years: the flows were not recorded, not nil). The sidecar starts
in 1981 for that reason; the earlier stock is not needed by the window.

**Numbers stored as text.** Recent columns of the causes-of-death,
firearm and assisted-suicide workbooks are strings with apostrophe
thousands separators ("1'774", "36'728"); "*" is not-applicable
(COVID-19 before 2020), "X" suppressed, "..." not available. All are
parsed explicitly; anything else raises.

**The wage-index base files are shorter than the manifest says.** The
base-2010 file starts in 2011, not 2006; 2006-2010 come from the
base-1993 file (NOGA 2002 sections). The three bases are written as
separate rows (`base` column), each with the year-on-year change BFS
publishes; where two bases overlap their changes are asserted equal.

**The GDP CSV does not repeat rows in this edition** (the research pass
saw 1995 twice). The fetcher still asserts one value per key.

**The life tables have no both-sexes figure.** BFS publishes period life
expectancy by sex only; Eurostat's demo_mlexpec total (84.2 in 2024) has
no BFS twin, so men and women are written and nothing is averaged.

**The assisted-suicide table starts in 2003**, not 2009: BFS counts the
cases back to 2003, but until 2008 those same deaths are ALSO inside
"Suizid" in the causes-of-death table. Never add the two.

Cross-checks against a second publisher, asserted every run:
  * BFS GDP (current prices) against Eurostat nama_10_gdp CP_MNAC for
    2006-2022, where the two vintages agree. For 2023+ Eurostat still
    carries the pre-2026-08-25 vintage, so the gap is printed, never
    asserted away and never used to mix a ratio.
  * BFS mean permanent population against the World Bank mid-year
    population (SP.POP.TOTL), within 0.5%.
  * Live births from the PxWeb balance against the DAM fertility
    workbook, to the birth (two BFS channels, one count).
  * Permanent population by permit (sum of permits) against the
    balance's 31 December stock, to the person.

Outputs are `_from_bfs` sidecars under data/. Nothing here overwrites a
curated file.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
import sys
import time
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

try:
    import openpyxl
except ImportError:  # the DAM workbooks cannot be read without it
    sys.exit("openpyxl is required: pip install -r requirements-fetch.txt")

import warnings
# BFS workbooks carry no default style; openpyxl warns on every open.
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
CACHE = ROOT / "cache"

DAM = "https://dam-api.bfs.admin.ch/hub/api/dam/assets"
PX = "https://www.pxweb.bfs.admin.ch/api/v1/de/{db}/{db}.px"
UA = {"User-Agent": "worldreports-fetch"}
GAP = {"dam": 1.5, "px": 5.0, "xcheck": 1.5}
Y0 = 2000                 # a short lead-in before the 2006 window
LAST_FULL = 2025          # country.yaml last_full_year
TODAY = dt.date.today()


# ---------------------------------------------------------------- transport

_last = {k: 0.0 for k in GAP}


def _get(url: str, channel: str, body: bytes | None = None,
         timeout: int = 120) -> bytes:
    """One request, spaced per channel and retried. Raises on failure."""
    last = None
    for attempt in range(4):
        wait = _last[channel] + GAP[channel] - time.monotonic()
        if wait > 0:
            time.sleep(wait)
        try:
            headers = dict(UA)
            if body is not None:
                headers["Content-Type"] = "application/json"
            req = urllib.request.Request(url, data=body, headers=headers)
            with urllib.request.urlopen(req, timeout=timeout) as r:
                blob = r.read()
            _last[channel] = time.monotonic()
            return blob
        except urllib.error.HTTPError as e:
            _last[channel] = time.monotonic()
            last = e
            if e.code in (400, 404):
                break  # a wrong table or query does not heal on retry
            time.sleep(20 * (attempt + 1) if e.code in (429, 403) else 5)
        except OSError as e:
            _last[channel] = time.monotonic()
            last = e
            time.sleep(5 * (attempt + 1))
    raise RuntimeError(f"{url}: {last}")


def cached(name: str, get, check) -> bytes:
    """Cached download; `check(blob)` raises before anything is cached."""
    CACHE.mkdir(exist_ok=True)
    local = CACHE / name
    if local.exists():
        return local.read_bytes()
    blob = get()
    check(blob)
    local.write_bytes(blob)
    return blob


def is_xlsx(blob: bytes) -> None:
    # An error page under HTTP 200 is not a zip; a truncated one fails testzip.
    if blob[:2] != b"PK":
        raise RuntimeError("payload is not an xlsx (error page?)")
    bad = zipfile.ZipFile(io.BytesIO(blob)).testzip()
    if bad:
        raise RuntimeError(f"xlsx member {bad} is corrupt")


def csv_header(first: str):
    def check(blob: bytes) -> None:
        head = blob.decode("utf-8-sig", "replace").splitlines()[:1]
        if not head or not head[0].strip('"').startswith(first):
            raise RuntimeError(f"CSV does not start with {first!r}: "
                               f"{head[:1]}")
        if blob.count(b"\n") < 20:
            raise RuntimeError("CSV has fewer than 20 rows")
    return check


# ---------------------------------------------------------------- DAM

def dam(order_nr: str, fmt: str, max_age_days: int = 400) -> bytes:
    """The CURRENT master file of one exact BFS asset number."""
    url = f"{DAM}?orderNr={order_nr}"
    print(f"  resolving {order_nr}")
    found = json.loads(_get(url, "dam"))
    if not isinstance(found, dict) or "data" not in found:
        raise RuntimeError(f"{order_nr}: DAM lookup answered no data list")
    exact = [a for a in found["data"] if a["shop"]["orderNr"] == order_nr]
    current = [a for a in exact if a["bfs"]["lifecycleGroup"] == "CURRENT"]
    if len(current) != 1:
        raise RuntimeError(
            f"{order_nr}: {len(current)} CURRENT assets among {len(exact)} "
            f"exact and {len(found['data'])} prefix matches")
    a = current[0]
    dam_id = a["ids"]["damId"]
    embargo = dt.date.fromisoformat(a["bfs"]["embargo"][:10])
    age = (TODAY - embargo).days
    if age > max_age_days:
        raise RuntimeError(f"{order_nr}: CURRENT edition {dam_id} is from "
                           f"{embargo}, {age} days old (> {max_age_days})")
    masters = [ln for ln in a["links"] if ln.get("rel") == "master"]
    if len(masters) != 1 or masters[0].get("format") != fmt:
        raise RuntimeError(f"{order_nr}: master is not one {fmt}: {masters}")
    dropped = len(found["data"]) - 1
    print(f"    damId {dam_id}, published {embargo}"
          + (f" ({dropped} prefix/old matches dropped)" if dropped else ""))
    DAM_SEEN[order_nr] = (dam_id, embargo)
    check = is_xlsx if fmt == "xlsx" else csv_header(CSV_FIRST[order_nr])
    return cached(f"bfs_{order_nr}_{dam_id}.{fmt}",
                  lambda: _get(masters[0]["href"], "dam", timeout=300),
                  check)


DAM_SEEN: dict[str, tuple[int, dt.date]] = {}
CSV_FIRST = {
    "ts-x-04.02.01.02": "PERIOD",
    "ts-x-03.02.01.01": "INDICATORS_HRCHY",
    "ts-x-03.04.03.00.04": "YEAR",
    "ts-x-03.04.03.00.03": "YEAR",
    "ts-x-03.04.03.00.06": "YEAR",
}


def dam_csv(order_nr: str) -> list[dict]:
    blob = dam(order_nr, "csv")
    return list(csv.DictReader(io.StringIO(blob.decode("utf-8-sig"))))


def dam_xlsx(order_nr: str, max_age_days: int = 400):
    blob = dam(order_nr, "xlsx", max_age_days)
    return openpyxl.load_workbook(io.BytesIO(blob), read_only=True,
                                  data_only=True)


def grid(wb, sheet: str) -> list[list]:
    ws = wb[sheet]
    ws.reset_dimensions()  # some BFS sheets declare a wrong used range
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    width = max((len(r) for r in rows), default=0)
    return [r + [None] * (width - len(r)) for r in rows]  # ragged in read_only


def ref(order_nr: str) -> str:
    dam_id, embargo = DAM_SEEN[order_nr]
    return f"BFS asset {order_nr} (damId {dam_id}, published {embargo})"


# ---------------------------------------------------------------- PxWeb

def px(db: str, query: list[dict], name: str, first: str,
       timeout: int = 180) -> list[list[str]]:
    """POST one PxWeb query; returns csv rows decoded from cp1252."""
    body = json.dumps({"query": query, "response": {"format": "csv"}},
                      ensure_ascii=False).encode("utf-8")

    def get() -> bytes:
        print(f"  querying PxWeb {db} ({name})")
        return _get(PX.format(db=db), "px", body, timeout)

    def check(blob: bytes) -> None:
        text = blob.decode("cp1252")
        head = text.splitlines()[:1]
        if not head or not head[0].startswith(f'"{first}"'):
            raise RuntimeError(f"{db}: CSV header {head[:1]} (error page?)")
        if text.count("\n") < 2:
            raise RuntimeError(f"{db}: CSV has no data rows")

    blob = cached(f"bfs_px_{name}.csv", get, check)
    return list(csv.reader(io.StringIO(blob.decode("cp1252"))))


# ---------------------------------------------------------------- parsing

MISSING = {"", "*", "X", "x", "...", "..", "…", "-", "–"}


def num(v) -> float | None:
    """A BFS cell as a number; holes are None. Unknown text raises."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("'", "").replace("’", "").replace(" ", "")
    if s in MISSING:
        return None
    try:
        return float(s)
    except ValueError:
        raise RuntimeError(f"unparseable BFS cell {v!r}") from None


def year_of(v) -> int | None:
    """A header cell such as 2010, '2010 1' or 'II\\n2006' as a year."""
    if isinstance(v, (int, float)) and 1900 <= v <= 2100:
        return int(v)
    m = re.search(r"\b(19|20)\d{2}\b", str(v or ""))
    return int(m[0]) if m else None


def label(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip()


def find_row(rows, test, what: str) -> int:
    for i, r in enumerate(rows):
        if r and test(label(r[0])):
            return i
    raise RuntimeError(f"row not found: {what}")


def fmt(v, nd: int = 1):
    if v is None:
        return ""
    if nd == 0:
        return str(int(round(v)))
    return f"{v:.{nd}f}"


def write(rel: str, comment: str, header: list[str], rows: list[list]) -> None:
    path = DATA / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        lines = comment.strip().splitlines() + [
            "Years are CALENDAR years (Swiss public finance and statistics"
            " use the calendar year; nothing here is an April fiscal year)."]
        for line in lines:
            fh.write(f"# {line}".rstrip() + "\n")
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(rows)
    span = f"{rows[0][0]}–{rows[-1][0]}" if rows else "EMPTY"
    print(f"  data/{rel}  {len(rows)} rows ({span})")


def expect(what: str, got, want, tol: float = 0.0) -> None:
    """Spot value from the research pass; raise when the payload disagrees."""
    if got is None or abs(got - want) > tol:
        raise RuntimeError(f"spot check failed: {what} = {got}, expected "
                           f"{want} (±{tol})")


# ---------------------------------------------------------------- macro

def gdp() -> dict[int, float]:
    o = "ts-x-04.02.01.02"
    rows = dam_csv(o)
    vals: dict[tuple, float | None] = {}
    for r in rows:
        key = (int(r["PERIOD"]), r["SECTOR"], r["CLASSIFICATION"],
               r["UNIT_MEAS"])
        v = None if r["STATUS"] == "O" or r["VALUE"] == "" else float(r["VALUE"])
        if key in vals and vals[key] != v:
            raise RuntimeError(f"GDP CSV: two values for {key}")
        vals[key] = v
    years = sorted({k[0] for k in vals})
    if years[-1] < LAST_FULL:
        raise RuntimeError(f"GDP CSV ends in {years[-1]}")
    cols = [("S1", "B1GQ", "MCHF"), ("S1", "B1GQ", "AC"),
            ("S1", "B1GQ", "ACPP"), ("S14_S15", "P3", "MCHF"),
            ("S13", "P3", "MCHF"), ("S1", "P51G", "MCHF"),
            ("S2", "P6", "MCHF"), ("S2", "P7", "MCHF"),
            ("S2", "P61", "MCHF"), ("S2", "P61XF11", "MCHF"),
            ("S2", "P71", "MCHF"), ("S2", "P71XF11", "MCHF")]
    have = {k[1:] for k in vals}
    for c in cols:
        if c not in have:
            raise RuntimeError(f"GDP CSV lacks {c}; has "
                               f"{sorted(x for x in have if x[2] == 'MCHF')}")
    out = []
    for y in years:
        if y < Y0:
            continue
        g = [vals.get((y,) + c) for c in cols]
        if g[0] is None:
            raise RuntimeError(f"GDP {y}: nominal B1GQ missing")
        out.append([y] + [fmt(v, 1) if i != 1 and i != 2 else fmt(v, 2)
                          for i, v in enumerate(g)])
    b1 = {y: vals[(y, "S1", "B1GQ", "MCHF")] for y in years}
    expect("GDP 2006 CHF mn", b1[2006], 562159.6, 0.1)
    expect("GDP 2024 CHF mn", b1[2024], 860932.1, 0.1)
    expect("GDP 2025 CHF mn", b1[2025], 881571.8, 0.1)
    expect("real GDP growth 2009", vals[(2009, "S1", "B1GQ", "ACPP")], -1.93, 0.005)
    expect("real GDP growth 2020", vals[(2020, "S1", "B1GQ", "ACPP")], -2.26, 0.005)
    expect("real GDP growth 2025", vals[(2025, "S1", "B1GQ", "ACPP")], 1.61, 0.005)
    write("macro/gdp_from_bfs.csv", f"""
{ref(o)}: Bruttoinlandprodukt nach Verwendungsarten, annual, ESA 2010.
CHF MILLIONS at current prices unless the column says pct. gdp_nominal =
B1GQ S1; nominal_growth_pct = AC (% change, current prices);
real_growth_pct = ACPP (% change at previous-year prices, the official
real growth). household_consumption = P3 S14_S15 (households + NPISH);
government_consumption = P3 S13; gfcf = P51G; exports = P6, imports = P7
(goods and services, S2 rest of world, gold INCLUDED). GOLD IN TRADE:
exports_goods = P61 includes non-monetary gold, exports_goods_ex_gold =
P61XF11 excludes it (2025: 517.5bn vs 365.0bn); same for imports P71 /
P71XF11. Separate columns, one total per line, named. VINTAGE: this is the
BFS revision of 2026-08-25. Eurostat nama_10_gdp and the World Bank still
carry the PREVIOUS vintage, which agrees through 2022 and parts from 2023
(asserted equal 2006-2022 each run; 2024: Eurostat 854.0bn vs BFS
860.9bn); a ratio built on this GDP must use BFS numerators only, and a
Eurostat ratio must never be redrawn on this denominator. 2025 is the
BFS first estimate. Blank = not published (STATUS O).
""", ["year", "gdp_nominal_chf_mn", "gdp_nominal_growth_pct",
      "gdp_real_growth_pct", "household_consumption_chf_mn",
      "government_consumption_chf_mn", "gfcf_chf_mn", "exports_chf_mn",
      "imports_chf_mn", "exports_goods_chf_mn", "exports_goods_ex_gold_chf_mn",
      "imports_goods_chf_mn", "imports_goods_ex_gold_chf_mn"], out)
    return b1


def cpi() -> None:
    o = "cc-d-05.02.08"
    wb = dam_xlsx(o, max_age_days=100)  # monthly release
    rows = grid(wb, "Index_y")
    h = find_row(rows, lambda s: s.startswith("Datum"), "CPI base header")
    head = rows[h]
    if label(head[-1]) != "%" and "%" not in [label(c) for c in head]:
        raise RuntimeError("CPI Index_y: no % column")
    pct_col = [label(c) for c in head].index("%")
    bases = {c.year * 100 + c.month: i for i, c in enumerate(head)
             if isinstance(c, dt.datetime)}
    for b in (200512, 202012):
        if b not in bases:
            raise RuntimeError(f"CPI Index_y: base {b} missing: {bases}")
    # The Dec-2025 = 100 basket must exist in the monthly sheet (rebased
    # early 2026); its annual column appears only once 2026 is complete.
    mhead = [c for r in grid(wb, "Index_m")[:8] for c in r]
    if not any(isinstance(c, dt.datetime) and (c.year, c.month) == (2025, 12)
               for c in mhead):
        raise RuntimeError("CPI Index_m: no December 2025 base column")
    out, rate = [], {}
    for r in rows[h + 1:]:
        y = r[0] if isinstance(r[0], (int, float)) else None
        if y is None:
            continue
        y = int(y)
        rate[y] = num(r[pct_col])
        if y < Y0:
            continue
        out.append([y, fmt(rate[y], 1), fmt(num(r[bases[200512]]), 1),
                    fmt(num(r[bases[202012]]), 1)])
    for y, v in {2008: 2.4, 2009: -0.5, 2015: -1.1, 2022: 2.8, 2023: 2.1,
                 2024: 1.1, 2025: 0.2}.items():
        expect(f"CPI inflation {y}", rate.get(y), v, 0.001)
    if out[-1][0] < LAST_FULL:
        raise RuntimeError(f"CPI annual ends {out[-1][0]}")
    write("macro/cpi_from_bfs.csv", f"""
{ref(o)}: Landesindex der Konsumentenpreise (LIK), sheet Index_y.
inflation_pct = the official annual average change (the % column).
index_dec2005_100 and index_dec2020_100 are the annual averages on two
of BFS's original bases, as published; they are LEVELS ON DIFFERENT
BASES and never spliced. BFS rebased the basket to December 2025 = 100
in early 2026; that base has no annual average yet. Take rates, not
levels across bases.
""", ["year", "cpi_inflation_pct", "cpi_index_dec2005_100",
      "cpi_index_dec2020_100"], out)


def unemployment() -> None:
    o = "je-d-03.03.00.01"
    rows = grid(dam_xlsx(o), "T3.3.0.1")
    h = next(i for i, r in enumerate(rows)
             if r and r[0] is None and year_of(r[1]) == 2000)
    years = {i: year_of(c) for i, c in enumerate(rows[h]) if year_of(c)}
    want = {"jobseekers": "Registrierte Stellensuchende gemäss SECO",
            "registered": "Registrierte Arbeitslose gemäss SECO",
            "ilo": "Erwerbslose (ILO) des BFS"}
    got = {}
    for k, text in want.items():
        i = find_row(rows, lambda s, t=text: s.startswith(t), text)
        got[k] = {y: num(rows[i][c]) for c, y in years.items()}
    notes = " ".join(label(r[0]) for r in rows if r and r[0]
                     and label(r[0])[:1] in "123")
    if "März 2018" not in notes or "2009" not in notes:
        raise RuntimeError("unemployment footnotes changed (break notes)")
    out = [[y, fmt(got["jobseekers"][y], 1), fmt(got["registered"][y], 1),
            fmt(got["ilo"][y], 1)] for y in sorted(years.values()) if y >= Y0]
    expect("SECO registered unemployed 2006 (k)", got["registered"][2006], 131.5, 0.05)
    expect("ILO unemployed 2006 (k)", got["ilo"][2006], 195.3, 0.05)
    expect("SECO job seekers 2025 (k)", got["jobseekers"][2025], 214.1, 0.05)
    expect("ILO unemployed 2025 (k)", got["ilo"][2025], 248.5, 0.05)
    write("macro/unemployment_from_bfs.csv", f"""
{ref(o)}: annual means, THOUSANDS of persons. Three different counts:
seco_registered_jobseekers and seco_registered_unemployed are
ADMINISTRATIVE (RAV registers, SECO); ilo_unemployed is the SURVEY count
(SAKE/ESPA, BFS, ILO definition). Two concepts, never one line. Breaks
(footnotes): SECO moved the RAV offices to a new semi-automated
registration system in March 2018; the 2017 SAKE revision broke the ILO
series between 2009 and 2010, which BFS has back-adjusted in this table.
""", ["year", "seco_registered_jobseekers_k", "seco_registered_unemployed_k",
      "ilo_unemployed_k"], out)


def unemployment_rate() -> None:
    o = "je-d-03.03.01.14"
    wb = dam_xlsx(o)

    def block(sheet: str, quarter: bool) -> dict[str, dict]:
        rows = grid(wb, sheet)
        h = find_row(rows, lambda s: s == "Total", "period header")
        cols = {}
        for c, v in enumerate(rows[h]):
            y = year_of(v)
            if y is None:
                continue
            q = label(v).split(" ")[0] if quarter else "Y"
            if quarter and q != "II":
                continue  # Q2 only: the one quarter surveyed since 1991
            cols[c] = y
        out = {}
        for key, test in (("total", lambda s: s == "Total"),
                          ("swiss", lambda s: s.startswith("Schweizer/innen")),
                          ("foreign", lambda s: s.startswith("Ausländer/innen"))):
            i = next(i for i, r in enumerate(rows) if i > h and r
                     and test(label(r[0])))
            out[key] = {y: num(rows[i][c]) for c, y in cols.items()}
        return out

    q2 = block("Quartalswerte", True)
    yr = block("Jahreswerte", False)
    years = sorted(set(q2["total"]) | set(yr["total"]))
    out = [[y, fmt(yr["total"].get(y), 2), fmt(q2["total"].get(y), 2),
            fmt(yr["swiss"].get(y), 2), fmt(yr["foreign"].get(y), 2)]
           for y in years if y >= Y0]
    expect("ILO unemployment rate 2024", yr["total"][2024], 4.34, 0.005)
    expect("ILO unemployment rate 2025", yr["total"][2025], 4.84, 0.005)
    if min(yr["total"]) != 2010:
        raise RuntimeError(f"annual ILO rate starts {min(yr['total'])}")
    write("macro/unemployment_rate_from_bfs.csv", f"""
{ref(o)}: Erwerbslosenquote gemäss ILO (SAKE), %, permanent resident
population. ilo_rate_annual_pct = annual average of the continuous
quarterly survey, 2010 onward (sheet Jahreswerte). ilo_rate_q2_pct = the
second-quarter value (sheet Quartalswerte): before 2010 the SAKE was a
Q2-only survey, so 2000-2009 exist ONLY as Q2 values; the Q2 column is
continued after 2010 for a like-for-like line. The two columns are
different statistics and are not spliced. swiss/foreign = annual by
nationality, 2010 onward. ILO SURVEY rate, not SECO's registered rate.
The last row may be a year in progress: its Q2 value is published, its
annual average is blank until the year is complete.
""", ["year", "ilo_rate_annual_pct", "ilo_rate_q2_pct",
      "ilo_rate_swiss_annual_pct", "ilo_rate_foreign_annual_pct"], out)


def employment() -> None:
    o = "ts-x-03.02.01.01"
    rows = dam_csv(o)
    kinds = {(r["INDICATORS_DE"], r["DETAILS_DE"], r["FREQ"]) for r in rows}
    if kinds != {("TOTAL", "Total", "Y")}:
        raise RuntimeError(f"employment CSV holds more than the total: {kinds}")
    vals = {}
    for r in rows:
        y = int(r["PERIOD"])
        if y in vals:
            raise RuntimeError(f"employment CSV: {y} twice")
        vals[y] = None if r["STATUS"] == "O" or r["VALUE"] == "" else float(r["VALUE"])
    out = [[y, fmt(vals[y], 0)] for y in sorted(vals) if y >= Y0]
    if any(v[1] == "" for v in out):
        raise RuntimeError("employment has holes inside the window")
    if 4.1e6 > vals[2006] or vals[2006] > 4.4e6 or not 5.2e6 < vals[2025] < 5.5e6:
        raise RuntimeError("employment magnitude off (expected ~4.2mn in 2006,"
                           " ~5.4mn in 2025)")
    write("macro/employment_from_bfs.csv", f"""
{ref(o)}: Erwerbstätigenstatistik (ETS), employed persons, annual mean,
PERSONS. DOMESTIC CONCEPT (Inlandkonzept): everyone working in
Switzerland, including about 400k cross-border commuters (Grenzgänger)
and short-stay workers, which the resident-population SAKE does not
count. Not comparable with an LFS employment rate.
""", ["year", "employed_persons_domestic_concept"], out)


def wages() -> None:
    files = [("ts-x-03.04.03.00.04", "1993", "C-O", "NOGA 2002"),
             ("ts-x-03.04.03.00.03", "2010", "B-S", "NOGA 2008"),
             ("ts-x-03.04.03.00.06", "2020", "B-S", "NOGA 2008")]
    out, pct = [], {}
    for o, base, section, noga in files:
        rows = dam_csv(o)
        got = {}
        for r in rows:
            if r["SECTION"] != section or r["SEX"] != "T":
                continue
            key = (int(r["YEAR"]), r["WAGE_TYPE"])
            if key in got:
                raise RuntimeError(f"{o}: {key} twice")
            got[key] = (num(r["VALUE"]), num(r["VALUE_P"]), r["OBS_STATUS_VALUE"])
        years = sorted({k[0] for k in got})
        if not years:
            raise RuntimeError(f"{o}: no {section} total rows")
        for y in years:
            n, rl = got.get((y, "N")), got.get((y, "R"))
            if n is None or rl is None:
                raise RuntimeError(f"{o} {y}: nominal or real row missing")
            for t, v in (("N", n[1]), ("R", rl[1])):
                if (y, t) in pct and abs(pct[(y, t)] - v) > 0.05:
                    raise RuntimeError(f"wage change {y}{t}: base {base} "
                                       f"says {v}, earlier base {pct[(y, t)]}")
                pct[(y, t)] = v
            if y >= Y0:
                # The base is written as a label ("2010=100"), not a bare
                # number: it is part of the row key, beside the year.
                out.append([y, f"{base}=100", noga, section, fmt(n[0], 1), fmt(n[1], 1),
                            fmt(rl[0], 1), fmt(rl[1], 1),
                            n[2] if n[2] == rl[2] else f"{n[2]}/{rl[2]}"])
        print(f"    base {base}: {years[0]}–{years[-1]}")
    expect("wage index base 2020, 2021 nominal", next(
        float(r[4]) for r in out if r[0] == 2021 and r[1] == "2020=100"), 99.8)
    expect("nominal wage change 2021", pct[(2021, "N")], -0.2)
    for y in range(2006, LAST_FULL + 1):
        if (y, "N") not in pct:
            raise RuntimeError(f"no wage index covers {y}")
    out.sort(key=lambda r: (r[0], r[1]))
    write("macro/wages_from_bfs.csv", """
BFS Schweizerischer Lohnindex (SLI), assets ts-x-03.04.03.00.04 (base
1993 = 100, NOGA 2002, total C-O), .00.03 (base 2010 = 100, NOGA 2008,
total B-S) and .00.06 (base 2020 = 100, total B-S); both sexes. One row
per year PER BASE: index LEVELS are on different bases and must never be
spliced; nominal_pct / real_pct are BFS's own year-on-year changes, and
where two bases cover a year their changes are asserted equal, so the
pct columns are the series to draw across 2006-2025. real = deflated by
the CPI (BFS). The base-2010 file starts in 2011 and the base-1993 total
stops in 2010. base = the index base as a label (1993=100, 2010=100,
2020=100); (year, base) is the row key. status: A = definitive, others
as published by BFS.
""", ["year", "base", "classification", "section", "nominal_index",
      "nominal_pct", "real_index", "real_pct", "status"], out)


# ---------------------------------------------------------------- social

def gini() -> None:
    o = "je-d-20.03.01.02.02"
    wb = dam_xlsx(o)
    out, seen = [], {}
    for sheet in wb.sheetnames:
        y = int(sheet)
        rows = grid(wb, sheet)
        g4 = next(i for i, r in enumerate(rows) if r and len(r) > 1
                  and label(r[1]) == "Gesamtbevölkerung")
        inc = [label(rows[g4 + 1][c]) for c in (1, 3, 5)]
        if not (inc[0].startswith("Primär") and inc[1].startswith("Brutto")
                and inc[2].startswith("Verfügbar")):
            raise RuntimeError(f"HABE {y}: income columns changed: {inc}")
        gi = find_row(rows, lambda s: s.startswith("Gini"), f"Gini {y}")
        si = find_row(rows, lambda s: s.startswith("S80/S20"), f"S80 {y}")
        g, s = rows[gi], rows[si]
        seen[y] = num(g[5])
        if y >= Y0:
            out.append([y, fmt(num(g[1]), 3), fmt(num(g[3]), 3),
                        fmt(num(g[5]), 3), fmt(num(g[6]), 3),
                        fmt(num(s[5]), 2)])
    out.sort()
    expect("HABE Gini disposable 2023", seen[2023], 0.305, 0.0005)
    expect("HABE Gini primary 2023", num(grid(wb, "2023")[find_row(
        grid(wb, "2023"), lambda s: s.startswith("Gini"), "Gini")][1]),
        0.458, 0.0005)
    write("social/gini_from_bfs.csv", f"""
{ref(o)}: income distribution before and after state transfers, from
the household budget survey (HABE), total population, equivalised
monthly household income. gini_primary = market (primary) income;
gini_gross = after transfers received; gini_disposable = after taxes
and compulsory contributions (the headline). ci95 = half-width of BFS's
95% confidence interval. Gini as a FRACTION (0.305), not points.
THREE GINIS EXIST for Switzerland and none is the other: HABE (this),
SILC (Eurostat ilc_di12, 31.5 in 2023) and the World Bank (33.5).
Survey years with no sheet (1999) are absent, not interpolated.
""", ["year", "gini_primary", "gini_gross", "gini_disposable",
      "gini_disposable_ci95", "s80_s20_disposable"], out)


def poverty() -> None:
    o = "je-d-20.03.02.01.01"
    wb = dam_xlsx(o)
    out, seen = [], {}
    for sheet in wb.sheetnames:
        y = int(sheet)
        rows = grid(wb, sheet)
        hdr = next(r for r in rows if r and any(label(c) == "in 1000" for c in r))
        if label(hdr[1]) != "in 1000" or label(hdr[3]) != "in %":
            raise RuntimeError(f"poverty {y}: columns changed {hdr[:5]}")
        t = rows[find_row(rows, lambda s: s == "Gesamtbevölkerung", f"total {y}")]
        old = [r for r in rows if r and re.match(r"(Ab 65|65 Jahre und|65\+)",
                                                 label(r[0]))]
        seen[y] = (num(t[1]), num(t[3]), num(old[0][3]) if old else None)
        out.append([y, fmt(num(t[1]), 1), fmt(num(t[3]), 1), fmt(num(t[4]), 1),
                    fmt(seen[y][2], 1)])
    out.sort()
    expect("absolute poverty rate 2024", seen[2024][1], 8.4, 0.05)
    expect("people in poverty 2024 (k)", seen[2024][0], 743.4, 0.05)
    expect("poverty rate 65+ 2024", seen[2024][2], 14.9, 0.05)
    write("social/poverty_from_bfs.csv", f"""
{ref(o)}: Armutsquote, SILC, permanent resident population in private
households. ABSOLUTE poverty against the national social-minimum line
(SKOS guidelines: basic needs + housing + CHF 100 per person 16+), NOT
the EU at-risk-of-poverty rate (60% of median, Eurostat ilc_li02): two
definitions, two lines. poor_persons_k in THOUSANDS; rate in %; ci95 =
half-width of the 95% confidence interval in points; rate_65plus = the
'Ab 65 Jahren' row. Income year = survey year - 1 in SILC; years as BFS
labels them.
""", ["year", "poor_persons_k", "poverty_rate_pct", "poverty_rate_ci95",
      "poverty_rate_65plus_pct"], out)


ARREARS = [  # (column, test on the lowercased header path)
    ("any_kind_pct", lambda p: "einer art von" in p),
    ("two_plus_kinds_pct", lambda p: "zwei arten von" in p),
    ("rent_tenants_pct", lambda p: "anteil der mieter" in p
     or re.search(r"/miete( \[\d+\])?$", p)),
    ("mortgage_owners_pct", lambda p: "anteil der eigentümer" in p
     or re.search(r"/hypothekarzinsen( \[\d+\])?$", p)),
    ("rent_or_mortgage_pct", lambda p: "miete oder hypothekarzinsen" in p
     and p.endswith("mindestens einmal")),
    ("utilities_pct", lambda p: "laufende rechnungen" in p
     and p.endswith("mindestens einmal")),
    ("loans_or_cards_pct", lambda p: "oder unbezahlte" in p),
    ("loans_pct", lambda p: "kreditrückzahlungen" in p and "oder unbezahlte"
     not in p and p.endswith("mindestens einmal")),
    ("credit_cards_pct", lambda p: p.split("/")[-2:-1] == ["kreditkartenrechnungen"]
     and p.endswith("mindestens einmal")),
    ("health_insurance_pct", lambda p: "krankenkassenprämien" in p
     and p.endswith("mindestens einmal")),
    ("taxes_pct", lambda p: "/steuern" in p and p.endswith("mindestens einmal")),
]


def arrears() -> None:
    o = "je-d-20.02.04.08"
    wb = dam_xlsx(o)
    out, seen = [], {}
    for sheet in wb.sheetnames:
        y = int(sheet)
        rows = grid(wb, sheet)
        pm = next(i for i, r in enumerate(rows)
                  if r and any(label(c) == "+/-" for c in r))
        groups = []
        for gi in (pm - 2, pm - 1):  # merged group labels: forward-fill
            ff, cur = [], None
            for c, v in enumerate(rows[gi]):
                cur = label(v) if c > 0 and v is not None else (cur if c > 0 else None)
                ff.append(cur)
            groups.append(ff)
        data = rows[find_row(rows, lambda s: s == "Gesamtbevölkerung", f"total {y}")]
        got = {}
        for c in range(1, len(rows[pm]) - 1):
            if label(rows[pm][c + 1]) != "+/-" or label(rows[pm][c]) == "+/-":
                continue  # a value column is followed by its CI column
            parts = [groups[0][c], groups[1][c], label(rows[pm][c]),
                     label(rows[pm + 1][c]) if pm + 1 < len(rows) else ""]
            path = "/".join(p for p in parts if p).lower()
            hits = [k for k, test in ARREARS if test(path)]
            if len(hits) > 1:
                raise RuntimeError(f"arrears {y} col {c}: {path!r} -> {hits}")
            if hits:
                if hits[0] in got:
                    raise RuntimeError(f"arrears {y}: {hits[0]} twice")
                got[hits[0]] = (num(data[c]), num(data[c + 1]))
        need = {"any_kind_pct", "two_plus_kinds_pct", "rent_tenants_pct",
                "mortgage_owners_pct", "rent_or_mortgage_pct", "utilities_pct",
                "loans_pct", "health_insurance_pct", "taxes_pct"}
        if need - set(got):
            raise RuntimeError(f"arrears {y}: missing {need - set(got)}")
        seen[y] = got
        out.append([y] + [fmt(got[k][0], 2) if k in got else ""
                          for k, _ in ARREARS] + [fmt(got["any_kind_pct"][1], 2)])
    out.sort()
    expect("arrears of any kind 2024", seen[2024]["any_kind_pct"][0], 15.4, 0.05)
    expect("arrears of two+ kinds 2024", seen[2024]["two_plus_kinds_pct"][0], 8.0, 0.05)
    expect("mortgage arrears among owners 2024",
           seen[2024]["mortgage_owners_pct"][0], 0.21, 0.005)
    # The research pass recorded "taxes 3.9%" for 2024; 3.86 is the
    # utilities column and taxes are 8.75. Asserted as the file has them.
    expect("utility-bill arrears 2024", seen[2024]["utilities_pct"][0], 3.86, 0.005)
    expect("tax arrears 2024", seen[2024]["taxes_pct"][0], 8.75, 0.005)
    write("social/arrears_from_bfs.csv", f"""
{ref(o)}: Zahlungsrückstände, SILC: % of PEOPLE living in a household
that, in the last 12 months and for financial reasons, fell behind at
least once on the item. A SURVEY of people, not a loan register;
Switzerland publishes no bank delinquency series. Denominators differ:
rent_tenants_pct is among tenants and mortgage_owners_pct among owners
(with or without a mortgage); every other column is among the whole
population. loans_pct: 2013-2016 'ausser Hypothekarzinsen für den
Hauptwohnsitz', from 2017 leasing, consumer credit, instalments,
second-home mortgages and loans from family outside the household - a
definitional seam in 2017, annotate. loans_or_cards_pct and
credit_cards_pct start in 2014 (blank = not asked, not zero).
any_kind_ci95 = half-width of the 95% CI for any_kind_pct. Survey year
as BFS labels it.
""", ["year"] + [k for k, _ in ARREARS] + ["any_kind_ci95"], out)


# ---------------------------------------------------------------- demography

BAL_DB = "px-x-0102020000_101"
BAL = {  # sidecar column -> PxWeb component label
    "pop_1jan": "Bestand am 1. Januar",
    "births": "Lebendgeburt",
    "deaths": "Todesfall",
    "natural_change": "Geburtenüberschuss",
    "immigration_incl_type_change": "Einwanderung inkl. Änderung des Bevölkerungstyps",
    "emigration": "Auswanderung",
    "net_migration_incl_type_change": "Wanderungssaldo inkl. Änderung des Bevölkerungstyps",
    "change_of_population_type": "Änderung des Bevölkerungstyps",
    "naturalisations": "Erwerb des Schweizer Bürgerrechts",
    "stock_adjustment": "Bestandesbereinigung",
    "pop_31dec": "Bestand am 31. Dezember",
}


def balance_rows() -> dict[str, dict[int, dict[str, int]]]:
    rows = px(BAL_DB, [
        {"code": "Kanton", "selection": {"filter": "item", "values": ["0"]}},
        {"code": "Staatsangehörigkeit (Kategorie)",
         "selection": {"filter": "item", "values": ["0", "1", "2"]}},
        {"code": "Geschlecht", "selection": {"filter": "item", "values": ["0"]}},
    ], "balance_ch", "Jahr")
    head = rows[0]
    for v in BAL.values():
        if v not in head:
            raise RuntimeError(f"balance CSV lacks component {v!r}")
    cit = {"Staatsangehörigkeit (Kategorie) - Total": "total",
           "Schweiz": "swiss", "Ausland": "foreign"}
    out: dict[str, dict[int, dict[str, int]]] = {k: {} for k in cit.values()}
    for r in rows[1:]:
        if not r:
            continue
        d = dict(zip(head, r))
        if d["Kanton"] != "Schweiz":
            raise RuntimeError(f"balance: unexpected canton {d['Kanton']}")
        out[cit[d["Staatsangehörigkeit (Kategorie)"]]][int(d["Jahr"])] = {
            k: int(num(d[v])) for k, v in BAL.items()}
    return out


def balance(bal) -> None:
    t = bal["total"]
    first = min(y for y, d in t.items() if d["immigration_incl_type_change"] > 0)
    if first != 1981:
        raise RuntimeError(f"balance flows begin {first}, expected 1981")
    out = []
    for y in sorted(t):
        if y < first:
            continue
        d = t[y]
        if d["immigration_incl_type_change"] <= 0 or d["emigration"] <= 0:
            raise RuntimeError(f"balance {y}: a migration flow is zero")
        calc = d["pop_1jan"] + d["natural_change"] + \
            d["net_migration_incl_type_change"] + d["stock_adjustment"]
        if calc != d["pop_31dec"]:
            raise RuntimeError(f"balance {y}: identity off by {calc - d['pop_31dec']}")
        sw = bal["swiss"][y]["naturalisations"]
        out.append([y] + [d[k] for k in BAL if k != "naturalisations"] + [sw])
    d = t[2025]
    expect("population 1 Jan 2025", d["pop_1jan"], 9051029)
    expect("births 2025", d["births"], 78153)
    expect("deaths 2025", d["deaths"], 71904)
    expect("immigration 2025", d["immigration_incl_type_change"], 204261)
    expect("emigration 2025", d["emigration"], 133829)
    expect("population 31 Dec 2025", d["pop_31dec"], 9127125)
    expect("naturalisations 2025", bal["swiss"][2025]["naturalisations"], 40672)
    write("demography/balance_from_bfs.csv", f"""
BFS STAT-TAB {BAL_DB} (STATPOP/BEVNAT), Demografische Bilanz, Schweiz,
both sexes, all citizenships. PERSONS. PERMANENT resident population
(since 2010 including asylum seekers and S-status holders resident 12+
months). immigration/emigration/net INCLUDE CHANGES OF POPULATION TYPE
('inkl. Änderung des Bevölkerungstyps'): a person who has lived 12 months
in the NON-permanent population enters the permanent one and is counted
as an immigrant then - the 2023 peak is largely Ukrainians with S status
crossing that line. change_of_population_type is that component alone.
naturalisations = acquisitions of Swiss citizenship (a transfer between
citizenships, zero in the national total, so taken from the Swiss row).
Identity asserted: pop_31dec = pop_1jan + natural_change + net + stock
adjustment. pop_1jan 1991 does not equal pop_31dec 1990 (a published
register revision). Starts 1981: BFS publishes 0 for the flows before
then, which are holes, not zeros.
""", ["year"] + [k for k in BAL if k != "naturalisations"] + ["naturalisations"],
        out)


def flows(bal) -> None:
    out = []
    for y in sorted(bal["total"]):
        if y < Y0:
            continue
        row = [y]
        for k in ("immigration_incl_type_change", "emigration",
                  "net_migration_incl_type_change"):
            parts = [bal[c][y][k] for c in ("total", "swiss", "foreign")]
            if parts[0] != parts[1] + parts[2]:
                raise RuntimeError(f"flows {y} {k}: total != swiss + foreign")
            row += parts
        row.append(bal["foreign"][y]["change_of_population_type"])
        out.append(row)
    f = bal["foreign"]
    expect("foreign immigration 2025", f[2025]["immigration_incl_type_change"], 181279)
    expect("foreign emigration 2025", f[2025]["emigration"], 104423)
    expect("foreign immigration 2023", f[2023]["immigration_incl_type_change"], 241040)
    expect("total immigration 2023",
           bal["total"][2023]["immigration_incl_type_change"], 263081)
    write("migration/flows_from_bfs.csv", f"""
BFS STAT-TAB {BAL_DB}, Demografische Bilanz, Schweiz: international
migration of the PERMANENT resident population by citizenship
(total = swiss + foreign, asserted). PERSONS. IMMIGRATION INCLUDES
CHANGES OF POPULATION TYPE - the label must say 'incl. change of
population type' on the plate: people who have spent 12 months in the
non-permanent population (asylum seekers, S status, L permits) are
counted as immigrants when they cross into the permanent one.
type_change_foreign = that component for foreigners alone (the
Ukrainian S-status wave shows here in 2023). Not the same quantity as a
count of arrivals. Eurostat migr_imm1ctz CH reproduces this definition
in 2000-2010 and 2023-2024 only; in 2011-2022 it leaves the type changes
out (lower by exactly type_change_foreign). Never one line.
""", ["year", "immigration_total", "immigration_swiss", "immigration_foreign",
      "emigration_total", "emigration_swiss", "emigration_foreign",
      "net_total", "net_swiss", "net_foreign", "type_change_foreign"], out)


def fertility(bal) -> None:
    o = "je-d-01.04.01.01"
    wb = dam_xlsx(o)
    if wb.sheetnames[:2] != ["2020-2025", "2002-2020"] and \
            not wb.sheetnames[0].startswith("2020-"):
        raise RuntimeError(f"fertility sheets changed: {wb.sheetnames}")
    got: dict[int, tuple] = {}
    for sheet in wb.sheetnames:
        rows = grid(wb, sheet)
        hy = next(i for i, r in enumerate(rows) if r and year_of(r[1]))
        years = {c: year_of(v) for c, v in enumerate(rows[hy]) if c and year_of(v)}
        tot = rows[find_row(rows, lambda s: s == "Total", "births total")]
        tfr = rows[find_row(rows, lambda s: s.startswith("Zusammengefasste"), "TFR")]
        try:
            age = rows[find_row(rows, lambda s: s.startswith("bei der Geburt des Kindes"),
                                "mean age")]
        except RuntimeError:
            age = [None] * len(tot)
        for c, y in years.items():
            v = (num(tot[c]), num(tfr[c]), num(age[c]) if c < len(age) else None)
            if y in got:  # 2020 sits on two sheets
                if got[y][0] != v[0] or abs((got[y][1] or 0) - (v[1] or 0)) > 0.001:
                    raise RuntimeError(f"fertility {y} differs between sheets")
                continue
            got[y] = v
    out = [[y, fmt(got[y][0], 0), fmt(got[y][1], 3), fmt(got[y][2], 2)]
           for y in sorted(got) if y >= Y0]
    for y in range(Y0, LAST_FULL + 1):
        if got[y][0] != bal["total"][y]["births"]:
            raise RuntimeError(f"births {y}: DAM {got[y][0]} vs PxWeb "
                               f"{bal['total'][y]['births']}")
    # Eurostat demo_find twin (research pass): 1.48 / 1.54 / 1.33 / 1.29.
    for y, v in {2008: 1.48, 2015: 1.54, 2023: 1.33, 2024: 1.29}.items():
        expect(f"TFR {y} vs Eurostat", got[y][1], v, 0.006)
    write("demography/fertility_from_bfs.csv", f"""
{ref(o)}: live births (PERSONS), total fertility rate (children per
woman) and mean age of the mother at birth (years), BEVNAT, permanent
resident population. Births equal the PxWeb demographic balance to the
birth (asserted). From 2010 the permanent population includes persons
in the asylum process resident 12+ months (BFS footnote) - a small
definitional seam in the denominator of the TFR.
""", ["year", "live_births", "tfr", "mean_age_mother"], out)


def life_expectancy() -> None:
    o = "je-d-01.04.02.03.01"
    wb = dam_xlsx(o)
    rows = grid(wb, wb.sheetnames[0])
    r2 = next(i for i, r in enumerate(rows) if any("Vollständige jährliche"
                                                   in label(c) for c in r))
    start = next(c for c, v in enumerate(rows[r2]) if "Vollständige jährliche" in label(v))
    sexrow = next(i for i in range(r2, r2 + 6) if label(rows[i][start]) == "Männer")
    yrow = sexrow + 1
    sex, cols = None, {}
    for c in range(start, len(rows[yrow])):
        if label(rows[sexrow][c]) in ("Männer", "Frauen"):
            sex = label(rows[sexrow][c])
        y = year_of(rows[yrow][c])
        if y and sex:
            cols[(sex, y)] = c
    e0 = rows[find_row(rows, lambda s: s == "Bei Geburt", "at birth")]
    e65 = rows[find_row(rows, lambda s: s == "65 Jahren", "at 65")]
    years = sorted({y for _, y in cols})
    out = [[y, fmt(num(e0[cols[("Männer", y)]]), 1), fmt(num(e0[cols[("Frauen", y)]]), 1),
            fmt(num(e65[cols[("Männer", y)]]), 1), fmt(num(e65[cols[("Frauen", y)]]), 1)]
           for y in years if y >= Y0]
    if years[0] != 2000 or years[-1] < LAST_FULL:
        raise RuntimeError(f"annual life tables span {years[0]}-{years[-1]}")
    m, w = float(out[-1][1]), float(out[-1][2])
    if not (80 < m < 84 and 84 < w < 88):
        raise RuntimeError(f"life expectancy {years[-1]}: men {m}, women {w}")
    write("demography/life_expectancy_from_bfs.csv", f"""
{ref(o)}: period life expectancy in YEARS, complete annual life tables
(the right-hand block of the sheet; the 5-year official tables 1948/53 -
1998/2003 to its left are another table type and are not read). By sex
ONLY: BFS publishes no both-sexes figure here, so Eurostat demo_mlexpec's
total (84.2 in 2024) has no BFS twin, and none is computed. 2020 is the
COVID dip.
""", ["year", "e0_men", "e0_women", "e65_men", "e65_women"], out)


PROJ_DB = "px-x-0104000000_101"


def projection() -> None:
    rows = px(PROJ_DB, [
        {"code": "Staatsangehörigkeit (Kategorie)",
         "selection": {"filter": "item", "values": ["0"]}},
        {"code": "Geschlecht", "selection": {"filter": "item", "values": ["0"]}},
        {"code": "Altersklasse", "selection": {"filter": "item", "values": ["0"]}},
        {"code": "Beobachtungseinheit",
         "selection": {"filter": "item", "values": ["0", "1", "2", "3", "4", "5", "6"]}},
    ], "projection_ch", "Szenario-Variante")
    head = rows[0]
    cols = ["Bevölkerungsstand am 1. Januar", "Lebendgeburten", "Todesfälle",
            "Einwanderungen", "Auswanderungen", "Erwerb des Schweizer Bürgerrechts",
            "Bevölkerungsstand am 31. Dezember"]
    for c in cols:
        if c not in head:
            raise RuntimeError(f"projection CSV lacks {c!r}")
    out, ref_pop = [], {}
    for r in rows[1:]:
        if not r:
            continue
        d = dict(zip(head, r))
        sc = d["Szenario-Variante"]
        m = re.search(r"\b([A-E]-\d\d-2025)\b", sc)
        if not m:
            raise RuntimeError(f"projection: no scenario code in {sc!r}")
        y = int(d["Jahr"])
        vals = [num(d[c]) for c in cols]
        out.append([y, m[1], sc] + [fmt(v, 0) for v in vals])
        if m[1] == "A-00-2025":
            ref_pop[y] = vals[-1]
    out.sort(key=lambda r: (r[0], r[1]))   # year first: years stay ordered
    if len({r[1] for r in out}) != 16:
        raise RuntimeError("projection: expected 16 scenario variants")
    if not 9.2e6 < ref_pop[2030] < 9.8e6 or not 10.0e6 < ref_pop[2055] < 11.5e6:
        raise RuntimeError(f"reference scenario magnitudes off: {ref_pop[2030]}, "
                           f"{ref_pop[2055]}")
    write("demography/projection_from_bfs.csv", f"""
BFS STAT-TAB {PROJ_DB}: Szenarien zur Bevölkerungsentwicklung der
Schweiz 2025-2055, all 16 scenario variants (A-00-2025 = reference,
B = 'high', C = 'low', D/E ageing, A-01..A-11 single-assumption
variants), total population, both sexes, all ages. PERSONS. kind =
PROJECTION: every year here is modelled, INCLUDING 2024-2025 (the
model's launch years, which do not equal the observed balance - e.g.
reference pop_1jan 2024 8,960,817 vs observed 31 Dec 2023 8,962,258).
Never splice onto the observed series; draw from the last observed
year as a separate line. Rows are sorted by year, then scenario_code.
""", ["year", "scenario_code", "scenario", "pop_1jan", "births", "deaths",
      "immigration", "emigration", "naturalisations", "pop_31dec"], out)


# ---------------------------------------------------------------- security

PKS_DB = "px-x-1903020100_101"


def homicide(bal) -> None:
    rows = px(PKS_DB, [
        {"code": "Straftat", "selection": {"filter": "item",
                                           "values": ["311.00.T0", "311.00.111.00"]}},
        {"code": "Kanton", "selection": {"filter": "item", "values": ["8100"]}},
        {"code": "Ausführungsgrad", "selection": {"filter": "item",
                                                  "values": ["0", "1", "2"]}},
        {"code": "Aufklärungsgrad", "selection": {"filter": "item", "values": ["0"]}},
    ], "pks_homicide_ch", "Straftat")
    head = rows[0]
    years = {i: int(c) for i, c in enumerate(head) if c.isdigit()}
    got = {}
    for r in rows[1:]:
        if not r:
            continue
        off = "all" if r[0] == "Straftat - Total" else (
            "hom" if r[0].startswith("Tötungsdelikte") else None)
        grade = {"Ausführungsgrad - Total": "total", "Vollendet": "completed",
                 "Versucht": "attempted"}[r[2]]
        if off is None or r[1] != "Schweiz":
            raise RuntimeError(f"PKS: unexpected row {r[:4]}")
        got[(off, grade)] = {y: num(r[i]) for i, y in years.items()}
    out = []
    for y in sorted(years.values()):
        c, a, t = (got[("hom", g)][y] for g in ("completed", "attempted", "total"))
        if c + a != t:
            raise RuntimeError(f"PKS {y}: completed + attempted != total")
        b = bal["total"][y]
        mean_pop = (b["pop_1jan"] + b["pop_31dec"]) / 2
        out.append([y, fmt(c, 0), fmt(a, 0), fmt(t, 0),
                    fmt(c / mean_pop * 1e5, 2), fmt(got[("all", "total")][y], 0)])
    h = got[("hom", "completed")]
    for y, v in {2009: 51, 2015: 57, 2024: 45, 2025: 55}.items():
        expect(f"completed homicides {y}", h[y], v)
    expect("attempted homicides 2025", got[("hom", "attempted")][2025], 251)
    expect("all StGB offences 2025", got[("all", "total")][2025], 554963)
    write("security/homicide_police_from_bfs.csv", f"""
BFS STAT-TAB {PKS_DB}: Polizeiliche Kriminalstatistik (PKS), offence
311.00.111.00 'Tötungsdelikte (Art. 111-113/116)' (intentional homicide:
murder, manslaughter, infanticide), Schweiz, police-RECORDED OFFENCES.
homicide_completed = Vollendet; homicide_attempted = Versucht - FOUR
TIMES the completed count. homicide_total is mostly attempts: draw
completed, label attempts separately, never the total as 'homicides'.
completed_per_100k is COMPUTED here on the mean of the BFS permanent
population on 1 Jan and 31 Dec (same publisher). all_stgb_offences =
every Criminal Code offence recorded. PKS was rebuilt in 2009; earlier
figures are not comparable and are not published in this table.
Eurostat crim_off_cat ICCS0101 CH reproduces the completed count.
""", ["year", "homicide_completed", "homicide_attempted", "homicide_total",
      "homicide_completed_per_100k", "all_stgb_offences"], out)


def firearm_deaths() -> None:
    o = "je-d-14.03.04.01.13"
    wb = dam_xlsx(o)
    rows = grid(wb, wb.sheetnames[0])
    hy = next(i for i, r in enumerate(rows) if r and year_of(r[1]) == 1995)
    years = {c: year_of(v) for c, v in enumerate(rows[hy]) if c and year_of(v)}

    def triple(start: str) -> list[dict]:
        i = find_row(rows, lambda s: s.startswith(start), start)
        if label(rows[i + 1][0]) != "Männer" or label(rows[i + 2][0]) != "Frauen":
            raise RuntimeError(f"firearm table: {start} not followed by sexes")
        return [{y: num(rows[j][c]) for c, y in years.items()} for j in (i, i + 1, i + 2)]

    tot, sui, oth = (triple(s) for s in ("Anzahl Schusswaffentote",
                                         "Schusswaffen-Suizide",
                                         "Andere Schusswaffen"))
    out = []
    for y in sorted(years.values()):
        for k in range(3):
            if tot[k][y] != sui[k][y] + oth[k][y]:
                raise RuntimeError(f"firearm {y}: total != suicides + other")
        if tot[0][y] != tot[1][y] + tot[2][y]:
            raise RuntimeError(f"firearm {y}: total != men + women")
        if y >= Y0:
            out.append([y] + [fmt(v[y], 0) for v in tot + sui + oth])
    if max(years.values()) < LAST_FULL - 1:
        raise RuntimeError("firearm table older than the last CoD year")
    write("security/firearm_deaths_from_bfs.csv", f"""
{ref(o)}: Schusswaffentodesfälle, cause-of-death statistics (CoD),
DEATHS of residents, by sex. firearm_suicides + other_firearm_deaths =
firearm_deaths (asserted). In Switzerland nearly nine in ten firearm
deaths are SUICIDES (militia rifles kept at home): this is a SUPERSET
of firearm homicide and is never drawn as the peer of a homicide count.
'Other' bundles homicides, accidents and undetermined intent. Last year
2024 (CoD lags one year).
""", ["year", "firearm_deaths", "firearm_deaths_men", "firearm_deaths_women",
      "firearm_suicides", "firearm_suicides_men", "firearm_suicides_women",
      "other_firearm_deaths", "other_firearm_deaths_men",
      "other_firearm_deaths_women"], out)


# ---------------------------------------------------------------- mortality

def cod_sheet(wb, sheet: str) -> dict[str, dict[int, float | None]]:
    rows = grid(wb, sheet)
    hy = next(i for i, r in enumerate(rows) if r and year_of(r[1]) == 1970)
    years = {c: year_of(v) for c, v in enumerate(rows[hy]) if c and year_of(v)}
    start = find_row(rows, lambda s: s == "Alle Todesursachen", "all causes")
    out = {}
    for r in rows[start:]:
        lab = label(r[0])
        if not lab or lab == "davon:":
            continue
        if re.match(r"\d\)", lab):
            break  # footnotes
        name = re.sub(r"\s*\d\)$", "", lab)
        out[name] = {y: num(r[c]) for c, y in years.items()}
    return out


def causes_of_death() -> None:
    om, ow = "je-d-14.03.04.01.03", "je-d-14.03.04.01.04"
    men, women = dam_xlsx(om), dam_xlsx(ow)
    for wb in (men, women):
        if "Anzahl Fälle" not in wb.sheetnames or "Sterbeziffern" not in wb.sheetnames:
            raise RuntimeError(f"CoD workbook sheets changed: {wb.sheetnames}")
    nm, nw = cod_sheet(men, "Anzahl Fälle"), cod_sheet(women, "Anzahl Fälle")
    rm, rw = cod_sheet(men, "Sterbeziffern"), cod_sheet(women, "Sterbeziffern")
    only = {"Prostata": "men", "Gebärmutterhals": "women"}
    if set(nm) ^ set(nw) != set(only) or "Prostata" not in nm:
        raise RuntimeError(f"CoD causes differ by sex: {set(nm) ^ set(nw)}")
    order = list(nm) + [c for c in nw if c not in nm]
    out = []
    years = sorted(nm["Alle Todesursachen"])
    for cause in order:
        for y in years:
            if y < Y0:
                continue
            a, b = nm.get(cause, {}).get(y), nw.get(cause, {}).get(y)
            if cause in only:  # a sex-specific cause: the total is that sex
                tot = a if only[cause] == "men" else b
            else:
                tot = a + b if a is not None and b is not None else None
            out.append([y, cause, fmt(a, 0), fmt(b, 0), fmt(tot, 0),
                        fmt(rm.get(cause, {}).get(y), 1),
                        fmt(rw.get(cause, {}).get(y), 1)])
    out.sort(key=lambda r: (r[0], order.index(r[1])))
    expect("women all causes 2024", nw["Alle Todesursachen"][2024], 36728)
    expect("women dementia 2024", nw["Demenz"][2024], 4389)
    expect("women suicide 2024", nw["Suizid"][2024], 277)
    if years[-1] != 2024:
        print(f"    NOTE: causes of death now end in {years[-1]}")
    write("mortality/causes_of_death_from_bfs.csv", f"""
{ref(om)} (men) and {ref(ow)} (women): Sterbefälle und Sterbeziffern
wichtiger Todesursachen, residents. deaths_* in PERSONS; deaths_total =
men + women, summed here (blank if either is a hole); Prostata appears
in the men's file only and Gebärmutterhals in the women's only, so their
total is that sex and the other sex's columns are blank. rate_* =
age-standardised death rates per 100,000 by sex as BFS publishes them;
they cannot be summed, so there is no total rate. Long format, cause
labels as BFS writes them (German). Years: 2000 onward (the file has
1970, 1980, 1990-). BREAKS (footnotes): ICD-8 -> ICD-10 between 1994 and
1995; 'Unfälle insgesamt' include accidental poisoning by psychotropic
substances from 2006; SUIZID INCLUDES ASSISTED SUICIDE UNTIL 2008 and
excludes it from 2009 (it is coded to the underlying disease and counted
in assisted_suicide_from_bfs.csv) - the 2008/09 fall is classification.
COVID-19 is blank before 2020 (not applicable). No homicide row: BFS
publishes 'Unfälle und Gewalteinwirkungen' only.
""", ["year", "cause", "deaths_men", "deaths_women", "deaths_total",
      "rate_men_per100k_std", "rate_women_per100k_std"], out)


def assisted_suicide() -> None:
    o = "je-d-14.03.04.01.14"
    wb = dam_xlsx(o)
    rows = grid(wb, wb.sheetnames[0])
    hy = next(i for i, r in enumerate(rows) if r and year_of(r[1]) == 2003)
    years = {c: year_of(v) for c, v in enumerate(rows[hy]) if c and year_of(v)}
    t = next(i for i, r in enumerate(rows) if i > hy
             and label(r[0]).startswith("Assistierter Suizid"))
    m = find_row(rows, lambda s: s == "Männer", "men")
    w = find_row(rows, lambda s: s == "Frauen", "women")
    if not (label(rows[m + 1][0]).startswith("unter 65")
            and label(rows[m + 2][0]).startswith("über 65")):
        raise RuntimeError("assisted suicide: age rows changed")
    pick = [t, m, w, m + 1, m + 2, w + 1, w + 2]
    out = []
    for c, y in sorted(years.items(), key=lambda kv: kv[1]):
        v = [num(rows[i][c]) for i in pick]
        if v[0] != v[1] + v[2] or v[1] != v[3] + v[4] or v[2] != v[5] + v[6]:
            raise RuntimeError(f"assisted suicide {y}: parts do not add up")
        out.append([y] + [fmt(x, 0) for x in v])
    expect("assisted suicide 2024", float(out[-1][1]) if out[-1][0] == 2024 else None, 1774)
    write("mortality/assisted_suicide_from_bfs.csv", f"""
{ref(o)}: Assistierter Suizid nach Geschlecht und Alter, residents only,
DEATHS. Published from 2003. Until 2008 these same deaths are ALSO
counted inside 'Suizid' in causes_of_death_from_bfs.csv; from 2009 they
are not. NEVER ADD this to suicide: draw it beside suicide, with the
2008/09 seam annotated. Parts asserted to add up (men + women; under/
over 65). 'über 65' as BFS labels it (65 and over).
""", ["year", "assisted_suicides", "men", "women", "men_under65",
      "men_65plus", "women_under65", "women_65plus"], out)


# ---------------------------------------------------------------- migration

PERMIT_DB = "px-x-0103010000_101"


def stock_by_permit(bal) -> None:
    rows = px(PERMIT_DB, [
        {"code": "Kanton", "selection": {"filter": "item", "values": ["8100"]}},
        {"code": "Anwesenheitsbewilligung", "selection": {"filter": "all", "values": ["*"]}},
        {"code": "Geschlecht", "selection": {"filter": "item", "values": ["-99999"]}},
        {"code": "Altersklasse", "selection": {"filter": "item", "values": ["-99999"]}},
        {"code": "Staatsangehörigkeit", "selection": {"filter": "item", "values": ["-99999"]}},
    ], "permit_ch", "Jahr", timeout=900)
    head = rows[0]
    if "Anwesenheitsbewilligung" not in head:
        raise RuntimeError(f"permit CSV header {head}")
    vcol = len(head) - 1
    got: dict[tuple, float] = {}
    for r in rows[1:]:
        if not r:
            continue
        d = dict(zip(head, r))
        key = (int(d["Jahr"]), d["Bevölkerungstyp"], d["Anwesenheitsbewilligung"])
        got[key] = num(r[vcol])
    years = sorted({k[0] for k in got})
    perm, nonp = "Ständige Wohnbevölkerung", "Nichtständige Wohnbevölkerung"
    total = "Anwesenheitsbewilligung - Total"
    permits = sorted({k[2] for k in got} - {total})
    out = []
    for y in years:
        for t in (perm, nonp):
            parts = sum(got.get((y, t, p)) or 0 for p in permits)
            if got.get((y, t, total)) != parts:
                raise RuntimeError(f"permits {y} {t}: total != sum of permits")
            for p in [total] + permits:
                v = got.get((y, t, p))
                if p == "Schutzbedürftiger (S)" and y < 2022 and v == 0:
                    v = None  # S status did not exist before March 2022
                out.append([y, "permanent" if t == perm else "non_permanent",
                            p, fmt(v, 0)])
        # STATPOP's year-end stock is the balance's NEXT 1 January stock
        # (equal to its 31 December everywhere except the 2010 register
        # seam: 31 Dec 2010 7,864,012 vs 1 Jan 2011 7,870,134).
        b = bal["total"]
        want = b[y + 1]["pop_1jan"] if y + 1 in b else b[y]["pop_31dec"]
        if got[(y, perm, total)] != want:
            raise RuntimeError(f"permits {y}: permanent total {got[(y, perm, total)]}"
                               f" != balance stock {want}")
    for p, v in {"Schweizer": 6599008, "Aufenthalter (B)": 955652,
                 "Niedergelassener (C)": 1406673, "Asylsuchender (N)": 11092,
                 "Schutzbedürftiger (S)": 65260}.items():
        expect(f"permanent {p} 2025", got.get((2025, perm, p)), v)
    expect("non-permanent N 2025", got.get((2025, nonp, "Asylsuchender (N)")), 7771)
    expect("non-permanent S 2025", got.get((2025, nonp, "Schutzbedürftiger (S)")), 5946)
    write("migration/stock_by_permit_from_bfs.csv", f"""
BFS STAT-TAB {PERMIT_DB} (STATPOP): resident population on 31 December
by population type and residence permit (Anwesenheitsbewilligung),
Schweiz, both sexes, all ages, all citizenships. PERSONS. Long format;
permit labels as BFS writes them: Schweizer, B (Aufenthalter), C
(Niedergelassener), Ci, F (vorläufig aufgenommen), L (Kurzaufenthalter),
N (Asylsuchender), S (Schutzbedürftiger), diplomats, unassigned.
2010 is the first STATPOP register year: its permanent total equals the
balance's 1 Jan 2011 stock (7,870,134), not the balance's 31 Dec 2010
(7,864,012); from 2011 on the two agree to the person (asserted).
PERMANENT and NON-PERMANENT are two populations: S-status and asylum
seekers move into the permanent one after 12 months (65,260 S-status
holders were PERMANENT at end-2025). Permanent total = the demographic
balance's year-end stock (asserted). S status was created in March
2022: its rows before 2022 are blank (published as 0, not applicable).
Counts of persons by permit; SEM's ZEMIS register counts differ by
definition and are never drawn on the same line.
""", ["year", "population_type", "permit", "persons"], out)


# ---------------------------------------------------------------- cross-checks

def cross_checks(b1: dict[int, float], bal) -> None:
    print("  cross-checks against a second publisher")
    url = ("https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/"
           "data/nama_10_gdp?geo=CH&unit=CP_MNAC&na_item=B1GQ")

    def es_check(blob: bytes) -> None:
        d = json.loads(blob)
        if not d.get("value"):
            raise RuntimeError("Eurostat nama_10_gdp CH: empty cube")

    d = json.loads(cached("bfs_xcheck_estat_nama_10_gdp.json",
                          lambda: _get(url, "xcheck"), es_check))
    idx = d["dimension"]["time"]["category"]["index"]
    es = {int(t): d["value"][str(i)] for t, i in idx.items()
          if str(i) in d["value"]}
    for y in range(2006, 2023):
        gap = es[y] / b1[y] - 1
        if abs(gap) > 0.001:
            raise RuntimeError(f"GDP {y}: Eurostat {es[y]} vs BFS {b1[y]:.1f} "
                               f"({gap:+.2%}) where the vintages should agree")
    print("    GDP 2006-2022: BFS = Eurostat within 0.1% (same vintage)")
    for y in range(2023, LAST_FULL + 1):
        if y in es:
            print(f"    GDP {y}: BFS {b1[y] / 1e3:,.1f}bn vs Eurostat "
                  f"{es[y] / 1e3:,.1f}bn ({es[y] / b1[y] - 1:+.2%}) "
                  "- vintage gap, never mixed")

    url = ("https://api.worldbank.org/v2/country/CHE/indicator/SP.POP.TOTL"
           "?format=json&per_page=100")

    def wb_check(blob: bytes) -> None:
        d = json.loads(blob)
        if not isinstance(d, list) or len(d) < 2 or not d[1]:
            raise RuntimeError("World Bank SP.POP.TOTL CHE: no data")

    rows = json.loads(cached("bfs_xcheck_wb_sp_pop_totl.json",
                             lambda: _get(url, "xcheck"), wb_check))[1]
    wbp = {int(r["date"]): r["value"] for r in rows if r["value"] is not None}
    for y in (2006, 2015, max(wbp)):
        b = bal["total"][y]
        mean = (b["pop_1jan"] + b["pop_31dec"]) / 2
        gap = wbp[y] / mean - 1
        if abs(gap) > 0.005:
            raise RuntimeError(f"population {y}: World Bank {wbp[y]} vs BFS "
                               f"mean {mean:.0f} ({gap:+.2%})")
        print(f"    population {y}: World Bank {wbp[y]:,} vs BFS mean "
              f"{mean:,.0f} ({gap:+.2%})")


def main() -> None:
    print("BFS — DAM assets and STAT-TAB PxWeb")
    b1 = gdp()
    cpi()
    unemployment()
    unemployment_rate()
    employment()
    wages()
    gini()
    poverty()
    arrears()
    bal = balance_rows()
    balance(bal)
    flows(bal)
    fertility(bal)
    life_expectancy()
    projection()
    homicide(bal)
    firearm_deaths()
    causes_of_death()
    assisted_suicide()
    stock_by_permit(bal)
    cross_checks(b1, bal)
    print("Done. Diff the sidecars before adopting.")


if __name__ == "__main__":
    main()
